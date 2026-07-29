"""Gom 20 bài bổ sung Toán 5 thành bản đọc và bản DBJSON để import."""

from __future__ import annotations

import base64
import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "doc"
ASSET_OUT = OUT / "assets" / "grade5-part2-import"
READ_DOCX = OUT / "ngan_hang_cau_hoi_lop_5_phan_2.docx"
IMPORT_DOCX = OUT / "ngan_hang_cau_hoi_lop_5_phan_2_import.docx"
IMPORT_TEX = ROOT / "data" / "grade5_question_bank_part2.tex"
LESSONS = (61, 66, 67, 68, 71, 73, 75, 76, 77, 78, 82, 83, 84, 85, 86, 87, 88, 89, 90, 91)
TITLES = {
    61: "Ôn tập chung", 66: "Ôn tập về thời gian", 67: "Ôn tập về chuyển động",
    68: "Ôn tập về các đơn vị đo thời gian", 71: "Ôn tập về các phép tính với số đo thời gian",
    73: "Vận tốc", 75: "Luyện tập", 76: "Luyện tập chung",
    77: "Em ôn lại những gì đã học", 78: "Em vui học Toán",
    82: "Ôn tập về số tự nhiên và các phép tính với số tự nhiên",
    83: "Ôn tập về phân số và các phép tính với phân số",
    84: "Ôn tập về số thập phân và các phép tính với số thập phân",
    85: "Ôn tập về tỉ số, tỉ số phần trăm", 86: "Ôn tập về hình học",
    87: "Ôn tập về đo lường", 88: "Ôn tập về một số yếu tố thống kê và xác suất",
    89: "Em ôn lại những gì đã học", 90: "Em vui học Toán", 91: "Ôn tập chung",
}
MANUAL = (
    ("grade5_b61_b66_b67.json", "Ngan_hang_cau_hoi_Toan_5_Bai_61_66_67_60_cau.docx"),
    ("grade5_b68_b71_b73.json", "Ngan_hang_cau_hoi_Toan_5_Bai_68_71_73_66_cau.docx"),
    ("grade5_b75_b76_b77.json", "Ngan_hang_cau_hoi_Toan_5_Bai_75_76_77_66_cau.docx"),
    ("grade5_b78_b82_b83.json", "Ngan_hang_cau_hoi_Toan_5_Bai_78_82_83_66_cau.docx"),
)
LEGACY = (
    ("Ngan_hang_cau_hoi_Toan_5_Canh_Dieu_Bai_82_83_84_Da_sua.docx", {84}, "classic"),
    ("Ngan_hang_cau_hoi_Toan_5_Canh_Dieu_Bai_85_86_87_Da_sua_anh.docx", {85, 86, 87}, "classic"),
    ("Ngan_hang_cau_hoi_Toan_5_Canh_Dieu_Bai_88.docx", {88}, "classic"),
    ("Ngan_hang_cau_hoi_Toan_5_Bai_89_90_91_22_cau.docx", {89, 90, 91}, "modern"),
)


def clean(text):
    return re.sub(r"\s+", " ", text or "").strip()


def image_blobs(paragraph, document):
    result = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if rid and rid in document.part.rels:
            part = document.part.rels[rid].target_part
            suffix = Path(str(part.partname)).suffix or ".png"
            result.append((part.blob, suffix))
    return result


def save_images(blobs, lesson, number):
    paths = []
    for index, (blob, suffix) in enumerate(blobs, 1):
        path = ASSET_OUT / f"g5-l{lesson:03d}-q{number:03d}-{index:02d}{suffix}"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(blob)
        paths.append(str(path.resolve()))
    return paths


def finish(records, question):
    if not question:
        return
    question["prompt"] = clean(question.get("prompt"))
    question["explanation"] = clean(question.get("explanation"))
    question["answer"] = clean(question.get("answer"))
    question["choices"] = question.get("choices", [])
    if question["prompt"] and question["lesson"] in LESSONS:
        records.append(question)


def parse_classic(path, wanted):
    doc = Document(path)
    records, question, lesson = [], None, None
    pending_blobs = []
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            p = Paragraph(child, doc)
            text, style = clean(p.text), p.style.name
            match = re.search(r"\bBÀI\s+(\d+)\b", text, re.I)
            if style.startswith("Heading 1") and match:
                finish(records, question); question = None
                lesson = int(match.group(1))
                continue
            if style == "Question" and re.match(r"Câu\s+\d+\.", text, re.I):
                finish(records, question)
                number = int(re.match(r"Câu\s+(\d+)", text, re.I).group(1))
                question = {"lesson": lesson, "number": number,
                            "prompt": re.sub(r"^Câu\s+\d+\.\s*", "", text, flags=re.I),
                            "choices": [], "answer": "", "explanation": "", "images": []}
            elif question and style == "Option" and re.match(r"^[A-D]\.", text):
                question["choices"].append({"key": text[0], "text": clean(text[2:])})
            if question:
                pending_blobs.extend(image_blobs(p, doc))
                if pending_blobs:
                    question["images"].extend(save_images(pending_blobs, lesson, question["number"]))
                    pending_blobs = []
        elif isinstance(child, CT_Tbl) and question:
            table = Table(child, doc)
            tokens = [clean(value) for value in table._tbl.xpath(".//w:t/text()") if clean(value)]
            upper = [value.upper() for value in tokens]
            if "ĐÁP ÁN" in upper:
                index = upper.index("ĐÁP ÁN")
                if index + 1 < len(tokens):
                    question["answer"] = tokens[index + 1]
            if "LỜI GIẢI" in upper:
                index = upper.index("LỜI GIẢI")
                question["explanation"] = clean(" ".join(tokens[index + 1:]))
            for row in table.rows:
                cells = [clean(cell.text) for cell in row.cells]
                if not cells:
                    continue
                label = cells[0].upper()
                value = cells[-1]
                if "ĐÁP ÁN" in label:
                    question["answer"] = value
                elif "LỜI GIẢI" in label:
                    question["explanation"] = value
                for cell in row.cells:
                    for p in cell.paragraphs:
                        blobs = image_blobs(p, doc)
                        if blobs:
                            question["images"].extend(save_images(blobs, lesson, question["number"]))
    finish(records, question)
    return [q for q in records if q["lesson"] in wanted]


def parse_modern(path, wanted):
    doc = Document(path)
    records, question, lesson = [], None, None
    for p in doc.paragraphs:
        text, style = clean(p.text), p.style.name
        match = re.search(r"\bBÀI\s+(\d+)\b", text, re.I)
        if style.startswith("Heading 1") and match:
            finish(records, question); question = None
            lesson = int(match.group(1)); continue
        if style.startswith("Heading 2") and re.match(r"Câu\s+\d+\.", text, re.I):
            finish(records, question)
            number = int(re.match(r"Câu\s+(\d+)", text, re.I).group(1))
            question = {"lesson": lesson, "number": number,
                        "prompt": re.sub(r"^Câu\s+\d+\.\s*", "", text, flags=re.I),
                        "choices": [], "answer": "", "explanation": "", "images": []}
        elif question and re.match(r"^[A-D]\.", text):
            question["choices"].append({"key": text[0], "text": clean(text[2:])})
        elif question and text.startswith("Đáp án:"):
            question["answer"] = clean(text.split(":", 1)[1])
        elif question and text.startswith("Lời giải:"):
            question["explanation"] = clean(text.split(":", 1)[1])
        if question:
            blobs = image_blobs(p, doc)
            if blobs:
                question["images"].extend(save_images(blobs, lesson, question["number"]))
    finish(records, question)
    return [q for q in records if q["lesson"] in wanted]


def manual_image_map(path):
    doc = Document(path)
    result = {}
    lesson = number = None
    for child in doc.element.body.iterchildren():
        if not isinstance(child, CT_P):
            continue
        paragraph = Paragraph(child, doc)
        text, style = clean(paragraph.text), paragraph.style.name
        lesson_match = re.match(r"Bài\s+(\d+)\.", text, re.I)
        question_match = re.match(r"Câu\s+(\d+)\.", text, re.I)
        if lesson_match:
            lesson = int(lesson_match.group(1))
            number = None
        elif style.startswith("Heading 2") and question_match:
            number = int(question_match.group(1))
        elif lesson in LESSONS and number is not None:
            blobs = image_blobs(paragraph, doc)
            if blobs:
                result.setdefault((lesson, number), []).extend(save_images(blobs, lesson, number))
    return result


def manual_records():
    records = []
    for filename, source_docx in MANUAL:
        payload = json.loads((ROOT / "data" / "manual_question_batches" / filename).read_text(encoding="utf-8"))
        image_map = manual_image_map(OUT / source_docx)
        for item in payload["questions"]:
            images = image_map.get((item["lesson"], item["number"]), [])
            choices = []
            for choice in item.get("choices", []):
                match = re.match(r"^([A-D])\.\s*(.*)", choice)
                if match:
                    choices.append({"key": match.group(1), "text": clean(match.group(2))})
            records.append({"lesson": item["lesson"], "number": item["number"],
                            "prompt": item["question"], "choices": choices,
                            "answer": item["answer"], "explanation": item["explanation"], "images": images,
                            "response": item.get("response", "")})
    return records


def all_records():
    records = manual_records()
    for filename, wanted, kind in LEGACY:
        path = OUT / filename
        records.extend(parse_classic(path, wanted) if kind == "classic" else parse_modern(path, wanted))
    records.sort(key=lambda q: (LESSONS.index(q["lesson"]), q["number"]))
    return records


def correct_key(record):
    match = re.match(r"^\s*([A-D])(?:\.|\b)", record["answer"])
    return match.group(1) if match else ""


def payload(record, bank_index):
    lesson, key = record["lesson"], correct_key(record)
    multiple = len(record["choices"]) == 4 and bool(key)
    images = [
        {
            "id": f"G5-L{lesson:03d}-P2-Q{bank_index:03d}-IMG-{i:02d}",
            "source_path": path,
            "package_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            "url": "",
            "alt_text": f"Hình minh họa Bài {lesson} Câu {record['number']}",
            "width_percent": 100,
        }
        for i, path in enumerate(record["images"], 1)
    ]
    return {
        "external_id": f"G5-L{lesson:03d}-P2-Q{bank_index:03d}", "grade": 5,
        "lesson_number": lesson, "lesson_title": TITLES[lesson],
        "question_type": "MULTIPLE_CHOICE" if multiple else "FILL_IN_THE_BLANK",
        "difficulty": "EASY" if record["number"] <= 8 else "MEDIUM" if record["number"] <= 16 else "HARD",
        "layout_template": "SPLIT_HORIZONTAL_LEFT_IMAGE" if images else "STACK_VERTICAL",
        "content": {"text": record["prompt"], "images": images},
        "choices": [{**choice, "images": []} for choice in record["choices"]] if multiple else [],
        "correct_answer": key if multiple else record["answer"],
        "explanation": {"text": record["explanation"], "images": []},
        "misconceptions": [], "source_kind": "part2_reviewed", "source_file": "ngan_hang_cau_hoi_lop_5_phan_2",
    }


def configure(doc):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.7)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 23), ("Heading 1", 17), ("Heading 2", 12)):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(31, 78, 121)


def cover(doc, title, records):
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"20 bài học • {len(records)} câu hỏi • Có đáp án, lời giải và hình minh họa").bold = True
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Bài học", "Số câu"
    for lesson in LESSONS:
        cells = table.add_row().cells
        cells[0].text = f"Bài {lesson}. {TITLES[lesson]}"
        cells[1].text = str(sum(q["lesson"] == lesson for q in records))
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_page_break()


def write_read_doc(records):
    doc = Document(); configure(doc)
    cover(doc, "NGÂN HÀNG CÂU HỎI LỚP 5 - PHẦN 2", records)
    counters = {}
    for record in records:
        lesson = record["lesson"]
        if lesson not in counters:
            if counters:
                doc.add_page_break()
            doc.add_heading(f"Bài {lesson}. {TITLES[lesson]}", level=1)
            counters[lesson] = 0
        counters[lesson] += 1
        doc.add_heading(f"Câu {record['number']}", level=2)
        p = doc.add_paragraph(); p.add_run(record["prompt"]).bold = True
        for path in record["images"]:
            if Path(path).exists():
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Cm(13.2))
        if record["choices"]:
            doc.add_paragraph("\n".join(f"{c['key']}. {c['text']}" for c in record["choices"]))
        elif record.get("response"):
            doc.add_paragraph(f"Phần trả lời: {record['response']}")
        p = doc.add_paragraph(); r = p.add_run(f"Đáp án: {record['answer']}"); r.bold = True; r.font.color.rgb = RGBColor(0, 112, 60)
        p = doc.add_paragraph(); p.add_run("Lời giải: ").bold = True; p.add_run(record["explanation"])
    doc.save(READ_DOCX)


def write_import_doc(records):
    doc = Document(); configure(doc)
    cover(doc, "DỮ LIỆU IMPORT NGÂN HÀNG CÂU HỎI LỚP 5 - PHẦN 2", records)
    lines = []
    for lesson in LESSONS:
        doc.add_heading(f"Bài {lesson}. {TITLES[lesson]}", level=1)
        lesson_records = [q for q in records if q["lesson"] == lesson]
        for index, record in enumerate(lesson_records, 1):
            data = payload(record, index)
            encoded = base64.b64encode(json.dumps(data, ensure_ascii=False, separators=(",", ":")).encode()).decode()
            line = f"% DBJSON {encoded}"
            doc.add_paragraph(line)
            lines.append(line)
    media = doc.add_paragraph()
    media.style = doc.styles["Normal"]
    for source_path in sorted({path for record in records for path in record["images"]}):
        run = media.add_run()
        run.font.hidden = True
        run.add_picture(source_path, width=Pt(1))
    doc.save(IMPORT_DOCX)
    IMPORT_TEX.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    if ASSET_OUT.exists():
        shutil.rmtree(ASSET_OUT)
    ASSET_OUT.mkdir(parents=True, exist_ok=True)
    records = all_records()
    counts = {lesson: sum(q["lesson"] == lesson for q in records) for lesson in LESSONS}
    if len(counts) != 20 or any(value < 20 for value in counts.values()):
        raise RuntimeError(f"Độ phủ chưa đạt: {counts}")
    if any(not q["answer"] or not q["explanation"] for q in records):
        raise RuntimeError("Có câu thiếu đáp án hoặc lời giải")
    write_read_doc(records)
    write_import_doc(records)
    print(json.dumps({"questions": len(records), "lessons": counts, "read": str(READ_DOCX),
                      "import_docx": str(IMPORT_DOCX), "import_tex": str(IMPORT_TEX)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
