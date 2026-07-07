import json
import math
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_grade2_theory_docx_detailed as base  # noqa: E402


OUTPUT_DIR = ROOT / "output" / "doc"
IMAGE_DIR = OUTPUT_DIR / "grade3_theory_images_detailed"
BLUEPRINT_PATH = ROOT / "content-theory" / "grade-3-theory-blueprint-detailed.json"
OUTPUT_PATH = OUTPUT_DIR / "ly-thuyet-lop-3-codex-chi-tiet.docx"
FALLBACK_OUTPUT_PATH = OUTPUT_DIR / "ly-thuyet-lop-3-codex-chi-tiet-fixed.docx"


CHAPTERS = [
    {
        "title": "Bảng nhân, bảng chia",
        "lessons": [
            "Mi-li-mét",
            "Bảng nhân 3",
            "Bảng nhân 4",
            "Bảng nhân 6",
            "Gấp một số lên một số lần",
            "Bảng nhân 7",
            "Bảng nhân 8",
            "Bảng nhân 9",
            "Gam",
            "Bảng chia 3",
            "Bảng chia 4",
            "Bảng chia 6",
            "Giảm một số đi một số lần",
            "Bảng chia 7",
            "Bảng chia 8",
            "Bảng chia 9",
            "Một phần hai, một phần tư",
            "Một phần ba, một phần năm, một phần sáu",
            "Một phần bảy. Một phần tám. Một phần chín",
        ],
    },
    {
        "title": "Nhân, chia các số trong phạm vi 1 000",
        "lessons": [
            "Nhân số tròn chục với số có một chữ số",
            "Nhân với số có một chữ số không nhớ",
            "Phép chia hết, phép chia có dư",
            "Chia số tròn chục, tròn trăm cho số có một chữ số",
            "Chia cho số có một chữ số",
            "So sánh số lớn gấp mấy lần số bé",
            "Giải bài toán có đến hai bước tính",
            "Làm quen với biểu thức số",
            "Tính giá trị của biểu thức số",
            "Tính giá trị của biểu thức số (tiếp theo)",
            "Tính giá trị của biểu thức số (tiếp theo) trang 93",
            "Mi-li-lít",
            "Nhiệt độ",
            "Góc vuông, góc không vuông",
            "Hình tam giác, hình tứ giác",
            "Chu vi hình tam giác, chu vi hình tứ giác",
            "Hình chữ nhật",
            "Hình vuông",
            "Chu vi hình chữ nhật, chu vi hình vuông",
        ],
    },
    {
        "title": "Các số trong phạm vi 100 000",
        "lessons": [
            "Các số trong phạm vi 10 000",
            "Các số trong phạm vi 10 000 (tiếp theo)",
            "Làm quen với chữ số La Mã",
            "Các số trong phạm vi 100 000",
            "Các số trong phạm vi 100 000 (tiếp theo)",
            "So sánh các số trong phạm vi 100 000",
            "Điểm ở giữa, trung điểm của đoạn thẳng",
            "Hình tròn, tâm, đường kính, bán kính",
            "Làm tròn số đến hàng chục, hàng trăm",
            "Làm tròn số đến hàng nghìn, hàng chục nghìn",
            "Khối hộp chữ nhật, khối lập phương",
            "Tháng - năm trang 46",
        ],
    },
    {
        "title": "Cộng, trừ, nhân, chia trong phạm vi 100 000",
        "lessons": [
            "Phép cộng trong phạm vi 100 000",
            "Phép trừ trong phạm vi 100 000",
            "Tiền Việt Nam",
            "Nhân với số có một chữ số không nhớ trong phạm vi 100 000",
            "Nhân với số có một chữ số có nhớ",
            "Chia cho số có một chữ số trong phạm vi 100 000",
            "Chia cho số có một chữ số trong phạm vi 100 000 (tiếp theo)",
            "Chia cho số có một chữ số trong phạm vi 100 000 (tiếp theo) trang 77",
            "Tìm thành phần chưa biết của phép tính",
            "Tìm thành phần chưa biết của phép tính (tiếp theo)",
            "Diện tích một hình",
            "Đơn vị đo diện tích. Xăng-ti-mét vuông",
            "Diện tích hình chữ nhật, diện tích hình vuông",
            "Thu thập, phân loại, ghi chép số liệu thống kê",
            "Bảng số liệu thống kê",
            "Khả năng xảy ra của một sự kiện",
        ],
    },
]


def ntext(text):
    return base.norm(text)


def table_number(name, default=3):
    for value in [9, 8, 7, 6, 5, 4, 3, 2]:
        if str(value) in name:
            return value
    return default


def lesson_kind(name):
    n = ntext(name)
    if "mi-li-met" in n:
        return "mm"
    if "gam" == n:
        return "gram"
    if "mi-li-lit" in n:
        return "ml"
    if "nhiet do" in n:
        return "temperature"
    if "bang nhan" in n:
        return "mul_table"
    if "bang chia" in n:
        return "div_table"
    if "gap mot so" in n:
        return "times_more"
    if "giam mot so" in n:
        return "times_less"
    if "mot phan" in n:
        return "fraction"
    if "nhan so tron chuc" in n:
        return "multiply_tens"
    if "nhan voi so co mot chu so" in n and "co nho" in n:
        return "multiply_carry"
    if "nhan voi so co mot chu so" in n and "100 000" in n:
        return "multiply_large"
    if "chia cho so co mot chu so" in n and "100 000" in n:
        return "divide_large"
    if "nhan so tron chuc" in n or "nhan voi so co mot chu so" in n:
        return "multiply"
    if "phep chia het" in n or "chia so tron" in n or n.startswith("chia cho so"):
        return "divide"
    if "so sanh so lon gap" in n:
        return "compare_times"
    if "hai buoc tinh" in n:
        return "two_step_problem"
    if "bieu thuc so" in n or "tinh gia tri" in n:
        return "expression"
    if "phep cong" in n:
        return "add_large"
    if "phep tru" in n:
        return "sub_large"
    if "goc vuong" in n:
        return "angle"
    if "tam giac" in n and "chu vi" not in n:
        return "triangle_quad"
    if "chu vi" in n:
        return "perimeter"
    if "hinh chu nhat" in n:
        return "rectangle"
    if "hinh vuong" in n:
        return "square"
    if "10 000" in n:
        return "number_10000"
    if "100 000" in n and "so sanh" not in n:
        return "number_100000"
    if "la ma" in n:
        return "roman"
    if "so sanh cac so" in n:
        return "compare_large"
    if "diem o giua" in n:
        return "midpoint"
    if "hinh tron" in n:
        return "circle"
    if "lam tron" in n:
        return "rounding"
    if "khoi hop" in n:
        return "box_cube"
    if "thang - nam" in n:
        return "month_year"
    if "tien viet nam" in n:
        return "money"
    if "tim thanh phan" in n:
        return "unknown"
    if "dien tich" in n:
        return "area"
    if "thu thap" in n:
        return "statistics"
    if "bang so lieu" in n:
        return "data_table"
    if "kha nang xay ra" in n:
        return "probability"
    return "expression"


def lesson_spec(name):
    k = lesson_kind(name)
    table = table_number(name)
    specs = {
        "mm": ("Nhận biết mi-li-mét là đơn vị đo độ dài nhỏ hơn xăng-ti-mét.", "1 cm = 10 mm.", "Đoạn thẳng dài 30 mm thì dài 3 cm.", "3 cm"),
        "gram": ("Nhận biết gam là đơn vị đo khối lượng nhỏ.", "Gói bánh 200 g nhẹ hơn túi gạo 1 kg.", "500 g + 300 g = 800 g.", "800 g"),
        "ml": ("Nhận biết mi-li-lít là đơn vị đo dung tích nhỏ.", "Chai nước có 500 ml nước.", "250 ml + 250 ml = 500 ml.", "500 ml"),
        "temperature": ("Đọc nhiệt độ trên nhiệt kế trong tình huống quen thuộc.", "Nhiệt kế chỉ 30 độ C.", "Nhiệt kế chỉ 25 độ C.", "25 độ C"),
        "mul_table": (f"Ghi nhớ và vận dụng bảng nhân {table}.", f"{table} x 4 = {table * 4} vì có 4 nhóm, mỗi nhóm {table}.", f"{table} x 6 = {table * 6}.", str(table * 6)),
        "div_table": (f"Ghi nhớ và vận dụng bảng chia {table}.", f"{table * 4} : {table} = 4 vì chia thành các nhóm {table}.", f"{table * 6} : {table} = 6.", "6"),
        "times_more": ("Hiểu gấp một số lên một số lần là nhân số đó với số lần.", "4 được gấp lên 3 lần: 4 x 3 = 12.", "6 gấp lên 4 lần bằng 24.", "24"),
        "times_less": ("Hiểu giảm một số đi một số lần là chia số đó cho số lần.", "18 giảm đi 3 lần: 18 : 3 = 6.", "24 giảm đi 4 lần bằng 6.", "6"),
        "fraction": ("Nhận biết một phần bằng nhau của một hình hoặc một nhóm đồ vật.", "Tô 1 trong 4 phần bằng nhau là một phần tư.", "Tô 1 trong 6 phần bằng nhau là một phần sáu.", "1/6"),
        "multiply_tens": ("Nhân số tròn chục với số có một chữ số bằng cách nhân phần chục.", "30 x 4 = 120 vì 3 chục x 4 = 12 chục.", "20 x 5 = 100.", "100"),
        "multiply": ("Nhân số có nhiều chữ số với số có một chữ số theo từng hàng.", "123 x 3 = 369.", "212 x 4 = 848.", "848"),
        "multiply_carry": ("Nhân với số có một chữ số có nhớ theo từng hàng.", "127 x 3 = 381: 7 x 3 = 21, viết 1 nhớ 2.", "248 x 3 = 744.", "744"),
        "multiply_large": ("Nhân số trong phạm vi 100 000 với số có một chữ số theo từng hàng.", "12 324 x 2 = 24 648.", "21 203 x 3 = 63 609.", "63 609"),
        "divide": ("Chia số có nhiều chữ số cho số có một chữ số theo từng bước.", "248 : 2 = 124.", "369 : 3 = 123.", "123"),
        "divide_large": ("Chia số trong phạm vi 100 000 cho số có một chữ số theo từng bước.", "24 864 : 2 = 12 432.", "36 963 : 3 = 12 321.", "12 321"),
        "compare_times": ("So sánh số lớn gấp mấy lần số bé bằng phép chia.", "12 gấp 4 được 3 lần vì 12 : 4 = 3.", "20 gấp 5 được 4 lần.", "4 lần"),
        "two_step_problem": ("Giải bài toán có hai bước tính bằng cách tách dữ kiện theo thứ tự.", "Có 12 bông hoa, mua thêm 8 bông rồi chia đều vào 4 lọ: (12 + 8) : 4 = 5.", "18 quả táo thêm 6 quả rồi chia 3 đĩa, mỗi đĩa 8 quả.", "8 quả"),
        "expression": ("Nhận biết biểu thức số và tính giá trị theo đúng thứ tự.", "Tính 18 + 6 x 2: làm nhân trước, 6 x 2 = 12, rồi 18 + 12 = 30.", "40 - 12 : 3 = 36.", "36"),
        "angle": ("Nhận biết góc vuông và góc không vuông.", "Góc vuông có hai cạnh tạo thành hình như góc của quyển vở.", "Góc được đánh dấu ô vuông là góc vuông.", "Góc vuông"),
        "triangle_quad": ("Nhận biết hình tam giác và hình tứ giác theo số cạnh.", "Hình tam giác có 3 cạnh, hình tứ giác có 4 cạnh.", "Hình có 3 cạnh là hình tam giác.", "Hình tam giác"),
        "perimeter": ("Tính chu vi bằng tổng độ dài các cạnh.", "Tam giác có cạnh 3 cm, 4 cm, 5 cm thì chu vi là 12 cm.", "Hình vuông cạnh 6 cm có chu vi 24 cm.", "24 cm"),
        "rectangle": ("Nhận biết hình chữ nhật có 4 góc vuông, hai cạnh dài bằng nhau và hai cạnh ngắn bằng nhau.", "Hình chữ nhật có chiều dài 6 cm, chiều rộng 3 cm.", "Hình chữ nhật có 4 góc vuông.", "4 góc vuông"),
        "square": ("Nhận biết hình vuông có 4 cạnh bằng nhau và 4 góc vuông.", "Hình vuông có cạnh 4 cm.", "Hình vuông cạnh 5 cm có 4 cạnh bằng nhau.", "4 cạnh bằng nhau"),
        "number_10000": ("Đọc, viết và phân tích số trong phạm vi 10 000.", "3 426 gồm 3 nghìn, 4 trăm, 2 chục, 6 đơn vị.", "5 037 gồm 5 nghìn, 0 trăm, 3 chục, 7 đơn vị.", "5 nghìn, 0 trăm, 3 chục, 7 đơn vị"),
        "number_100000": ("Đọc, viết và phân tích số trong phạm vi 100 000.", "42 615 gồm 4 chục nghìn, 2 nghìn, 6 trăm, 1 chục, 5 đơn vị.", "70 408 gồm 7 chục nghìn, 0 nghìn, 4 trăm, 0 chục, 8 đơn vị.", "70 408"),
        "roman": ("Làm quen với chữ số La Mã I, V, X.", "VI = 6 vì V là 5 và I là 1.", "IX = 9.", "9"),
        "compare_large": ("So sánh các số lớn theo từng hàng từ trái sang phải.", "42 615 > 39 850 vì 4 chục nghìn lớn hơn 3 chục nghìn.", "70 408 > 69 999.", ">"),
        "midpoint": ("Nhận biết điểm ở giữa và trung điểm của đoạn thẳng.", "M là trung điểm của AB nếu M ở giữa A, B và AM = MB.", "Nếu AM = MB thì M là trung điểm của AB.", "M là trung điểm"),
        "circle": ("Nhận biết tâm, bán kính và đường kính của hình tròn.", "Đường kính đi qua tâm và dài gấp đôi bán kính.", "Bán kính là đoạn nối tâm với một điểm trên đường tròn.", "Bán kính"),
        "rounding": ("Làm tròn số đến hàng được yêu cầu.", "3 426 làm tròn đến hàng trăm là 3 400.", "7 683 làm tròn đến hàng nghìn là 8 000.", "8 000"),
        "box_cube": ("Nhận biết khối hộp chữ nhật và khối lập phương.", "Hộp sữa giống khối hộp chữ nhật, khối rubik giống khối lập phương.", "Khối lập phương có các mặt là hình vuông.", "Khối lập phương"),
        "month_year": ("Nhận biết tháng, năm và số ngày trong năm.", "Một năm thường có 12 tháng và 365 ngày.", "Một năm có 12 tháng.", "12 tháng"),
        "add_large": ("Cộng các số trong phạm vi 100 000 theo từng hàng.", "42 615 + 13 204 = 55 819.", "24 305 + 12 480 = 36 785.", "36 785"),
        "sub_large": ("Trừ các số trong phạm vi 100 000 theo từng hàng.", "55 819 - 13 204 = 42 615.", "36 785 - 12 480 = 24 305.", "24 305"),
        "money": ("Nhận biết và tính với tiền Việt Nam trong tình huống mua bán đơn giản.", "10 000 đồng + 5 000 đồng = 15 000 đồng.", "20 000 đồng - 8 000 đồng = 12 000 đồng.", "12 000 đồng"),
        "unknown": ("Tìm thành phần chưa biết bằng phép tính ngược.", "x + 15 = 42 nên x = 42 - 15 = 27.", "x x 3 = 24 nên x = 8.", "8"),
        "area": ("Hiểu diện tích là phần mặt phẳng được phủ kín bởi các ô vuông đơn vị.", "Hình chữ nhật 4 cm x 3 cm có diện tích 12 cm2.", "Hình vuông cạnh 5 cm có diện tích 25 cm2.", "25 cm2"),
        "statistics": ("Thu thập, phân loại và ghi chép số liệu thống kê.", "Ghi mỗi lựa chọn vào đúng nhóm rồi đếm số lượng.", "Có 6 bạn chọn bóng đá.", "6 bạn"),
        "data_table": ("Đọc bảng số liệu thống kê đơn giản.", "Bảng cho biết số sách mỗi tổ đọc được.", "Tổ 2 đọc 18 quyển sách.", "18 quyển"),
        "probability": ("Nhận biết khả năng xảy ra của một sự kiện.", "Bốc thẻ màu đỏ từ hộp có thẻ đỏ và xanh là có thể.", "Bốc được thẻ vàng từ hộp không có thẻ vàng là không thể.", "Không thể"),
    }
    objective, model, quick, answer = specs[k]
    return {"kind": k, "objective": objective, "model": model, "quick": quick, "answer": answer}


def header(draw, lesson, card_title):
    base.rounded(draw, (34, 34, 1166, 641), "#ffffff", "#cbd5e1", 3, 30)
    draw.text((78, 65), "Toán lớp 3", font=base.font(27, True), fill="#1f4e79")
    draw.text((78, 103), lesson, font=base.font(29, True), fill="#0f172a")
    base.label_box(draw, (860, 62, 1122, 116), card_title, fill="#e0f2fe", outline="#0369a1", color="#075985", size=24)
    draw.line((78, 145, 1122, 145), fill="#e2e8f0", width=3)


def draw_grid_number(draw, x, y, number, labels):
    colors = ["#ddd6fe", "#bfdbfe", "#fde68a", "#bbf7d0", "#fecaca"]
    s = str(number).replace(" ", "")
    for i, digit in enumerate(s):
        bx = x + i * 95
        base.rounded(draw, (bx, y, bx + 70, y + 70), colors[i % len(colors)], "#334155", 2, 12)
        base.draw_center(draw, (bx + 35, y + 35), digit, base.font(34, True), "#0f172a")
        draw.text((bx - 5, y + 82), labels[i], font=base.font(18, True), fill="#334155")


def draw_mul_groups(draw, table, groups):
    for g in range(groups):
        x = 115 + g * 170
        base.rounded(draw, (x, 245, x + 135, 455), "#f8fafc", "#94a3b8", 3, 20)
        draw.text((x + 28, 265), f"Nhóm {g + 1}", font=base.font(21, True), fill="#0f172a")
        if table <= 6:
            for i in range(table):
                cx = x + 33 + (i % 3) * 36
                cy = 325 + (i // 3) * 38
                base.draw_apple(draw, cx, cy, 0.65)
        else:
            for i in range(6):
                cx = x + 33 + (i % 3) * 36
                cy = 320 + (i // 3) * 36
                base.draw_apple(draw, cx, cy, 0.55)
            base.label_box(draw, (x + 28, 400, x + 108, 438), f"{table} quả", "#dcfce7", "#15803d", "#166534", 16)
    draw.text((790, 335), f"{table} x {groups} = {table * groups}", font=base.font(48, True), fill="#15803d")


def draw_div_groups(draw, table, groups):
    total = table * groups
    base.rounded(draw, (115, 235, 430, 495), "#fefce8", "#a16207", 3, 20)
    draw.text((140, 255), f"{total} đồ vật ban đầu", font=base.font(24, True), fill="#0f172a")
    for i in range(min(total, 18)):
        cx = 155 + (i % 6) * 40
        cy = 325 + (i // 6) * 40
        base.draw_apple(draw, cx, cy, 0.55)
    if total > 18:
        base.label_box(draw, (260, 445, 390, 485), f"tất cả {total}", "#fff7ed", "#ea580c", "#9a3412", 18)
    for g in range(groups):
        x = 520 + (g % 3) * 180
        y = 250 + (g // 3) * 125
        base.rounded(draw, (x, y, x + 145, y + 105), "#eff6ff", "#1d4ed8", 2, 16)
        draw.text((x + 17, y + 10), f"Nhóm {g + 1}", font=base.font(18, True), fill="#0f172a")
        if table <= 5:
            for i in range(table):
                cx = x + 28 + (i % 3) * 35
                cy = y + 56 + (i // 3) * 30
                base.draw_apple(draw, cx, cy, 0.45)
        else:
            for i in range(4):
                cx = x + 35 + (i % 2) * 42
                cy = y + 54 + (i // 2) * 28
                base.draw_apple(draw, cx, cy, 0.42)
            draw.text((x + 92, y + 60), f"{table}", font=base.font(22, True), fill="#1d4ed8")
    draw.text((760, 520), f"{total} : {table} = {groups}", font=base.font(44, True), fill="#1d4ed8")


def draw_vertical(draw, x, y, lines, result, op):
    for i, line in enumerate(lines):
        prefix = op if i == len(lines) - 1 else " "
        draw.text((x, y + i * 55), f"{prefix} {line}", font=base.font(48, True), fill="#0f172a")
    draw.line((x, y + len(lines) * 55 + 8, x + 250, y + len(lines) * 55 + 8), fill="#0f172a", width=5)
    draw.text((x + 30, y + len(lines) * 55 + 22), str(result), font=base.font(48, True), fill="#1d4ed8")


def draw_scene_by_kind(draw, name, k, card_index):
    table = table_number(name)
    if k == "mm":
        base.rounded(draw, (170, 340, 930, 410), "#fef3c7", "#a16207", 4, 8)
        for i in range(31):
            x = 190 + i * 24
            h = 55 if i % 10 == 0 else 35
            draw.line((x, 340, x, 340 + h), fill="#78350f", width=2)
            if i % 10 == 0:
                base.draw_center(draw, (x, 430), str(i // 10), base.font(18, True), "#78350f")
        draw.text((410, 255), "1 cm = 10 mm", font=base.font(54, True), fill="#1d4ed8")
    elif k in {"gram", "ml", "temperature", "month_year"}:
        draw_measurement(draw, k)
    elif k == "mul_table":
        draw_mul_groups(draw, table, 4 if card_index != 3 else 6)
    elif k == "div_table":
        draw_div_groups(draw, table, 4 if card_index != 3 else 6)
    elif k in {"times_more", "times_less", "compare_times"}:
        draw_times(draw, k)
    elif k == "fraction":
        draw_fraction(draw, name)
    elif k in {"multiply_tens", "multiply", "multiply_carry", "multiply_large"}:
        draw.text((135, 250), "Nhân theo từng hàng", font=base.font(36, True), fill="#0f172a")
        if k == "multiply_tens":
            draw_vertical(draw, 470, 270, ["30", "4"], "120", "x")
            base.label_box(draw, (130, 405, 390, 470), "3 chục x 4 = 12 chục", "#fef3c7", "#ca8a04", "#92400e", 22)
        elif k == "multiply_carry":
            draw_vertical(draw, 470, 270, ["127", "3"], "381", "x")
            base.label_box(draw, (130, 405, 390, 470), "7 x 3 = 21, nhớ 2", "#fee2e2", "#dc2626", "#991b1b", 22)
        elif k == "multiply_large":
            draw_vertical(draw, 450, 270, ["12324", "2"], "24648", "x")
        else:
            draw_vertical(draw, 470, 270, ["123", "3"], "369", "x")
            base.draw_group_panel(draw, (125, 330, 390, 505), 9, "Minh họa tích", color="#f0fdf4")
    elif k in {"divide", "divide_large"}:
        draw.text((155, 270), "Chia từng bước từ trái sang phải", font=base.font(34, True), fill="#0f172a")
        if k == "divide_large":
            draw.text((500, 310), "24 864 : 2 = 12 432", font=base.font(44, True), fill="#1d4ed8")
            parts = ["2 chục nghìn : 2", "4 nghìn : 2", "8 trăm : 2"]
        else:
            draw.text((520, 310), "248 : 2 = 124", font=base.font(54, True), fill="#1d4ed8")
            parts = ["2 trăm : 2", "4 chục : 2", "8 đơn vị : 2"]
        for i, part in enumerate(parts):
            base.label_box(draw, (185 + i * 290, 430, 430 + i * 290, 490), part, "#eff6ff", "#1d4ed8", "#1e3a8a", 21)
    elif k == "two_step_problem":
        base.draw_group_panel(draw, (120, 245, 390, 455), 12, "Có 12 hoa")
        base.draw_group_panel(draw, (440, 245, 650, 455), 8, "Thêm 8 hoa", color="#f0fdf4")
        draw.text((720, 275), "B1: 12 + 8 = 20", font=base.font(34, True), fill="#1d4ed8")
        draw.text((720, 340), "B2: 20 : 4 = 5", font=base.font(34, True), fill="#15803d")
        draw.text((720, 405), "Mỗi lọ 5 hoa", font=base.font(38, True), fill="#dc2626")
    elif k == "expression":
        draw.text((230, 260), "18 + 6 x 2", font=base.font(70, True), fill="#0f172a")
        base.label_box(draw, (210, 390, 520, 455), "Làm nhân trước: 6 x 2 = 12", "#fef3c7", "#ca8a04", "#92400e", 23)
        base.label_box(draw, (580, 390, 910, 455), "Rồi cộng: 18 + 12 = 30", "#dcfce7", "#15803d", "#166534", 23)
    elif k in {"angle", "triangle_quad", "perimeter", "rectangle", "square", "midpoint", "circle", "box_cube", "area"}:
        draw_geometry(draw, k)
    elif k in {"number_10000", "number_100000", "roman", "compare_large", "rounding"}:
        draw_number_topic(draw, k)
    elif k in {"add_large", "sub_large"}:
        if k == "add_large":
            draw_grid_number(draw, 120, 250, "42615", ["chục nghìn", "nghìn", "trăm", "chục", "đơn vị"])
            draw_vertical(draw, 745, 260, ["42615", "13204"], "55819", "+")
        else:
            draw_grid_number(draw, 120, 250, "55819", ["chục nghìn", "nghìn", "trăm", "chục", "đơn vị"])
            draw_vertical(draw, 745, 260, ["55819", "13204"], "42615", "-")
    elif k == "money":
        draw_money(draw)
    elif k == "unknown":
        draw.text((260, 255), "x + 15 = 42", font=base.font(64, True), fill="#0f172a")
        draw.text((300, 365), "x = 42 - 15 = 27", font=base.font(48, True), fill="#1d4ed8")
        base.label_box(draw, (690, 325, 1010, 405), "Dùng phép tính ngược", "#fef3c7", "#ca8a04", "#92400e", 25)
    elif k in {"statistics", "data_table", "probability"}:
        draw_data_probability(draw, k)


def draw_measurement(draw, k):
    if k == "gram":
        base.rounded(draw, (260, 280, 470, 465), "#fef3c7", "#a16207", 4, 25)
        draw.text((315, 340), "200 g", font=base.font(42, True), fill="#92400e")
        base.rounded(draw, (620, 250, 835, 490), "#dbeafe", "#1d4ed8", 4, 25)
        draw.text((665, 345), "1 kg", font=base.font(46, True), fill="#1d4ed8")
        draw.text((875, 360), "g đo vật nhẹ", font=base.font(31, True), fill="#0f172a")
    elif k == "ml":
        base.rounded(draw, (275, 230, 450, 515), "#dbeafe", "#1d4ed8", 4, 22)
        draw.rectangle((300, 360, 425, 490), fill="#60a5fa")
        draw.text((305, 285), "500 ml", font=base.font(36, True), fill="#1d4ed8")
        draw.text((610, 345), "ml đo lượng nước nhỏ", font=base.font(36, True), fill="#0f172a")
    elif k == "temperature":
        base.rounded(draw, (360, 205, 450, 520), "#fee2e2", "#dc2626", 4, 35)
        draw.rectangle((397, 270, 413, 460), fill="#ef4444")
        draw.ellipse((370, 445, 440, 515), fill="#ef4444", outline="#991b1b", width=3)
        draw.text((560, 330), "30 độ C", font=base.font(56, True), fill="#dc2626")
    else:
        base.rounded(draw, (245, 215, 805, 520), "#ffffff", "#1d4ed8", 5, 20)
        draw.rectangle((245, 215, 805, 285), fill="#bfdbfe", outline="#1d4ed8", width=5)
        draw.text((380, 235), "MỘT NĂM", font=base.font(38, True), fill="#1e3a8a")
        for i in range(12):
            x = 285 + (i % 4) * 125
            y = 315 + (i // 4) * 55
            base.label_box(draw, (x, y, x + 90, y + 38), f"Tháng {i + 1}", "#f8fafc", "#94a3b8", "#0f172a", 16)
        draw.text((850, 345), "12 tháng", font=base.font(48, True), fill="#dc2626")


def draw_times(draw, k):
    base.draw_group_panel(draw, (150, 270, 390, 455), 4, "Số bé: 4")
    if k == "times_less":
        base.draw_group_panel(draw, (545, 245, 850, 495), 18, "18 chia thành 3 phần", color="#fefce8")
        draw.text((890, 350), "18 : 3 = 6", font=base.font(42, True), fill="#1d4ed8")
    elif k == "compare_times":
        base.draw_group_panel(draw, (530, 245, 855, 495), 12, "Số lớn: 12", color="#f0fdf4")
        draw.text((890, 350), "12 : 4 = 3 lần", font=base.font(39, True), fill="#1d4ed8")
    else:
        for i in range(3):
            base.draw_group_panel(draw, (510 + i * 160, 280, 640 + i * 160, 450), 4, f"Lần {i + 1}", color="#f0fdf4")
        draw.text((780, 500), "4 x 3 = 12", font=base.font(42, True), fill="#15803d")


def draw_fraction(draw, name):
    n = ntext(name)
    denom = 4
    if "ba" in n:
        denom = 3
    if "nam" in n:
        denom = 5
    if "sau" in n:
        denom = 6
    if "bay" in n:
        denom = 7
    if "tam" in n:
        denom = 8
    if "chin" in n:
        denom = 9
    cx, cy, r = 420, 365, 145
    start = -90
    for i in range(denom):
        end = start + 360 / denom
        fill = "#facc15" if i == 0 else "#e0f2fe"
        draw.pieslice((cx - r, cy - r, cx + r, cy + r), start, end, fill=fill, outline="#1d4ed8", width=3)
        start = end
    draw.text((670, 320), f"Tô 1 trong {denom} phần", font=base.font(42, True), fill="#0f172a")
    draw.text((740, 390), f"1/{denom}", font=base.font(62, True), fill="#dc2626")


def draw_geometry(draw, k):
    if k == "angle":
        draw.line((250, 450, 250, 260), fill="#1d4ed8", width=8)
        draw.line((250, 450, 470, 450), fill="#1d4ed8", width=8)
        draw.rectangle((250, 410, 290, 450), outline="#dc2626", width=5)
        draw.line((690, 450, 900, 305), fill="#15803d", width=8)
        draw.line((690, 450, 950, 450), fill="#15803d", width=8)
        draw.text((210, 500), "Góc vuông", font=base.font(32, True), fill="#1d4ed8")
        draw.text((705, 500), "Góc không vuông", font=base.font(32, True), fill="#15803d")
    elif k == "triangle_quad":
        draw.polygon([(190, 470), (360, 250), (530, 470)], fill="#bfdbfe", outline="#1d4ed8")
        draw.polygon([(690, 270), (930, 250), (1010, 455), (740, 500)], fill="#dcfce7", outline="#15803d")
        draw.text((245, 510), "3 cạnh", font=base.font(34, True), fill="#1d4ed8")
        draw.text((800, 510), "4 cạnh", font=base.font(34, True), fill="#15803d")
    elif k == "perimeter":
        pts = [(210, 460), (410, 260), (610, 460)]
        draw.polygon(pts, fill="#bfdbfe", outline="#1d4ed8")
        draw.text((280, 335), "3 cm", font=base.font(27, True), fill="#0f172a")
        draw.text((480, 340), "4 cm", font=base.font(27, True), fill="#0f172a")
        draw.text((360, 475), "5 cm", font=base.font(27, True), fill="#0f172a")
        draw.text((720, 350), "3 + 4 + 5 = 12 cm", font=base.font(38, True), fill="#1d4ed8")
    elif k == "rectangle":
        draw.rectangle((250, 270, 720, 480), fill="#bfdbfe", outline="#1d4ed8", width=5)
        draw.text((440, 230), "6 cm", font=base.font(30, True), fill="#0f172a")
        draw.text((735, 360), "3 cm", font=base.font(30, True), fill="#0f172a")
        draw.text((780, 335), "4 góc vuông", font=base.font(40, True), fill="#1d4ed8")
    elif k == "square":
        draw.rectangle((315, 245, 625, 555), fill="#dcfce7", outline="#15803d", width=5)
        draw.text((420, 205), "4 cm", font=base.font(30, True), fill="#0f172a")
        draw.text((720, 345), "4 cạnh bằng nhau", font=base.font(40, True), fill="#15803d")
    elif k == "midpoint":
        draw.line((210, 365, 920, 365), fill="#1d4ed8", width=7)
        for x, label in [(210, "A"), (565, "M"), (920, "B")]:
            draw.ellipse((x - 11, 354, x + 11, 376), fill="#dc2626")
            draw.text((x - 12, 390), label, font=base.font(32, True), fill="#0f172a")
        draw.text((410, 465), "AM = MB", font=base.font(42, True), fill="#1d4ed8")
    elif k == "circle":
        cx, cy, r = 500, 365, 155
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), outline="#1d4ed8", width=6)
        draw.ellipse((cx - 7, cy - 7, cx + 7, cy + 7), fill="#dc2626")
        draw.line((cx, cy, cx + r, cy), fill="#15803d", width=5)
        draw.line((cx - r, cy, cx + r, cy), fill="#dc2626", width=4)
        draw.text((705, 300), "Tâm O", font=base.font(34, True), fill="#dc2626")
        draw.text((705, 360), "Bán kính", font=base.font(34, True), fill="#15803d")
        draw.text((705, 420), "Đường kính", font=base.font(34, True), fill="#1d4ed8")
    elif k == "box_cube":
        draw.rectangle((220, 300, 460, 490), fill="#bfdbfe", outline="#1d4ed8", width=5)
        draw.polygon([(220, 300), (285, 245), (525, 245), (460, 300)], fill="#dbeafe", outline="#1d4ed8")
        draw.polygon([(460, 300), (525, 245), (525, 435), (460, 490)], fill="#93c5fd", outline="#1d4ed8")
        draw.rectangle((700, 280, 920, 500), fill="#fecaca", outline="#dc2626", width=5)
        draw.text((205, 525), "Khối hộp chữ nhật", font=base.font(28, True), fill="#1d4ed8")
        draw.text((720, 525), "Khối lập phương", font=base.font(28, True), fill="#dc2626")
    else:
        for r in range(3):
            for c in range(4):
                x = 270 + c * 55
                y = 260 + r * 55
                draw.rectangle((x, y, x + 50, y + 50), fill="#bbf7d0", outline="#15803d", width=2)
        draw.text((570, 320), "4 x 3 = 12 cm2", font=base.font(46, True), fill="#15803d")


def draw_number_topic(draw, k):
    if k == "roman":
        entries = [("I", "1"), ("V", "5"), ("X", "10"), ("VI", "6"), ("IX", "9")]
        for i, (roman, value) in enumerate(entries):
            x = 160 + i * 180
            base.label_box(draw, (x, 285, x + 130, 360), roman, "#fef3c7", "#ca8a04", "#92400e", 34)
            draw.text((x + 43, 390), value, font=base.font(32, True), fill="#0f172a")
    elif k == "compare_large":
        draw_grid_number(draw, 120, 275, "42615", ["chục nghìn", "nghìn", "trăm", "chục", "đơn vị"])
        draw.text((555, 330), ">", font=base.font(70, True), fill="#dc2626")
        draw_grid_number(draw, 650, 275, "39850", ["chục nghìn", "nghìn", "trăm", "chục", "đơn vị"])
    elif k == "rounding":
        draw.text((230, 275), "3 426", font=base.font(76, True), fill="#0f172a")
        draw.line((240, 390, 770, 390), fill="#1d4ed8", width=5)
        draw.polygon([(770, 390), (745, 375), (745, 405)], fill="#1d4ed8")
        draw.text((815, 355), "3 400", font=base.font(58, True), fill="#dc2626")
        draw.text((270, 455), "Làm tròn đến hàng trăm", font=base.font(31, True), fill="#0f172a")
    else:
        number = "3426" if k == "number_10000" else "42615"
        labels = ["nghìn", "trăm", "chục", "đơn vị"] if k == "number_10000" else ["chục nghìn", "nghìn", "trăm", "chục", "đơn vị"]
        draw_grid_number(draw, 180, 280, number, labels)
        draw.text((720, 315), "Phân tích theo hàng", font=base.font(40, True), fill="#1d4ed8")


def draw_money(draw):
    notes = [("10 000 đ", "#dcfce7"), ("5 000 đ", "#fef3c7"), ("2 000 đ", "#dbeafe")]
    for i, (text, fill) in enumerate(notes):
        base.rounded(draw, (170 + i * 270, 275, 395 + i * 270, 400), fill, "#334155", 3, 18)
        draw.text((205 + i * 270, 320), text, font=base.font(31, True), fill="#0f172a")
    draw.text((350, 480), "10 000 + 5 000 = 15 000 đồng", font=base.font(38, True), fill="#1d4ed8")


def draw_data_probability(draw, k):
    if k == "statistics":
        rows = [("Bóng đá", 6), ("Cầu lông", 4), ("Bơi", 3)]
        for i, (label, count) in enumerate(rows):
            y = 265 + i * 85
            draw.text((210, y), label, font=base.font(30, True), fill="#0f172a")
            for j in range(count):
                draw.line((500 + j * 28, y, 500 + j * 28, y + 45), fill="#1d4ed8", width=5)
            draw.text((760, y), f"{count} bạn", font=base.font(28, True), fill="#1d4ed8")
    elif k == "data_table":
        headers = ["Tổ", "Số sách"]
        data = [("Tổ 1", "15"), ("Tổ 2", "18"), ("Tổ 3", "12")]
        for c, h in enumerate(headers):
            base.label_box(draw, (360 + c * 210, 240, 550 + c * 210, 300), h, "#dbeafe", "#1d4ed8", "#1e3a8a", 25)
        for r, row in enumerate(data):
            for c, value in enumerate(row):
                base.label_box(draw, (360 + c * 210, 315 + r * 65, 550 + c * 210, 370 + r * 65), value, "#f8fafc", "#94a3b8", "#0f172a", 24)
    else:
        boxes = [("Chắc chắn", "#dcfce7"), ("Có thể", "#fef3c7"), ("Không thể", "#fee2e2")]
        for i, (label, fill) in enumerate(boxes):
            base.label_box(draw, (150 + i * 320, 275, 410 + i * 320, 430), label, fill, "#334155", "#0f172a", 30)
        draw.text((235, 485), "Phân loại sự kiện theo khả năng xảy ra", font=base.font(32, True), fill="#1d4ed8")


def build_cards(lesson_name):
    spec = lesson_spec(lesson_name)
    k = spec["kind"]
    return [
        {
            "type": "concept_scene",
            "title": "Quan sát",
            "kind": k,
            "display_text": f"Quan sát tình huống minh họa cho bài: {lesson_name}.",
            "caption": spec["objective"],
            "student_task": "Nêu dữ kiện nhìn thấy trong tranh và gọi tên kiến thức chính.",
            "interaction": "choose_or_say",
            "expected_answer": spec["objective"],
            "wrong_hint": "Cho học sinh chỉ trực tiếp vào số, nhóm đồ vật, nhãn hoặc hình trong tranh trước khi trả lời.",
        },
        {
            "type": "worked_model",
            "title": "Làm mẫu",
            "kind": k,
            "display_text": spec["model"],
            "caption": spec["model"],
            "student_task": "Đọc ví dụ mẫu và nhắc lại bước làm chính.",
            "interaction": "none",
            "expected_answer": spec["model"],
            "wrong_hint": "Tách bài thành từng bước nhỏ, sau đó đối chiếu từng bước với tranh minh họa.",
        },
        {
            "type": "quick_try",
            "title": "Thử nhanh",
            "kind": k,
            "display_text": spec["quick"],
            "caption": f"Đáp án cần đạt: {spec['answer']}",
            "student_task": "Tự làm một câu ngắn cùng dạng với ví dụ mẫu.",
            "interaction": "fill_blank",
            "expected_answer": spec["answer"],
            "wrong_hint": "Quay lại thẻ làm mẫu, thay số mới vào đúng vị trí rồi làm lại theo thứ tự.",
        },
    ]


def build_blueprint():
    chapters = []
    for chapter in CHAPTERS:
        lessons = []
        for name in chapter["lessons"]:
            spec = lesson_spec(name)
            lessons.append({"lesson": name, "kind": spec["kind"], "objective": spec["objective"], "cards": build_cards(name)})
        chapters.append({"title": chapter["title"], "lessons": lessons})
    return {
        "grade": 3,
        "source_textbook": "SGK Toán 3 Cánh Diều",
        "version_note": "Bản chi tiết: mỗi bài có 3 tranh minh họa riêng cho quan sát, làm mẫu và thử nhanh.",
        "cognitive_profile": {
            "reading_level": "Học sinh lớp 3 đọc hiểu tốt hơn lớp 1-2, nhưng vẫn cần tranh, mô hình và bước làm rõ ràng.",
            "learning_style": "Quan sát tình huống, nhận ra mô hình toán, đọc ví dụ mẫu, sau đó tự làm một câu cùng dạng.",
            "ai_policy": "Có thể dùng AI ở mức gợi ý từng bước khi học sinh làm sai; không đưa ngay đáp án.",
        },
        "chapters": chapters,
    }


def set_run_font(run, size=10.5, bold=False, italic=False, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_meta(doc, label, value):
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True)
    r = p.add_run(str(value))
    set_run_font(r)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(str(item))
        set_run_font(r)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    r = cell.paragraphs[0].add_run(str(text))
    set_run_font(r, size=9, bold=bold)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_overview(doc, data):
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(["STT", "Chủ đề", "Số bài", "Số thẻ", "Số ảnh"]):
        set_cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i], "D9EAF7")
    for i, chapter in enumerate(data["chapters"], 1):
        row = table.add_row().cells
        card_count = sum(len(lesson["cards"]) for lesson in chapter["lessons"])
        for j, value in enumerate([i, chapter["title"], len(chapter["lessons"]), card_count, card_count]):
            set_cell_text(row[j], value)
            row[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def render_card_image(chapter_index, lesson_index, card_index, lesson_name, card):
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1200, 675), "#f1f5f9")
    draw = ImageDraw.Draw(img)
    header(draw, lesson_name, card["title"])
    draw_scene_by_kind(draw, lesson_name, card["kind"], card_index)
    base.rounded(draw, (78, 560, 1122, 618), "#f8fafc", "#e2e8f0", 2, 18)
    base.draw_wrapped(draw, card["caption"], (105, 575), base.font(22, True), fill="#0f172a", width=84, spacing=4)
    out = IMAGE_DIR / f"grade3-c{chapter_index:02d}-l{lesson_index:02d}-card{card_index:02d}.png"
    img.save(out)
    return out


def build_docx(data):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    base.set_defaults(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("LÝ THUYẾT TOÁN LỚP 3")
    set_run_font(r, size=24, bold=True, color="1F4E79")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Bản chi tiết theo thẻ, có tranh tình huống, đáp án và gợi ý sai - Codex")
    set_run_font(r, size=13, italic=True)
    add_meta(doc, "Nguồn tham khảo", data["source_textbook"])
    add_meta(doc, "Ghi chú phiên bản", data["version_note"])
    doc.add_heading("Định hướng thiết kế lớp 3", level=1)
    add_bullets(
        doc,
        [
            data["cognitive_profile"]["reading_level"],
            data["cognitive_profile"]["learning_style"],
            data["cognitive_profile"]["ai_policy"],
            "Mỗi bài có 3 tranh: quan sát tình huống, làm mẫu theo bước và thử nhanh có đáp án.",
        ],
    )
    doc.add_heading("Tổng quan nội dung", level=1)
    add_overview(doc, data)
    for chapter_index, chapter in enumerate(data["chapters"], 1):
        doc.add_page_break()
        doc.add_heading(f"Chủ đề {chapter_index}: {chapter['title']}", level=1)
        for lesson_index, lesson in enumerate(chapter["lessons"], 1):
            doc.add_heading(f"{chapter_index}.{lesson_index}. {lesson['lesson']}", level=2)
            add_meta(doc, "Mục tiêu bài học", lesson["objective"])
            for card_index, card in enumerate(lesson["cards"], 1):
                image_path = render_card_image(chapter_index, lesson_index, card_index, lesson["lesson"], card)
                doc.add_heading(f"Thẻ {card_index}: {card['title']}", level=3)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path), width=Inches(6.45))
                add_bullets(
                    doc,
                    [
                        f"Nội dung hiển thị: {card['display_text']}",
                        f"Nhiệm vụ học sinh: {card['student_task']}",
                        f"Tương tác đề xuất: {card['interaction']}",
                        f"Đáp án/Kết quả mong đợi: {card['expected_answer']}",
                        f"Gợi ý khi học sinh sai: {card['wrong_hint']}",
                    ],
                )
            doc.add_paragraph()
    try:
        doc.save(OUTPUT_PATH)
        return OUTPUT_PATH
    except PermissionError:
        doc.save(FALLBACK_OUTPUT_PATH)
        return FALLBACK_OUTPUT_PATH


def build_contact_sheets():
    paths = sorted(IMAGE_DIR.glob("grade3-*.png"))
    font = base.font(16, True)
    outs = []
    for start in range(0, len(paths), 60):
        chunk = paths[start : start + 60]
        cols = 5
        rows = math.ceil(len(chunk) / cols)
        sheet = Image.new("RGB", (cols * 240, rows * 170), "#ffffff")
        draw = ImageDraw.Draw(sheet)
        for i, path in enumerate(chunk):
            img = Image.open(path).resize((240, 135))
            x = (i % cols) * 240
            y = (i // cols) * 170
            sheet.paste(img, (x, y))
            draw.text((x + 6, y + 140), path.stem.replace("grade3-", ""), font=font, fill="#0f172a")
        out = IMAGE_DIR / f"contact-sheet-{start + 1}-{start + len(chunk)}.png"
        sheet.save(out)
        outs.append(out)
    return outs


def main():
    data = build_blueprint()
    BLUEPRINT_PATH.parent.mkdir(parents=True, exist_ok=True)
    BLUEPRINT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    output = build_docx(data)
    sheets = build_contact_sheets()
    print(BLUEPRINT_PATH)
    print(output)
    for sheet in sheets:
        print(sheet)
    print(f"lessons={sum(len(c['lessons']) for c in data['chapters'])}")
    print(f"cards={sum(len(l['cards']) for c in data['chapters'] for l in c['lessons'])}")
    print(f"images={len(list(IMAGE_DIR.glob('grade3-*.png')))}")


if __name__ == "__main__":
    main()
