import argparse
import hashlib
import json
import re
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "doc"
DEFAULT_OUTPUT = OUTPUT_DIR / "tong_hop_cau_hoi_crawl_cong_khai.docx"
DEFAULT_IMAGE_DIR = OUTPUT_DIR / "images"
DEFAULT_JSON_OUTPUT = OUTPUT_DIR / "crawled_questions.json"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
}

INDEX_URLS = {
    1: [
        "https://vietjack.com/toan-1-canh-dieu/trac-nghiem-toan-lop-1-tap-1.jsp",
        "https://vietjack.com/toan-1-canh-dieu/trac-nghiem-toan-lop-1-tap-2.jsp",
    ],
    2: [
        "https://vietjack.com/toan-2-canh-dieu/trac-nghiem-toan-lop-2-tap-1.jsp",
        "https://vietjack.com/toan-2-canh-dieu/trac-nghiem-toan-lop-2-tap-2.jsp",
    ],
    3: [
        "https://vietjack.com/toan-3-cd/trac-nghiem-toan-lop-3-tap-1.jsp",
        "https://vietjack.com/toan-3-cd/trac-nghiem-toan-lop-3-tap-2.jsp",
    ],
    4: [
        "https://vietjack.com/toan-4-cd/trac-nghiem-toan-lop-4.jsp",
    ],
    5: [
        "https://vietjack.com/toan-5-cd/trac-nghiem-toan-lop-5-tap-1.jsp",
        "https://vietjack.com/toan-5-cd/trac-nghiem-toan-lop-5-tap-2.jsp",
        "https://vietjack.com/toan-5-cd/trac-nghiem-toan-lop-5-canh-dieu.jsp",
    ],
    6: [
        "https://vietjack.com/toan-6-canh-dieu/bai-tap-trac-nghiem-toan-lop-6.jsp",
        "https://vietjack.com/toan-6-canh-dieu/bai-tap-trac-nghiem-toan-lop-6-nam-2022.jsp",
    ],
    7: [
        "https://vietjack.com/toan-7-cd/trac-nghiem-toan-lop-7-canh-dieu.jsp",
    ],
}


def clean_text(value):
    if not value:
        return ""
    value = value.replace("\xa0", " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def normalize_url(url):
    parsed = urlparse(url)
    return parsed._replace(fragment="").geturl()


def fetch_soup(session, url, timeout=20):
    response = session.get(url, headers=HEADERS, timeout=timeout)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")


def is_question_link(base_url, href, title):
    if not href:
        return False
    full_url = normalize_url(urljoin(base_url, href))
    parsed = urlparse(full_url)
    base_path = urlparse(base_url).path.rsplit("/", 1)[0]
    if "vietjack.com" not in parsed.netloc:
        return False
    cross_series_link = (
        (base_path.endswith("/toan-1-canh-dieu") and parsed.path.startswith("/toan-1-ket-noi/"))
        or (base_path.endswith("/toan-5-cd") and parsed.path.startswith("/toan-5-kn/"))
    )
    if not parsed.path.startswith(base_path) and not cross_series_link:
        return False
    if full_url == normalize_url(base_url):
        return False
    if clean_text(title).lower() in {"trang trước", "trang sau", "trang truoc", "trang sau"}:
        return False
    file_name = Path(parsed.path).name.lower()
    if re.match(r"trac-nghiem-toan-lop-\d+(-tap-\d+)?(-canh-dieu)?\.jsp$", file_name):
        return False
    if re.match(r"bai-tap-trac-nghiem-toan-lop-\d+(-nam-\d+)?\.jsp$", file_name):
        return False
    lowered = f"{title} {parsed.path}".lower()
    if "trac-nghiem" not in lowered and "bai-tap-trac-nghiem" not in lowered:
        return False
    if any(skip in lowered for skip in ["online", "de-thi", "ly-thuyet", "sgk"]):
        return False
    return True


def discover_links(session, grade, index_url):
    soup = fetch_soup(session, index_url)
    content = soup.select_one(".content") or soup.select_one("main") or soup.body or soup
    current_chapter = "Chưa phân loại"
    links = []

    for element in content.find_all(["h2", "h3", "h4", "a"]):
        if element.name in {"h2", "h3", "h4"}:
            text = clean_text(element.get_text(" "))
            if re.search(r"(chương|chủ đề|học kì|tập\s+\d)", text, re.I):
                current_chapter = text
            continue

        title = clean_text(element.get_text(" "))
        href = element.get("href")
        if not is_question_link(index_url, href, title):
            continue

        url = normalize_url(urljoin(index_url, href))
        links.append(
            {
                "grade": grade,
                "chapter": current_chapter,
                "title": title or Path(urlparse(url).path).stem,
                "url": url,
                "index_url": index_url,
            }
        )

    return links


def extract_text_with_math(tag):
    for node in tag.find_all(["script", "style", "ins"]):
        node.decompose()
    return clean_text(tag.get_text(" "))


def extract_images(tag, base_url):
    images = []
    img_tags = [tag] if getattr(tag, "name", None) == "img" else tag.find_all("img")
    for img in img_tags:
        src = img.get("data-src") or img.get("data-original") or img.get("src")
        if not src:
            continue
        url = normalize_url(urljoin(base_url, src))
        lowered = url.lower()
        if any(skip in lowered for skip in ["loading-cg.gif", "logo", "google-badge", "ios-store", "apple_store", "google_play"]):
            continue
        alt = clean_text(img.get("alt") or "")
        images.append({
            "url": url,
            "alt": alt,
            "width": clean_text(img.get("width") or ""),
            "height": clean_text(img.get("height") or ""),
        })
    return images


def image_extension(url, content_type):
    path_ext = Path(urlparse(url).path).suffix.lower()
    if path_ext in {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}:
        return path_ext
    if "png" in content_type:
        return ".png"
    if "jpeg" in content_type or "jpg" in content_type:
        return ".jpg"
    if "gif" in content_type:
        return ".gif"
    if "webp" in content_type:
        return ".webp"
    return ".img"


def download_image(session, image_url, assets_dir, cache):
    if image_url in cache:
        return cache[image_url]

    assets_dir.mkdir(parents=True, exist_ok=True)
    response = session.get(image_url, headers=HEADERS, timeout=30)
    response.raise_for_status()
    content_type = response.headers.get("content-type", "").lower()
    if not content_type.startswith("image/"):
        raise ValueError(f"URL không phải ảnh: {content_type}")

    ext = image_extension(image_url, content_type)
    file_name = hashlib.sha1(image_url.encode("utf-8")).hexdigest() + ext
    file_path = assets_dir / file_name
    if not file_path.exists():
        file_path.write_bytes(response.content)
    cache[image_url] = file_path
    return file_path


def parse_answer(text):
    patterns = [
        r"Đáp án đúng là\s*:?\s*([A-D])",
        r"Đáp án\s*:?\s*([A-D])",
        r"Chọn\s+([A-D])\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return match.group(1).upper()
    return ""


def has_explanation_context(tag):
    if tag is None:
        return False
    for node in [tag, *tag.find_parents()[:3]]:
        classes = set(node.get("class") or [])
        if node.name == "section" or "toggle" in classes or "toggle-content" in classes:
            return True
    return False


def append_unique_images(target, images):
    existing = {item["url"] for item in target}
    for image in images:
        if image["url"] not in existing:
            target.append(image)
            existing.add(image["url"])


def parse_question_page(session, task, max_questions=None):
    soup = fetch_soup(session, task["url"])
    content = soup.select_one(".content") or soup.select_one("main") or soup.body or soup
    for node in content.find_all(["script", "style", "ins", "iframe"]):
        node.decompose()

    blocks = content.find_all(["p", "div", "section", "img"], recursive=True)
    questions = []
    current = None

    for block in blocks:
        classes = set(block.get("class") or [])
        if block.name == "img":
            if current:
                target = current["explanation_images"] if has_explanation_context(block) else current["images"]
                append_unique_images(target, extract_images(block, task["url"]))
            continue

        text = extract_text_with_math(block)
        if not text or text in {"Quảng cáo", "TRẮC NGHIỆM ONLINE", "Hiển thị đáp án"}:
            continue
        if len(text) > 2500:
            continue

        q_match = re.match(r"^Câu\s+(\d+)[\.:]?\s*(.*)$", text, re.I)
        if q_match:
            if current and current["choices"]:
                questions.append(current)
                if max_questions and len(questions) >= max_questions:
                    break
            current = {
                "number": int(q_match.group(1)),
                "text": clean_text(q_match.group(2)),
                "images": extract_images(block, task["url"]),
                "choices": [],
                "correct_answer": "",
                "explanation": "",
                "explanation_images": [],
            }
            continue

        if not current:
            continue

        if block.name == "section" or "toggle" in classes or "toggle-content" in classes:
            append_unique_images(current["explanation_images"], extract_images(block, task["url"]))
            answer = parse_answer(text)
            if answer:
                current["correct_answer"] = answer
            explanation = re.sub(r"Hiển thị đáp án", "", text, flags=re.I)
            explanation = re.sub(r"Đáp án đúng là\s*:?\s*[A-D]\.?", "", explanation, flags=re.I)
            explanation = re.sub(r"Đáp án\s*:?\s*[A-D]\.?", "", explanation, flags=re.I)
            current["explanation"] = clean_text(explanation)
            continue

        append_unique_images(current["images"], extract_images(block, task["url"]))

        choice_match = re.match(r"^([A-D])[\.\)]\s*(.+)$", text, re.I)
        if choice_match:
            key = choice_match.group(1).upper()
            value = clean_text(choice_match.group(2))
            if not any(choice["key"] == key for choice in current["choices"]):
                current["choices"].append({"key": key, "text": value})
            continue

        if current["choices"]:
            answer = parse_answer(text)
            if answer and not current["correct_answer"]:
                current["correct_answer"] = answer
                continue
        elif text and not text.startswith(("A.", "B.", "C.", "D.")):
            if text not in current["text"]:
                current["text"] = clean_text(f"{current['text']} {text}")

    if current and current["choices"] and (not max_questions or len(questions) < max_questions):
        questions.append(current)

    return questions


def add_metadata(document, stats, source_urls):
    document.add_heading("Tổng hợp câu hỏi crawl công khai", 0)
    document.add_paragraph(f"Thời điểm tạo: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    document.add_paragraph(
        "Tài liệu này chỉ dùng để rà soát trước khi insert vào database. "
        "Dữ liệu được lấy từ các trang công khai, không cần đăng nhập. "
        "Khi sử dụng chính thức cần kiểm tra lại chất lượng, bản quyền và độ phù hợp."
    )
    document.add_paragraph(
        f"Tổng số lớp: {len(stats['grades'])}; tổng số bài crawl được: {stats['lessons']}; "
        f"tổng số câu hỏi trích xuất: {stats['questions']}; "
        f"tổng số ảnh phát hiện: {stats['images']}."
    )
    document.add_heading("Nguồn đã dùng", level=1)
    for url in source_urls:
        document.add_paragraph(url, style="List Bullet")


def add_embedded_image(document, session, image, assets_dir, cache):
    try:
        image_path = download_image(session, image["url"], assets_dir, cache)
        width_px = int(image.get("width") or 0)
        width_in = min(5.7, max(1.0, width_px / 96)) if width_px else 4.8
        document.add_picture(str(image_path), width=Inches(width_in))
        caption = image.get("alt") or image["url"]
        document.add_paragraph(f"Nguồn ảnh: {caption} - {image['url']}")
        return True
    except Exception as error:
        document.add_paragraph(f"Không tải được ảnh: {image['url']} ({error})", style="List Bullet")
        return False


def add_questions(document, grouped, embed_images=True, assets_dir=DEFAULT_IMAGE_DIR):
    image_cache = {}
    embedded_count = 0
    failed_count = 0
    with requests.Session() as image_session:
        for grade in sorted(grouped):
            document.add_page_break()
            document.add_heading(f"Lớp {grade}", level=1)
            chapters = grouped[grade]
            for chapter, lessons in chapters.items():
                document.add_heading(chapter, level=2)
                for lesson in lessons:
                    document.add_heading(lesson["title"], level=3)
                    document.add_paragraph(f"Nguồn: {lesson['url']}")
                    document.add_paragraph(f"Số câu trích xuất: {len(lesson['questions'])}")
                    for question in lesson["questions"]:
                        document.add_paragraph(
                            f"Câu {question['number']}. {question['text']}",
                            style="List Number",
                        )
                        if question["images"]:
                            document.add_paragraph("Ảnh trong đề:")
                            for image in question["images"]:
                                if embed_images:
                                    if add_embedded_image(document, image_session, image, assets_dir, image_cache):
                                        embedded_count += 1
                                    else:
                                        failed_count += 1
                                else:
                                    label = image["alt"] or image["url"]
                                    document.add_paragraph(f"{label}: {image['url']}", style="List Bullet")
                        for choice in question["choices"]:
                            document.add_paragraph(f"{choice['key']}. {choice['text']}")
                        if question["correct_answer"]:
                            document.add_paragraph(f"Đáp án: {question['correct_answer']}")
                        if question.get("explanation_images"):
                            document.add_paragraph("Ảnh trong lời giải:")
                            for image in question["explanation_images"]:
                                if embed_images:
                                    if add_embedded_image(document, image_session, image, assets_dir, image_cache):
                                        embedded_count += 1
                                    else:
                                        failed_count += 1
                                else:
                                    label = image["alt"] or image["url"]
                                    document.add_paragraph(f"{label}: {image['url']}", style="List Bullet")
                        if question["explanation"]:
                            document.add_paragraph(f"Lời giải: {question['explanation']}")
    return {"embedded_images": embedded_count, "failed_images": failed_count}


def style_document(document):
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Arial"


def calculate_stats(grouped):
    return {
        "grades": set(grouped.keys()),
        "lessons": sum(len(lessons) for chapters in grouped.values() for lessons in chapters.values()),
        "questions": sum(
            len(lesson["questions"])
            for chapters in grouped.values()
            for lessons in chapters.values()
            for lesson in lessons
        ),
        "images": sum(
            len(question["images"]) + len(question.get("explanation_images", []))
            for chapters in grouped.values()
            for lessons in chapters.values()
            for lesson in lessons
            for question in lesson["questions"]
        ),
    }


def save_docx(grouped, source_urls, output_path, embed_images=True, assets_dir=DEFAULT_IMAGE_DIR):
    stats = calculate_stats(grouped)
    document = Document()
    style_document(document)
    add_metadata(document, stats, source_urls)
    image_stats = add_questions(document, grouped, embed_images=embed_images, assets_dir=assets_dir)
    stats.update(image_stats)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return stats


def make_json_export(grouped, source_urls, stats):
    lessons = []
    for grade in sorted(grouped):
        for chapter, chapter_lessons in grouped[grade].items():
            for lesson in chapter_lessons:
                lessons.append({
                    "grade": grade,
                    "chapter": chapter,
                    "title": lesson["title"],
                    "url": lesson["url"],
                    "index_url": lesson.get("index_url", ""),
                    "questions": lesson["questions"],
                })

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_urls": source_urls,
        "stats": {
            "grades": sorted(list(stats["grades"])),
            "lessons": stats["lessons"],
            "questions": stats["questions"],
            "images": stats["images"],
            "embedded_images": stats.get("embedded_images", 0),
            "failed_images": stats.get("failed_images", 0),
        },
        "lessons": lessons,
    }


def save_json_export(grouped, source_urls, stats, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    payload = make_json_export(grouped, source_urls, stats)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def parse_grades(value):
    if not value:
        return sorted(INDEX_URLS)
    grades = []
    for part in value.split(","):
        part = part.strip()
        if "-" in part:
            start, end = [int(item) for item in part.split("-", 1)]
            grades.extend(range(start, end + 1))
        elif part:
            grades.append(int(part))
    return [grade for grade in sorted(set(grades)) if grade in INDEX_URLS]


def main():
    parser = argparse.ArgumentParser(description="Crawl câu hỏi công khai và xuất file Word để duyệt.")
    parser.add_argument("--grades", default="1-5", help="Ví dụ: 1-5 hoặc 3,4,5")
    parser.add_argument("--max-lessons", type=int, default=0, help="Giới hạn số bài để test. 0 là không giới hạn.")
    parser.add_argument("--max-questions-per-lesson", type=int, default=0, help="Giới hạn câu mỗi bài. 0 là không giới hạn.")
    parser.add_argument("--delay", type=float, default=0.35, help="Độ trễ giữa các request.")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Đường dẫn file .docx đầu ra.")
    parser.add_argument("--json-output", default=str(DEFAULT_JSON_OUTPUT), help="Đường dẫn file JSON dùng để import.")
    parser.add_argument("--image-dir", default=str(DEFAULT_IMAGE_DIR), help="Thư mục cache ảnh đã tải.")
    parser.add_argument("--no-images", action="store_true", help="Chỉ ghi URL ảnh, không nhúng ảnh vào Word.")
    parser.add_argument("--skip-docx", action="store_true", help="Chỉ xuất JSON, không tạo file Word.")
    args = parser.parse_args()

    grades = parse_grades(args.grades)
    output_path = Path(args.output)
    grouped = defaultdict(lambda: defaultdict(list))
    source_urls = []
    seen_task_urls = set()

    with requests.Session() as session:
        tasks = []
        for grade in grades:
            for index_url in INDEX_URLS[grade]:
                print(f"Đang đọc mục lục lớp {grade}: {index_url}")
                source_urls.append(index_url)
                try:
                    for task in discover_links(session, grade, index_url):
                        if task["url"] not in seen_task_urls:
                            seen_task_urls.add(task["url"])
                            tasks.append(task)
                except Exception as error:
                    print(f"  Lỗi đọc mục lục: {error}")
                time.sleep(args.delay)

        if args.max_lessons:
            tasks = tasks[: args.max_lessons]

        print(f"Tìm thấy {len(tasks)} bài/trang câu hỏi cần crawl.")

        for index, task in enumerate(tasks, start=1):
            print(f"[{index}/{len(tasks)}] {task['title']}")
            try:
                questions = parse_question_page(
                    session,
                    task,
                    max_questions=args.max_questions_per_lesson or None,
                )
            except Exception as error:
                print(f"  Lỗi crawl trang bài: {error}")
                questions = []
            if questions:
                grouped[task["grade"]][task["chapter"]].append({**task, "questions": questions})
                print(f"  Trích xuất {len(questions)} câu.")
            else:
                print("  Không trích xuất được câu hỏi.")
            time.sleep(args.delay)

    if args.skip_docx:
        stats = calculate_stats(grouped)
        stats.update({"embedded_images": 0, "failed_images": 0})
    else:
        stats = save_docx(
            grouped,
            sorted(set(source_urls)),
            output_path,
            embed_images=not args.no_images,
            assets_dir=Path(args.image_dir),
        )
    json_output_path = Path(args.json_output)
    save_json_export(grouped, sorted(set(source_urls)), stats, json_output_path)
    print("Hoàn tất.")
    if not args.skip_docx:
        print(f"File Word: {output_path}")
    print(f"File JSON: {json_output_path}")
    print(f"Số bài: {stats['lessons']}")
    print(f"Số câu hỏi: {stats['questions']}")
    print(f"Số ảnh phát hiện: {stats['images']}")
    print(f"Số ảnh đã nhúng: {stats.get('embedded_images', 0)}")
    print(f"Số ảnh lỗi: {stats.get('failed_images', 0)}")


if __name__ == "__main__":
    main()
