from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MANIFEST = ROOT / "database" / "question_supplements" / "supplement_manifest.json"
DEFAULT_OUTPUT = ROOT / "documents" / "question_supplements"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "667085"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Xuất câu hỏi bổ sung thành DOCX, tối đa 5 bài mỗi tệp.")
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def set_run_font(run, size=11, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_style(style, size, color, before, after, line_spacing=1.25, bold=False):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def shade_paragraph(paragraph, fill):
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def configure_document(doc: Document, grade: int, chapter_name: str):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    configure_style(styles["Normal"], 11, "000000", 0, 6, 1.25)
    configure_style(styles["Title"], 24, DARK_BLUE, 0, 8, 1.0, True)
    configure_style(styles["Heading 1"], 16, BLUE, 18, 10, 1.0, True)
    configure_style(styles["Heading 2"], 13, BLUE, 14, 7, 1.0, True)
    configure_style(styles["Heading 3"], 12, DARK_BLUE, 10, 5, 1.0, True)
    if "Question Label" not in styles:
        style = styles.add_style("Question Label", WD_STYLE_TYPE.PARAGRAPH)
    configure_style(styles["Question Label"], 11.5, DARK_BLUE, 8, 4, 1.0, True)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    run = header_paragraph.add_run(f"BỔ SUNG CÂU HỎI LỚP {grade}  |  {chapter_name}")
    set_run_font(run, size=9, bold=True, color=MUTED)
    add_page_field(section.footer.paragraphs[0])


def add_title_block(doc, grade, chapter_name, lesson_names, group_index):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    run = kicker.add_run("NGÂN HÀNG CÂU HỎI BỔ SUNG")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.add_run(f"Bổ sung câu hỏi lớp {grade}")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(chapter_name)
    set_run_font(run, size=13, bold=True, color=DARK_BLUE)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    text = f"Nhóm {group_index:02d} • {len(lesson_names)} bài học • " + " | ".join(lesson_names)
    run = meta.add_run(text)
    set_run_font(run, size=9.5, color=MUTED)


def add_picture_with_alt(paragraph, image_path: Path, alt_text: str):
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(3.65))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt_text[:250])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(4)


def add_question(doc, question, number):
    label = doc.add_paragraph(style="Question Label")
    label.add_run(f"Câu {number:02d}  •  {question['difficulty']}")

    text = question.get("content", {}).get("text", "")
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=True, color="172554")

    images = question.get("content", {}).get("images", [])
    if images:
        image_url = images[0].get("url", "")
        image_path = ROOT / "public" / Path(*image_url.lstrip("/").split("/"))
        if image_path.exists():
            picture_paragraph = doc.add_paragraph()
            add_picture_with_alt(picture_paragraph, image_path, images[0].get("alt_text", text))

    choices = question.get("choices") or []
    for choice in choices:
        option = doc.add_paragraph()
        option.paragraph_format.left_indent = Inches(0.25)
        option.paragraph_format.first_line_indent = Inches(-0.18)
        option.paragraph_format.space_after = Pt(1)
        run = option.add_run(f"{choice.get('key', '')}. {choice.get('text', '')}")
        set_run_font(run, size=10)

    answer = doc.add_paragraph()
    answer.paragraph_format.space_before = Pt(3)
    answer.paragraph_format.space_after = Pt(5)
    shade_paragraph(answer, LIGHT_GRAY)
    run = answer.add_run(f"Đáp án: {question.get('correct_answer', '')}. ")
    set_run_font(run, size=9.5, bold=True, color=DARK_BLUE)
    explanation = question.get("explanation", {}).get("text", "")
    run = answer.add_run(f"Lời giải: {explanation}")
    set_run_font(run, size=9.5, color="344054")


def group_manifest(payload):
    grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    metadata = {}
    for question in payload.get("questions", []):
        grade = int(question["grade"])
        chapter_id = int(question["chapter_id"])
        lesson_id = int(question["lesson_id"])
        grouped[grade][chapter_id][lesson_id].append(question)
        metadata[(grade, chapter_id, lesson_id)] = {
            "chapter_name": question["chapter_name"],
            "chapter_order": int(question["chapter_order"]),
            "lesson_name": question["lesson_name"],
            "lesson_order": int(question["lesson_order"]),
        }
    return grouped, metadata


def safe_part(value):
    value = re.sub(r"[<>:\"/\\|?*]", "-", str(value))
    return re.sub(r"\s+", " ", value).strip().rstrip(".")


def build_documents(manifest_path: Path, output_dir: Path):
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    grouped, metadata = group_manifest(payload)
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = []

    for grade in sorted(grouped):
        chapter_ids = sorted(
            grouped[grade],
            key=lambda chapter_id: min(
                metadata[(grade, chapter_id, lesson_id)]["chapter_order"]
                for lesson_id in grouped[grade][chapter_id]
            ),
        )
        for chapter_id in chapter_ids:
            lesson_ids = sorted(
                grouped[grade][chapter_id],
                key=lambda lesson_id: metadata[(grade, chapter_id, lesson_id)]["lesson_order"],
            )
            chapter_name = metadata[(grade, chapter_id, lesson_ids[0])]["chapter_name"]
            for group_index, start in enumerate(range(0, len(lesson_ids), 5), start=1):
                chunk = lesson_ids[start : start + 5]
                lesson_names = [metadata[(grade, chapter_id, lesson_id)]["lesson_name"] for lesson_id in chunk]
                doc = Document()
                configure_document(doc, grade, chapter_name)
                add_title_block(doc, grade, chapter_name, lesson_names, group_index)

                question_number = 1
                for lesson_id in chunk:
                    lesson_name = metadata[(grade, chapter_id, lesson_id)]["lesson_name"]
                    heading = doc.add_paragraph(style="Heading 1")
                    heading.paragraph_format.keep_with_next = True
                    heading.add_run(lesson_name)
                    summary = doc.add_paragraph()
                    summary.paragraph_format.space_after = Pt(8)
                    run = summary.add_run(f"Số câu bổ sung trong tài liệu: {len(grouped[grade][chapter_id][lesson_id])}")
                    set_run_font(run, size=9.5, italic=True, color=MUTED)
                    for question in grouped[grade][chapter_id][lesson_id]:
                        add_question(doc, question, question_number)
                        question_number += 1

                first = metadata[(grade, chapter_id, chunk[0])]["lesson_order"]
                last = metadata[(grade, chapter_id, chunk[-1])]["lesson_order"]
                filename = safe_part(
                    f"Bổ sung câu hỏi lớp {grade} - Chương {chapter_id:02d} - Nhóm {group_index:02d} - Bài {first:02d}-{last:02d}.docx"
                )
                output = output_dir / filename
                doc.core_properties.title = f"Bổ sung câu hỏi lớp {grade}"
                doc.core_properties.subject = f"{chapter_name}; nhóm tối đa 5 bài"
                doc.core_properties.author = "Hệ thống ôn luyện Toán Tiểu học"
                doc.save(output)
                outputs.append({
                    "path": str(output.resolve()),
                    "grade": grade,
                    "chapter_id": chapter_id,
                    "chapter_name": chapter_name,
                    "lesson_ids": chunk,
                    "question_count": sum(len(grouped[grade][chapter_id][lesson_id]) for lesson_id in chunk),
                })

    output_manifest = output_dir / "manifest.json"
    output_manifest.write_text(
        json.dumps({"document_count": len(outputs), "documents": outputs}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return outputs, output_manifest


def main():
    args = parse_args()
    outputs, output_manifest = build_documents(args.manifest.resolve(), args.output_dir.resolve())
    print(json.dumps({"document_count": len(outputs), "manifest": str(output_manifest)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
