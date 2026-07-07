import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "doc"
SOURCE_JSON = OUTPUT_DIR / "grade-1-generated-questions.json"
IMAGE_DIR = OUTPUT_DIR / "grade1_question_images_ai_chapter1"
JSON_OUT = OUTPUT_DIR / "grade-1-generated-questions-ai-chuong-1.json"
DOCX_OUT = OUTPUT_DIR / "cau-hoi-lop-1-codex-ai-chuong-1.docx"
CONTACT_SHEET_OUT = IMAGE_DIR / "contact-sheet-chapter-1.png"

GEN_DIR = Path.home() / ".codex" / "generated_images" / "019f31ab-a671-73a2-9741-ca61b8bf2110"

SHEETS = [
    "ig_08d6ddc13a18c46c016a4c3dd9d71c8191bbecd5570f3cbdc0.png",
    "ig_08d6ddc13a18c46c016a4c3e16c4608191b8b6369667b77281.png",
    "ig_08d6ddc13a18c46c016a4c3e521d8081918fd8ce60efcb623a.png",
    "ig_08d6ddc13a18c46c016a4c3e8f13e88191b06d9f2b1fe19ef6.png",
    "ig_08d6ddc13a18c46c016a4c3eb716bc8191bdd8fb6c578317e8.png",
    "ig_08d6ddc13a18c46c016a4c3ef3913c8191b05dba63d0a03400.png",
    "ig_08d6ddc13a18c46c016a4c3f2cfcac8191b3d301d439c29de1.png",
    "ig_08d6ddc13a18c46c016a4c3f6e2e508191bbf433366ecfdc72.png",
    "ig_08d6ddc13a18c46c016a4c3faf12708191aead2fd7107c5cb3.png",
]


def crop_panels():
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    panel_paths = []

    for lesson_index, sheet_name in enumerate(SHEETS, start=1):
        sheet_path = GEN_DIR / sheet_name
        if not sheet_path.exists():
            raise FileNotFoundError(sheet_path)

        image = Image.open(sheet_path).convert("RGB")
        w, h = image.size
        boxes = [
            (0, 0, w // 2, h // 2),
            (w // 2, 0, w, h // 2),
            (0, h // 2, w // 2, h),
        ]

        for q_index, box in enumerate(boxes, start=1):
            crop = image.crop(box)
            out = IMAGE_DIR / f"q1_{lesson_index:02d}_{q_index}.png"
            crop.save(out)
            panel_paths.append(out)

    return panel_paths


def build_contact_sheet(image_paths):
    thumb_w, thumb_h = 220, 220
    label_h = 34
    cols = 3
    rows = 9
    gap = 18
    sheet = Image.new(
        "RGB",
        (
            cols * thumb_w + (cols + 1) * gap,
            rows * (thumb_h + label_h) + (rows + 1) * gap,
        ),
        "#f8fafc",
    )
    draw = ImageDraw.Draw(sheet)

    for index, path in enumerate(image_paths):
        row, col = divmod(index, cols)
        x = gap + col * (thumb_w + gap)
        y = gap + row * (thumb_h + label_h + gap)
        img = Image.open(path).convert("RGB")
        img.thumbnail((thumb_w, thumb_h))
        px = x + (thumb_w - img.width) // 2
        py = y + (thumb_h - img.height) // 2
        sheet.paste(img, (px, py))
        draw.text((x + 4, y + thumb_h + 8), f"Cau {index + 1:02d}", fill="#0f172a")

    sheet.save(CONTACT_SHEET_OUT)


def set_doc_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)


def add_answer_line(doc, q):
    answer_text = next(
        choice["text"] for choice in q["choices"] if choice["key"] == q["correct_answer"]
    )
    p = doc.add_paragraph()
    run = p.add_run(f"Đáp án: {q['correct_answer']}. {answer_text}")
    run.bold = True
    run.font.color.rgb = RGBColor(22, 101, 52)

    exp = doc.add_paragraph()
    exp.add_run("Lời giải: ").bold = True
    exp.add_run(q["explanation"]["text"])


def build_docx(questions, image_paths):
    doc = Document()
    set_doc_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_heading("Bộ câu hỏi Toán lớp 1 - Chương 1", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Bản dùng ảnh AI sinh động theo từng câu hỏi, dùng để duyệt chất lượng trước khi mở rộng sang toàn bộ lớp 1."
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    last_lesson = None
    for idx, (q, image_path) in enumerate(zip(questions, image_paths), start=1):
        if q["lesson"] != last_lesson:
            doc.add_heading(f"Bài {q['lesson_index']}. {q['lesson']}", level=2)
            last_lesson = q["lesson"]

        p = doc.add_paragraph()
        p.add_run(f"Câu {idx}. ").bold = True
        p.add_run(q["content"]["text"])

        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(image_path), width=Inches(4.8))

        choices = "    ".join(f"{c['key']}. {c['text']}" for c in q["choices"])
        doc.add_paragraph(choices)
        add_answer_line(doc, q)

    doc.save(DOCX_OUT)


def main():
    panels = crop_panels()
    build_contact_sheet(panels)

    data = json.loads(SOURCE_JSON.read_text(encoding="utf-8"))
    questions = [q for q in data["questions"] if q["chapter_index"] == 1]
    if len(questions) != len(panels):
        raise RuntimeError(f"Expected {len(panels)} questions, got {len(questions)}")

    for q, image_path in zip(questions, panels):
        rel = image_path.relative_to(ROOT).as_posix()
        q["content"]["images"] = [
            {
                "id": "question_image_ai_1",
                "url": rel,
                "alt": "Hình minh họa AI cho câu hỏi",
                "width_percent": 90,
            }
        ]
        q["image_path"] = str(image_path)

    output = {
        "grade": 1,
        "source": "generated_by_codex_with_ai_images",
        "scope": "chapter_1_preview",
        "question_count": len(questions),
        "questions": questions,
    }
    JSON_OUT.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    build_docx(questions, panels)

    print(DOCX_OUT)
    print(JSON_OUT)
    print(CONTACT_SHEET_OUT)


if __name__ == "__main__":
    main()
