import json
import math
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_grade2_theory_docx_detailed as base  # noqa: E402
import build_grade4_theory_docx_detailed as g4  # noqa: E402


OUTPUT_DIR = ROOT / "output" / "doc"
IMAGE_DIR = OUTPUT_DIR / "grade5_theory_images_detailed"
BLUEPRINT_PATH = ROOT / "content-theory" / "grade-5-theory-blueprint-detailed.json"
OUTPUT_PATH = OUTPUT_DIR / "ly-thuyet-lop-5-codex-chi-tiet.docx"
FALLBACK_OUTPUT_PATH = OUTPUT_DIR / "ly-thuyet-lop-5-codex-chi-tiet-fixed.docx"


CHAPTERS = [
    {
        "title": "Ôn tập và bổ sung về số tự nhiên, phân số, số thập phân",
        "lessons": [
            "Giới thiệu về tỉ số",
            "Tìm hai số khi biết tổng và tỉ số của hai số đó",
            "Tìm hai số khi biết hiệu và tỉ số của hai số đó",
            "Bài toán liên quan đến quan hệ phụ thuộc",
            "Hỗn số",
            "Phân số thập phân",
            "Số thập phân",
            "Số thập phân (tiếp theo)",
            "Số thập phân (tiếp theo 2)",
            "Số thập phân (tiếp theo 3)",
            "Số thập phân bằng nhau",
            "So sánh các số thập phân",
            "Làm tròn số thập phân",
            "Héc-ta",
            "Ki-lô-mét vuông",
        ],
    },
    {
        "title": "Các phép tính với số thập phân",
        "lessons": [
            "Cộng các số thập phân",
            "Trừ các số thập phân",
            "Nhân một số thập phân với 10, 100, 1 000",
            "Nhân một số thập phân với một số tự nhiên",
            "Nhân một số thập phân với một số thập phân",
            "Chia một số thập phân cho 10, 100, 1 000",
            "Chia một số thập phân cho một số tự nhiên",
            "Chia một số thập phân cho một số thập phân",
            "Viết các số đo đại lượng dưới dạng số thập phân",
            "Tỉ số phần trăm",
            "Tìm tỉ số phần trăm của hai số",
            "Tìm giá trị phần trăm của một số cho trước",
            "Sử dụng máy tính cầm tay",
            "Tỉ lệ bản đồ",
        ],
    },
    {
        "title": "Hình học và đo lường",
        "lessons": [
            "Hình tam giác",
            "Diện tích hình tam giác",
            "Hình thang",
            "Diện tích hình thang",
            "Hình tròn. Đường tròn",
            "Chu vi hình tròn",
            "Diện tích hình tròn",
            "Hình hộp chữ nhật. Hình lập phương. Hình trụ",
            "Hình khai triển của hình hộp chữ nhật, hình lập phương, hình trụ",
            "Diện tích xung quanh, diện tích toàn phần của hình hộp chữ nhật và hình lập phương",
            "Thể tích của một hình",
            "Xăng-ti-mét khối. Đề-xi-mét khối",
            "Mét khối",
            "Thể tích hình hộp chữ nhật, hình lập phương",
            "Cộng số đo thời gian. Trừ số đo thời gian",
            "Nhân số đo thời gian với một số. Chia số đo thời gian cho một số",
            "Vận tốc",
            "Quãng đường, thời gian trong chuyển động đều",
        ],
    },
    {
        "title": "Thống kê và xác suất. Ôn tập cuối năm",
        "lessons": [
            "Biểu đồ hình quạt",
            "Một số cách biểu diễn số liệu thống kê",
            "Mô tả số lần lặp lại của một kết quả có thể xảy ra trong một số trò chơi đơn giản",
        ],
    },
]


def ntext(text):
    return base.norm(text)


def lesson_kind(name):
    n = ntext(name)
    if "tong va ti so" in n:
        return "sum_ratio"
    if "hieu va ti so" in n:
        return "diff_ratio"
    if "ti so phan tram cua hai so" in n:
        return "percent_ratio"
    if "gia tri phan tram" in n:
        return "percent_value"
    if n == "ti so phan tram":
        return "percent_intro"
    if "ti so" in n and "ban do" not in n:
        return "ratio"
    if "quan he phu thuoc" in n:
        return "dependency"
    if "hon so" in n:
        return "mixed_number"
    if "phan so thap phan" in n:
        return "decimal_fraction"
    if "cong cac so thap phan" in n:
        return "decimal_add"
    if "tru cac so thap phan" in n:
        return "decimal_sub"
    if "nhan mot so thap phan voi 10" in n:
        return "decimal_mul10"
    if "nhan mot so thap phan voi mot so tu nhien" in n:
        return "decimal_mul_nat"
    if "nhan mot so thap phan voi mot so thap phan" in n:
        return "decimal_mul_dec"
    if "chia mot so thap phan cho 10" in n:
        return "decimal_div10"
    if "chia mot so thap phan cho mot so tu nhien" in n:
        return "decimal_div_nat"
    if "chia mot so thap phan cho mot so thap phan" in n:
        return "decimal_div_dec"
    if "so do dai luong" in n:
        return "decimal_measure"
    if "so thap phan bang nhau" in n:
        return "decimal_equal"
    if "so sanh" in n and "thap phan" in n:
        return "decimal_compare"
    if "lam tron" in n:
        return "decimal_round"
    if "so thap phan" in n:
        return "decimal_number"
    if "hec-ta" in n:
        return "hectare"
    if "ki-lo-met vuong" in n:
        return "km2"
    if "may tinh" in n:
        return "calculator"
    if "ti le ban do" in n:
        return "map_scale"
    if "dien tich hinh tam giac" in n:
        return "triangle_area"
    if "hinh tam giac" in n:
        return "triangle"
    if "dien tich hinh thang" in n:
        return "trapezoid_area"
    if "hinh thang" in n:
        return "trapezoid"
    if "chu vi hinh tron" in n:
        return "circle_perimeter"
    if "dien tich hinh tron" in n:
        return "circle_area"
    if "hinh tron" in n:
        return "circle"
    if "khai trien" in n:
        return "net_3d"
    if "dien tich xung quanh" in n:
        return "surface_area"
    if "the tich hinh hop" in n:
        return "box_volume"
    if "the tich cua mot hinh" in n:
        return "volume_intro"
    if "xang-ti-met khoi" in n:
        return "cm3_dm3"
    if "met khoi" in n:
        return "m3"
    if "hinh hop" in n or "hinh lap phuong" in n or "hinh tru" in n:
        return "solid_shapes"
    if "cong so do thoi gian" in n:
        return "time_add_sub"
    if "nhan so do thoi gian" in n:
        return "time_mul_div"
    if "van toc" in n:
        return "speed"
    if "quang duong" in n or "chuyen dong deu" in n:
        return "distance_time"
    if "bieu do hinh quat" in n:
        return "pie_chart"
    if "bieu dien so lieu" in n:
        return "data_representation"
    if "lap lai" in n or "tro choi" in n:
        return "frequency_event"
    return "decimal_number"


def lesson_spec(name):
    k = lesson_kind(name)
    specs = {
        "ratio": ("Nhận biết tỉ số của hai đại lượng cùng loại.", "Tỉ số của 3 và 5 là 3 : 5 hay 3/5.", "Tỉ số của 4 và 7 là 4/7.", "4/7"),
        "sum_ratio": ("Tìm hai số khi biết tổng và tỉ số bằng sơ đồ đoạn thẳng.", "Tổng 48, tỉ số 1 : 3. Tổng số phần là 4, mỗi phần 12; hai số là 12 và 36.", "Tổng 60, tỉ số 2 : 3 thì hai số là 24 và 36.", "24 và 36"),
        "diff_ratio": ("Tìm hai số khi biết hiệu và tỉ số bằng sơ đồ đoạn thẳng.", "Hiệu 24, tỉ số 1 : 3. Hiệu số phần là 2, mỗi phần 12; hai số là 12 và 36.", "Hiệu 18, tỉ số 2 : 5 thì hai số là 12 và 30.", "12 và 30"),
        "dependency": ("Nhận biết quan hệ phụ thuộc giữa hai đại lượng.", "1 hộp có 6 bút, 4 hộp có 24 bút.", "1 quyển 8 000 đồng, 5 quyển hết 40 000 đồng.", "40 000 đồng"),
        "mixed_number": ("Nhận biết hỗn số gồm phần nguyên và phần phân số.", "2 1/3 gồm phần nguyên 2 và phần phân số 1/3.", "3 2/5 có phần nguyên là 3.", "3"),
        "decimal_fraction": ("Nhận biết phân số thập phân có mẫu số là 10, 100, 1000.", "7/10 là phân số thập phân.", "35/100 là phân số thập phân.", "35/100"),
        "decimal_number": ("Nhận biết số thập phân và giá trị các hàng sau dấu phẩy.", "3,47 gồm 3 đơn vị, 4 phần mười, 7 phần trăm.", "5,08 gồm 5 đơn vị, 0 phần mười, 8 phần trăm.", "5 đơn vị, 0 phần mười, 8 phần trăm"),
        "decimal_equal": ("Nhận biết các số thập phân bằng nhau khi thêm hoặc bỏ chữ số 0 ở tận cùng phần thập phân.", "3,5 = 3,50.", "7,20 = 7,2.", "7,2"),
        "decimal_compare": ("So sánh số thập phân theo phần nguyên rồi từng hàng thập phân.", "4,35 > 4,28 vì 3 phần mười lớn hơn 2 phần mười.", "6,08 < 6,1.", "<"),
        "decimal_round": ("Làm tròn số thập phân đến hàng yêu cầu.", "3,47 làm tròn đến hàng phần mười là 3,5.", "8,263 làm tròn đến hàng phần trăm là 8,26.", "8,26"),
        "hectare": ("Nhận biết héc-ta là đơn vị đo diện tích lớn.", "1 ha = 10 000 m2.", "3 ha = 30 000 m2.", "30 000 m2"),
        "km2": ("Nhận biết ki-lô-mét vuông là đơn vị đo diện tích rất lớn.", "1 km2 = 1 000 000 m2.", "2 km2 = 2 000 000 m2.", "2 000 000 m2"),
        "decimal_add": ("Cộng số thập phân bằng cách đặt thẳng hàng dấu phẩy.", "12,35 + 4,6 = 16,95.", "7,25 + 3,8 = 11,05.", "11,05"),
        "decimal_sub": ("Trừ số thập phân bằng cách đặt thẳng hàng dấu phẩy.", "12,35 - 4,6 = 7,75.", "9,5 - 2,36 = 7,14.", "7,14"),
        "decimal_mul10": ("Nhân số thập phân với 10, 100, 1000 bằng cách chuyển dấu phẩy sang phải.", "3,47 x 100 = 347.", "5,08 x 10 = 50,8.", "50,8"),
        "decimal_mul_nat": ("Nhân số thập phân với số tự nhiên như nhân số tự nhiên rồi đặt dấu phẩy.", "2,35 x 4 = 9,4.", "3,6 x 5 = 18.", "18"),
        "decimal_mul_dec": ("Nhân hai số thập phân rồi đếm tổng số chữ số ở phần thập phân.", "1,2 x 0,3 = 0,36.", "2,5 x 0,4 = 1,0.", "1,0"),
        "decimal_div10": ("Chia số thập phân cho 10, 100, 1000 bằng cách chuyển dấu phẩy sang trái.", "34,7 : 10 = 3,47.", "508 : 100 = 5,08.", "5,08"),
        "decimal_div_nat": ("Chia số thập phân cho số tự nhiên theo từng bước và giữ dấu phẩy ở thương.", "9,6 : 4 = 2,4.", "12,5 : 5 = 2,5.", "2,5"),
        "decimal_div_dec": ("Chia số thập phân cho số thập phân bằng cách chuyển về chia cho số tự nhiên.", "4,8 : 1,2 = 48 : 12 = 4.", "7,5 : 2,5 = 3.", "3"),
        "decimal_measure": ("Viết số đo đại lượng dưới dạng số thập phân.", "2 m 35 cm = 2,35 m.", "3 kg 250 g = 3,25 kg.", "3,25 kg"),
        "percent_intro": ("Nhận biết tỉ số phần trăm là tỉ số có mẫu số 100.", "25% = 25/100.", "40% = 40/100.", "40/100"),
        "percent_ratio": ("Tìm tỉ số phần trăm của hai số bằng cách lấy số thứ nhất chia số thứ hai rồi viết dạng phần trăm.", "15 : 60 = 0,25 = 25%.", "18 : 50 = 36%.", "36%"),
        "percent_value": ("Tìm giá trị phần trăm của một số cho trước.", "20% của 150 là 150 x 20 : 100 = 30.", "15% của 200 là 30.", "30"),
        "calculator": ("Biết dùng máy tính cầm tay để kiểm tra phép tính.", "Nhập 12,5 + 7,8 được 20,3.", "Kiểm tra 3,6 x 5 được 18.", "18"),
        "map_scale": ("Nhận biết tỉ lệ bản đồ và tính độ dài thực tế.", "Tỉ lệ 1 : 1000, đoạn 5 cm trên bản đồ ứng với 5000 cm ngoài thực tế.", "Tỉ lệ 1 : 100, đoạn 3 cm ứng với 300 cm.", "300 cm"),
        "triangle": ("Nhận biết đáy và chiều cao của hình tam giác.", "Tam giác có đáy 6 cm và chiều cao 4 cm.", "Đáy là cạnh được chọn để tính diện tích.", "Đáy"),
        "triangle_area": ("Tính diện tích tam giác bằng đáy nhân chiều cao chia 2.", "S = 6 x 4 : 2 = 12 cm2.", "Tam giác đáy 8 cm, cao 5 cm có diện tích 20 cm2.", "20 cm2"),
        "trapezoid": ("Nhận biết hình thang có một cặp cạnh đối song song.", "Hai đáy của hình thang là hai cạnh song song.", "Hình thang có đáy lớn và đáy bé.", "Hai đáy"),
        "trapezoid_area": ("Tính diện tích hình thang bằng tổng hai đáy nhân chiều cao chia 2.", "S = (8 + 4) x 5 : 2 = 30 cm2.", "Đáy 10 cm, 6 cm, cao 4 cm thì diện tích 32 cm2.", "32 cm2"),
        "circle": ("Nhận biết tâm, bán kính, đường kính của hình tròn.", "Đường kính dài gấp đôi bán kính.", "Bán kính 4 cm thì đường kính 8 cm.", "8 cm"),
        "circle_perimeter": ("Tính chu vi hình tròn bằng đường kính nhân 3,14.", "C = 10 x 3,14 = 31,4 cm.", "Đường kính 8 cm thì chu vi 25,12 cm.", "25,12 cm"),
        "circle_area": ("Tính diện tích hình tròn bằng bán kính nhân bán kính nhân 3,14.", "S = 5 x 5 x 3,14 = 78,5 cm2.", "Bán kính 4 cm thì diện tích 50,24 cm2.", "50,24 cm2"),
        "solid_shapes": ("Nhận biết hình hộp chữ nhật, hình lập phương và hình trụ.", "Hộp sữa giống hình hộp chữ nhật, xúc xắc giống hình lập phương, lon nước giống hình trụ.", "Lon nước có dạng hình trụ.", "Hình trụ"),
        "net_3d": ("Nhận biết hình khai triển của khối hộp, khối lập phương và hình trụ.", "Hình lập phương có hình khai triển gồm 6 hình vuông.", "Hình hộp chữ nhật có 6 mặt.", "6 mặt"),
        "surface_area": ("Tính diện tích xung quanh và toàn phần của hình hộp chữ nhật, hình lập phương.", "Hộp chữ nhật dài 6, rộng 4, cao 3 có diện tích xung quanh (6+4)x2x3 = 60 cm2.", "Hình lập phương cạnh 5 có diện tích toàn phần 150 cm2.", "150 cm2"),
        "volume_intro": ("Hiểu thể tích là phần không gian vật chiếm chỗ.", "Khối được ghép từ 12 hình lập phương đơn vị có thể tích 12 đơn vị khối.", "Có 8 khối lập phương đơn vị thì thể tích là 8 đơn vị khối.", "8 đơn vị khối"),
        "cm3_dm3": ("Nhận biết xăng-ti-mét khối và đề-xi-mét khối.", "1 dm3 = 1000 cm3.", "2 dm3 = 2000 cm3.", "2000 cm3"),
        "m3": ("Nhận biết mét khối là đơn vị đo thể tích lớn.", "1 m3 = 1000 dm3.", "3 m3 = 3000 dm3.", "3000 dm3"),
        "box_volume": ("Tính thể tích hình hộp chữ nhật bằng dài nhân rộng nhân cao.", "V = 6 x 4 x 3 = 72 cm3.", "Hình lập phương cạnh 5 cm có thể tích 125 cm3.", "125 cm3"),
        "time_add_sub": ("Cộng, trừ số đo thời gian theo từng đơn vị.", "2 giờ 35 phút + 1 giờ 40 phút = 4 giờ 15 phút.", "5 giờ 20 phút - 2 giờ 45 phút = 2 giờ 35 phút.", "2 giờ 35 phút"),
        "time_mul_div": ("Nhân, chia số đo thời gian với một số.", "1 giờ 20 phút x 3 = 4 giờ.", "6 giờ 30 phút : 3 = 2 giờ 10 phút.", "2 giờ 10 phút"),
        "speed": ("Tính vận tốc bằng quãng đường chia thời gian.", "Đi 120 km trong 3 giờ thì vận tốc là 40 km/giờ.", "Đi 90 km trong 2 giờ thì vận tốc 45 km/giờ.", "45 km/giờ"),
        "distance_time": ("Tính quãng đường hoặc thời gian trong chuyển động đều.", "Quãng đường = vận tốc x thời gian: 40 x 3 = 120 km.", "Vận tốc 50 km/giờ đi 2 giờ được 100 km.", "100 km"),
        "pie_chart": ("Đọc biểu đồ hình quạt để nhận biết tỉ lệ từng phần.", "Phần Táo chiếm 50% hình quạt.", "Cam chiếm 25% nghĩa là một phần tư.", "25%"),
        "data_representation": ("Nhận biết nhiều cách biểu diễn số liệu: bảng, biểu đồ cột, biểu đồ hình quạt.", "Bảng số liệu và biểu đồ đều biểu diễn cùng một tập dữ liệu.", "Có thể biểu diễn số học sinh bằng bảng hoặc biểu đồ.", "Bảng hoặc biểu đồ"),
        "frequency_event": ("Mô tả số lần lặp lại của kết quả trong trò chơi đơn giản.", "Tung đồng xu 10 lần, mặt ngửa xuất hiện 6 lần.", "Mặt sấp xuất hiện 4 lần.", "4 lần"),
    }
    objective, model, quick, answer = specs[k]
    return {"kind": k, "objective": objective, "model": model, "quick": quick, "answer": answer}


def header(draw, lesson, card_title):
    base.rounded(draw, (34, 34, 1166, 641), "#ffffff", "#cbd5e1", 3, 30)
    draw.text((78, 65), "Toán lớp 5", font=base.font(27, True), fill="#1f4e79")
    draw.text((78, 103), lesson, font=base.font(28, True), fill="#0f172a")
    base.label_box(draw, (860, 62, 1122, 116), card_title, fill="#e0f2fe", outline="#0369a1", color="#075985", size=24)
    draw.line((78, 145, 1122, 145), fill="#e2e8f0", width=3)


def draw_quick_scene(draw, card):
    base.rounded(draw, (130, 225, 1070, 485), "#f8fafc", "#cbd5e1", 3, 24)
    draw.text((185, 260), "Bài thử nhanh", font=base.font(34, True), fill="#0f172a")
    base.draw_wrapped(draw, card["display_text"], (185, 320), base.font(35, True), fill="#1d4ed8", width=36, spacing=10)
    base.label_box(draw, (650, 370, 1010, 445), f"Đáp án: {card['expected_answer']}", "#dcfce7", "#15803d", "#166534", 24)


def draw_decimal_place(draw, text="3,47"):
    digits = [("3", "đơn vị"), (",", ""), ("4", "phần mười"), ("7", "phần trăm")]
    x = 310
    for digit, label in digits:
        if digit == ",":
            draw.text((x, 290), ",", font=base.font(70, True), fill="#0f172a")
            x += 48
            continue
        base.label_box(draw, (x, 270, x + 90, 350), digit, "#dbeafe", "#1d4ed8", "#1e3a8a", 38)
        draw.text((x - 10, 370), label, font=base.font(20, True), fill="#334155")
        x += 118
    draw.text((700, 315), text, font=base.font(56, True), fill="#dc2626")


def draw_fraction_grid(draw, num=3, den=4, label="3/4"):
    x0, y0 = 250, 250
    cols = den
    for i in range(den):
        fill = "#facc15" if i < num else "#e0f2fe"
        draw.rectangle((x0 + i * 85, y0, x0 + (i + 1) * 85, y0 + 150), fill=fill, outline="#1d4ed8", width=3)
    draw.text((650, 305), label, font=base.font(62, True), fill="#dc2626")


def draw_vertical_decimal(draw, expr, result, note="Thẳng hàng dấu phẩy"):
    draw.text((330, 260), expr, font=base.font(58, True), fill="#0f172a")
    draw.line((350, 365, 790, 365), fill="#1d4ed8", width=5)
    draw.text((520, 395), result, font=base.font(58, True), fill="#dc2626")
    base.label_box(draw, (330, 475, 850, 535), note, "#fef3c7", "#ca8a04", "#92400e", 24)


def draw_geometry(draw, k):
    if k in {"triangle", "triangle_area"}:
        pts = [(250, 470), (560, 220), (850, 470)]
        draw.polygon(pts, fill="#bfdbfe", outline="#1d4ed8")
        draw.line((560, 220, 560, 470), fill="#dc2626", width=5)
        draw.text((440, 490), "đáy 6 cm", font=base.font(28, True), fill="#0f172a")
        draw.text((580, 330), "cao 4 cm", font=base.font(28, True), fill="#dc2626")
        if k == "triangle_area":
            draw.text((690, 345), "S = 6 x 4 : 2", font=base.font(36, True), fill="#1d4ed8")
    elif k in {"trapezoid", "trapezoid_area"}:
        draw.polygon([(260, 470), (850, 470), (700, 250), (390, 250)], fill="#dcfce7", outline="#15803d")
        draw.text((460, 220), "đáy bé 4 cm", font=base.font(26, True), fill="#0f172a")
        draw.text((450, 490), "đáy lớn 8 cm", font=base.font(26, True), fill="#0f172a")
        draw.line((700, 250, 700, 470), fill="#dc2626", width=5)
        if k == "trapezoid_area":
            draw.text((735, 350), "S = (8+4)x5:2", font=base.font(32, True), fill="#1d4ed8")
    elif k in {"circle", "circle_perimeter", "circle_area"}:
        cx, cy, r = 460, 365, 150
        draw.ellipse((cx-r, cy-r, cx+r, cy+r), outline="#1d4ed8", width=6)
        draw.ellipse((cx-7, cy-7, cx+7, cy+7), fill="#dc2626")
        draw.line((cx-r, cy, cx+r, cy), fill="#dc2626", width=4)
        draw.line((cx, cy, cx+r, cy), fill="#15803d", width=5)
        text = "d = 10 cm" if k == "circle_perimeter" else ("r = 5 cm" if k == "circle_area" else "Tâm, bán kính, đường kính")
        draw.text((690, 335), text, font=base.font(38, True), fill="#1d4ed8")
    elif k in {"solid_shapes", "net_3d", "surface_area", "volume_intro", "cm3_dm3", "m3", "box_volume"}:
        draw.rectangle((210, 300, 430, 480), fill="#bfdbfe", outline="#1d4ed8", width=5)
        draw.polygon([(210, 300), (280, 245), (500, 245), (430, 300)], fill="#dbeafe", outline="#1d4ed8")
        draw.polygon([(430, 300), (500, 245), (500, 425), (430, 480)], fill="#93c5fd", outline="#1d4ed8")
        draw.rectangle((620, 285, 800, 465), fill="#dcfce7", outline="#15803d", width=5)
        draw.ellipse((900, 265, 1060, 465), outline="#dc2626", width=6)
        label = {
            "solid_shapes": "Hộp chữ nhật - lập phương - trụ",
            "net_3d": "Hình khai triển gồm các mặt",
            "surface_area": "Diện tích các mặt",
            "volume_intro": "Thể tích là phần chiếm chỗ",
            "cm3_dm3": "1 dm3 = 1000 cm3",
            "m3": "1 m3 = 1000 dm3",
            "box_volume": "V = dài x rộng x cao",
        }[k]
        draw.text((310, 520), label, font=base.font(32, True), fill="#1d4ed8")


def draw_scene_by_kind(draw, name, k):
    if k in {"ratio", "sum_ratio", "diff_ratio"}:
        draw.text((180, 265), "Sơ đồ tỉ số", font=base.font(42, True), fill="#0f172a")
        for i in range(3 if k != "ratio" else 5):
            draw.rectangle((245 + i * 90, 350, 325 + i * 90, 410), fill="#bfdbfe", outline="#1d4ed8", width=3)
        draw.text((260, 460), "Tỉ số 1 : 3 hoặc 3 : 5", font=base.font(34, True), fill="#1d4ed8")
    elif k == "dependency":
        base.draw_group_panel(draw, (180, 270, 430, 465), 6, "1 hộp: 6 bút", item="book")
        draw.text((550, 340), "4 hộp: 6 x 4 = 24 bút", font=base.font(38, True), fill="#1d4ed8")
    elif k == "mixed_number":
        for j, label in enumerate(["1", "1"]):
            x = 190 + j * 190
            draw.rectangle((x, 275, x + 145, 420), fill="#facc15", outline="#1d4ed8", width=3)
            base.draw_center(draw, (x + 72, 450), label, base.font(28, True), "#0f172a")
        x = 570
        for i in range(3):
            fill = "#facc15" if i == 0 else "#e0f2fe"
            draw.rectangle((x + i * 65, 275, x + (i + 1) * 65, 420), fill=fill, outline="#1d4ed8", width=3)
        draw.text((825, 320), "2 1/3", font=base.font(60, True), fill="#dc2626")
        draw.text((200, 500), "2 phần nguyên và 1 phần trong 3 phần bằng nhau", font=base.font(30, True), fill="#1d4ed8")
    elif k in {"decimal_fraction", "decimal_number", "decimal_equal", "decimal_compare", "decimal_round"}:
        draw_decimal_place(draw)
    elif k in {"hectare", "km2"}:
        unit = "1 ha = 10 000 m2" if k == "hectare" else "1 km2 = 1 000 000 m2"
        for r in range(4):
            for c in range(6):
                draw.rectangle((250+c*45, 260+r*45, 292+c*45, 302+r*45), fill="#bbf7d0", outline="#15803d", width=2)
        draw.text((600, 330), unit, font=base.font(42, True), fill="#15803d")
    elif k.startswith("decimal_"):
        mapping = {
            "decimal_add": ("12,35 + 4,6", "16,95"),
            "decimal_sub": ("12,35 - 4,6", "7,75"),
            "decimal_mul10": ("3,47 x 100", "347"),
            "decimal_mul_nat": ("2,35 x 4", "9,4"),
            "decimal_mul_dec": ("1,2 x 0,3", "0,36"),
            "decimal_div10": ("34,7 : 10", "3,47"),
            "decimal_div_nat": ("9,6 : 4", "2,4"),
            "decimal_div_dec": ("4,8 : 1,2", "4"),
            "decimal_measure": ("2 m 35 cm", "2,35 m"),
        }
        draw_vertical_decimal(draw, *mapping[k])
    elif k in {"percent_intro", "percent_ratio", "percent_value"}:
        draw.pieslice((260, 235, 560, 535), 0, 90, fill="#facc15", outline="#1d4ed8", width=4)
        draw.pieslice((260, 235, 560, 535), 90, 360, fill="#e0f2fe", outline="#1d4ed8", width=4)
        text = "25% = 25/100" if k == "percent_intro" else ("15 : 60 = 25%" if k == "percent_ratio" else "20% của 150 = 30")
        draw.text((650, 340), text, font=base.font(46, True), fill="#dc2626")
    elif k in {"calculator", "map_scale"}:
        if k == "calculator":
            base.rounded(draw, (330, 230, 650, 520), "#f8fafc", "#334155", 4, 24)
            draw.text((380, 280), "12,5 + 7,8", font=base.font(30, True), fill="#0f172a")
            draw.text((430, 355), "= 20,3", font=base.font(46, True), fill="#1d4ed8")
        else:
            draw.line((230, 385, 720, 385), fill="#1d4ed8", width=7)
            draw.text((350, 330), "5 cm trên bản đồ", font=base.font(32, True), fill="#0f172a")
            draw.text((760, 350), "Tỉ lệ 1 : 1000", font=base.font(40, True), fill="#dc2626")
    elif k in {"triangle", "triangle_area", "trapezoid", "trapezoid_area", "circle", "circle_perimeter", "circle_area", "solid_shapes", "net_3d", "surface_area", "volume_intro", "cm3_dm3", "m3", "box_volume"}:
        draw_geometry(draw, k)
    elif k in {"time_add_sub", "time_mul_div"}:
        expr = "2 giờ 35 phút + 1 giờ 40 phút" if k == "time_add_sub" else "1 giờ 20 phút x 3"
        result = "4 giờ 15 phút" if k == "time_add_sub" else "4 giờ"
        draw.text((190, 285), expr, font=base.font(42, True), fill="#0f172a")
        draw.text((470, 390), result, font=base.font(52, True), fill="#1d4ed8")
    elif k in {"speed", "distance_time"}:
        draw.line((180, 430, 1020, 430), fill="#64748b", width=14)
        base.label_box(draw, (210, 300, 360, 370), "A", "#dcfce7", "#15803d", "#166534", 32)
        base.label_box(draw, (840, 300, 990, 370), "B", "#fee2e2", "#dc2626", "#991b1b", 32)
        draw.text((460, 335), "120 km, 3 giờ", font=base.font(40, True), fill="#1d4ed8")
        draw.text((430, 485), "v = s : t", font=base.font(42, True), fill="#dc2626")
    elif k == "pie_chart":
        colors = ["#ef4444", "#f97316", "#eab308", "#22c55e"]
        start = 0
        for color, angle in zip(colors, [180, 90, 54, 36]):
            draw.pieslice((300, 220, 620, 540), start, start + angle, fill=color, outline="#ffffff", width=3)
            start += angle
        draw.text((700, 320), "Biểu đồ hình quạt", font=base.font(40, True), fill="#1d4ed8")
    elif k in {"data_representation", "frequency_event"}:
        if k == "data_representation":
            for i, count in enumerate([6, 4, 5]):
                x = 300 + i * 150
                draw.rectangle((x, 500-count*35, x+70, 500), fill=["#ef4444", "#f97316", "#eab308"][i], outline="#334155", width=2)
            draw.text((650, 330), "Bảng - cột - quạt", font=base.font(40, True), fill="#1d4ed8")
        else:
            for i, (label, count) in enumerate([("Mặt ngửa", 6), ("Mặt sấp", 4)]):
                y = 310 + i * 90
                draw.text((260, y), label, font=base.font(30, True), fill="#0f172a")
                for j in range(count):
                    draw.line((520+j*30, y, 520+j*30, y+45), fill="#1d4ed8", width=5)
                draw.text((780, y), f"{count} lần", font=base.font(28, True), fill="#1d4ed8")


def build_cards(lesson_name):
    spec = lesson_spec(lesson_name)
    return [
        {
            "type": "concept_scene",
            "title": "Quan sát",
            "kind": spec["kind"],
            "display_text": f"Quan sát tình huống minh họa cho bài: {lesson_name}.",
            "caption": spec["objective"],
            "student_task": "Nêu dữ kiện nhìn thấy trong tranh và gọi tên kiến thức chính.",
            "interaction": "choose_or_say",
            "expected_answer": spec["objective"],
            "wrong_hint": "Cho học sinh chỉ vào mô hình, số liệu, hình vẽ hoặc nhãn trong tranh rồi đọc lại yêu cầu.",
        },
        {
            "type": "worked_model",
            "title": "Làm mẫu",
            "kind": spec["kind"],
            "display_text": spec["model"],
            "caption": spec["model"],
            "student_task": "Đọc ví dụ mẫu và nhắc lại bước làm chính.",
            "interaction": "none",
            "expected_answer": spec["model"],
            "wrong_hint": "Tách bài thành từng bước nhỏ, đối chiếu với tranh minh họa rồi làm lại.",
        },
        {
            "type": "quick_try",
            "title": "Thử nhanh",
            "kind": spec["kind"],
            "display_text": spec["quick"],
            "caption": f"Đáp án cần đạt: {spec['answer']}",
            "student_task": "Tự làm một câu ngắn cùng dạng với ví dụ mẫu.",
            "interaction": "fill_blank",
            "expected_answer": spec["answer"],
            "wrong_hint": "Quay lại thẻ làm mẫu, thay dữ kiện mới vào đúng vị trí rồi tính lại.",
        },
    ]


def build_blueprint():
    chapters = []
    for chapter in CHAPTERS:
        lessons = []
        for name in chapter["lessons"]:
            spec = lesson_spec(name)
            lessons.append({"lesson": name, "kind": spec["kind"], "objective": spec["objective"], "cards": build_cards(name)})
        chapters.append({"title": chapter["title"], "lessons": lessons})
    return {
        "grade": 5,
        "source_textbook": "SGK Toán 5 Cánh Diều",
        "version_note": "Bản chi tiết: mỗi bài có 3 tranh minh họa riêng cho quan sát, làm mẫu và thử nhanh.",
        "cognitive_profile": {
            "reading_level": "Học sinh lớp 5 đọc hiểu tốt, có thể theo dõi công thức và bài nhiều bước, nhưng vẫn cần sơ đồ và mô hình trực quan.",
            "learning_style": "Quan sát mô hình, nhận ra công thức/quy tắc, theo ví dụ mẫu, sau đó tự làm một câu cùng dạng.",
            "ai_policy": "Có thể dùng AI ở mức gợi ý từng bước, hỏi dẫn dắt và kiểm tra lỗi; không đưa ngay đáp án.",
        },
        "chapters": chapters,
    }


def render_card_image(chapter_index, lesson_index, card_index, lesson_name, card):
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1200, 675), "#f1f5f9")
    draw = ImageDraw.Draw(img)
    header(draw, lesson_name, card["title"])
    if card_index == 3:
        draw_quick_scene(draw, card)
    else:
        draw_scene_by_kind(draw, lesson_name, card["kind"])
    base.rounded(draw, (78, 560, 1122, 618), "#f8fafc", "#e2e8f0", 2, 18)
    base.draw_wrapped(draw, card["caption"], (105, 575), base.font(22, True), fill="#0f172a", width=84, spacing=4)
    out = IMAGE_DIR / f"grade5-c{chapter_index:02d}-l{lesson_index:02d}-card{card_index:02d}.png"
    img.save(out)
    return out


def build_docx(data):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    base.set_defaults(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("LÝ THUYẾT TOÁN LỚP 5")
    g4.set_run_font(r, size=24, bold=True, color="1F4E79")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Bản chi tiết theo thẻ, có tranh tình huống, đáp án và gợi ý sai - Codex")
    g4.set_run_font(r, size=13, italic=True)
    g4.add_meta(doc, "Nguồn tham khảo", data["source_textbook"])
    g4.add_meta(doc, "Ghi chú phiên bản", data["version_note"])
    doc.add_heading("Định hướng thiết kế lớp 5", level=1)
    g4.add_bullets(
        doc,
        [
            data["cognitive_profile"]["reading_level"],
            data["cognitive_profile"]["learning_style"],
            data["cognitive_profile"]["ai_policy"],
            "Mỗi bài có 3 tranh: quan sát tình huống, làm mẫu theo bước và thử nhanh có đáp án.",
        ],
    )
    doc.add_heading("Tổng quan nội dung", level=1)
    g4.add_overview(doc, data)
    for chapter_index, chapter in enumerate(data["chapters"], 1):
        doc.add_page_break()
        doc.add_heading(f"Chủ đề {chapter_index}: {chapter['title']}", level=1)
        for lesson_index, lesson in enumerate(chapter["lessons"], 1):
            doc.add_heading(f"{chapter_index}.{lesson_index}. {lesson['lesson']}", level=2)
            g4.add_meta(doc, "Mục tiêu bài học", lesson["objective"])
            for card_index, card in enumerate(lesson["cards"], 1):
                image_path = render_card_image(chapter_index, lesson_index, card_index, lesson["lesson"], card)
                doc.add_heading(f"Thẻ {card_index}: {card['title']}", level=3)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path), width=Inches(6.45))
                g4.add_bullets(
                    doc,
                    [
                        f"Nội dung hiển thị: {card['display_text']}",
                        f"Nhiệm vụ học sinh: {card['student_task']}",
                        f"Tương tác đề xuất: {card['interaction']}",
                        f"Đáp án/Kết quả mong đợi: {card['expected_answer']}",
                        f"Gợi ý khi học sinh sai: {card['wrong_hint']}",
                    ],
                )
            doc.add_paragraph()
    try:
        doc.save(OUTPUT_PATH)
        return OUTPUT_PATH
    except PermissionError:
        doc.save(FALLBACK_OUTPUT_PATH)
        return FALLBACK_OUTPUT_PATH


def build_contact_sheets():
    paths = sorted(IMAGE_DIR.glob("grade5-*.png"))
    outs = []
    for start in range(0, len(paths), 60):
        chunk = paths[start : start + 60]
        cols = 5
        rows = math.ceil(len(chunk) / cols)
        sheet = Image.new("RGB", (cols * 240, rows * 170), "#ffffff")
        draw = ImageDraw.Draw(sheet)
        for i, path in enumerate(chunk):
            img = Image.open(path).resize((240, 135))
            x = (i % cols) * 240
            y = (i // cols) * 170
            sheet.paste(img, (x, y))
            draw.text((x + 6, y + 140), path.stem.replace("grade5-", ""), font=base.font(16, True), fill="#0f172a")
        out = IMAGE_DIR / f"contact-sheet-{start + 1}-{start + len(chunk)}.png"
        sheet.save(out)
        outs.append(out)
    return outs


def main():
    data = build_blueprint()
    BLUEPRINT_PATH.parent.mkdir(parents=True, exist_ok=True)
    BLUEPRINT_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    output = build_docx(data)
    sheets = build_contact_sheets()
    print(BLUEPRINT_PATH)
    print(output)
    for sheet in sheets:
        print(sheet)
    print(f"lessons={sum(len(c['lessons']) for c in data['chapters'])}")
    print(f"cards={sum(len(l['cards']) for c in data['chapters'] for l in c['lessons'])}")
    print(f"images={len(list(IMAGE_DIR.glob('grade5-*.png')))}")


if __name__ == "__main__":
    main()
