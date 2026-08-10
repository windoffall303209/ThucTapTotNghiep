# Script rebuild docx with images hỗ trợ nhập, xuất, kiểm tra hoặc bảo trì dữ liệu và cấu hình của dự án.
import os
import sys
import json
import hashlib
import requests
from pathlib import Path
from urllib.parse import urlparse
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure UTF-8 console output
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
CRAWLED_THEORY_JSON = ROOT / "output" / "doc" / "crawled_theory_cards.json"
GRADE4_THEORY_JSON = ROOT / "output" / "doc" / "grade4_theory_cards.json"
IMAGE_DIR = ROOT / "output" / "doc" / "theory_images"
OUTPUT_DOCX = ROOT / "output" / "doc" / "toan_4_ly_thuyet_canh_dieu_co_hinh_ve.docx"

# Hàm download_image dùng để lấy dữ liệu và xử lý trường hợp không tìm thấy kết quả; cần bảo toàn hợp đồng đầu vào và giá trị trả về của luồng gọi.

def download_image(session, image_url, cache_dir):
    cache_dir.mkdir(parents=True, exist_ok=True)
    h = hashlib.sha1(image_url.encode('utf-8')).hexdigest()
    ext = os.path.splitext(urlparse(image_url).path)[1]
    if not ext or len(ext) > 5:
        ext = ".png"
    file_path = cache_dir / f"{h}{ext}"
    if file_path.exists():
        return file_path
    
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        }
        r = session.get(image_url, headers=headers, timeout=20)
        if r.status_code == 200:
            file_path.write_bytes(r.content)
            return file_path
    except Exception as e:
        print(f"      [Lỗi tải ảnh] {image_url}: {e}")
    return None

# Hàm build_docx_with_images dùng để xây dựng kết quả từ các nguồn dữ liệu và quy tắc liên quan; cần bảo toàn hợp đồng đầu vào và giá trị trả về của luồng gọi.

def build_docx_with_images(lessons_theory, output_path):
    print(f"Đang xây dựng file Word có hình vẽ tại: {output_path}...")
    doc = Document()
    session = requests.Session()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("TỔNG HỢP LÝ THUYẾT TOÁN LỚP 4\n")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(20)
    title_run.bold = True
    
    subtitle_run = title_p.add_run("Bộ Sách Giáo Khoa Cánh Diều — Gợi Ý Học Tập & Hình Minh Họa")
    subtitle_run.font.name = 'Arial'
    subtitle_run.font.size = Pt(12)
    subtitle_run.italic = True
    
    doc.add_paragraph("\n")
    
    # Group by Chapter
    chapters = {}
    for item in lessons_theory:
        ch_name = item.get("chapter_name", "Khác")
        if ch_name not in chapters:
            chapters[ch_name] = []
        chapters[ch_name].append(item)
        
    for ch_name, ch_lessons in chapters.items():
        # Chapter Heading
        ch_h = doc.add_heading(ch_name, level=1)
        ch_h.style.font.name = 'Arial'
        ch_h.style.font.size = Pt(16)
        ch_h.style.font.bold = True
        
        for lesson in ch_lessons:
            # Lesson Heading
            l_h = doc.add_heading(lesson["lesson_name"], level=2)
            l_h.style.font.name = 'Arial'
            l_h.style.font.size = Pt(13)
            l_h.style.font.bold = True
            
            cards = lesson.get("theory_cards", [])
            if not cards:
                p = doc.add_paragraph("Chưa có lý thuyết cho bài học này.")
                p.italic = True
                continue
                
            for idx, card in enumerate(cards, 1):
                # Card Title
                c_p = doc.add_paragraph()
                c_run = c_p.add_run(f"• Thẻ {idx}: {card.get('title', '')}")
                c_run.font.name = 'Arial'
                c_run.font.size = Pt(11.5)
                c_run.bold = True
                
                # Card Body
                body = card.get("body", "")
                body_p = doc.add_paragraph()
                body_p.paragraph_format.left_indent = Inches(0.2)
                
                if isinstance(body, list):
                    body_text = "\n".join([f"- {item}" for item in body])
                else:
                    body_text = str(body)
                    
                body_run = body_p.add_run(body_text)
                body_run.font.name = 'Arial'
                
                # Card Formula
                formula = card.get("formula", "").strip()
                if formula:
                    form_p = doc.add_paragraph()
                    form_p.paragraph_format.left_indent = Inches(0.4)
                    form_run = form_p.add_run(f"Công thức/Quy tắc: {formula}")
                    form_run.font.name = 'Arial'
                    form_run.italic = True
                    form_run.bold = True
                    
                # Card Example
                example = card.get("example", "").strip()
                if example:
                    ex_p = doc.add_paragraph()
                    ex_p.paragraph_format.left_indent = Inches(0.4)
                    ex_run = ex_p.add_run(f"Ví dụ: {example}")
                    ex_run.font.name = 'Arial'
                    ex_run.italic = True
                
                # Card Images
                card_images = card.get("images", [])
                for img in card_images:
                    img_url = img.get("url")
                    if img_url:
                        # Try download
                        img_path = download_image(session, img_url, IMAGE_DIR)
                        if img_path and img_path.exists():
                            try:
                                img_p = doc.add_paragraph()
                                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                img_p.paragraph_format.left_indent = Inches(0.4)
                                doc.add_picture(str(img_path), width=Inches(3.8))
                                
                                caption_p = doc.add_paragraph()
                                caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                caption_run = caption_p.add_run(f"Hình vẽ: {img.get('alt_text', 'Hình minh họa lý thuyết')}")
                                caption_run.font.name = 'Arial'
                                caption_run.font.size = Pt(9.5)
                                caption_run.italic = True
                            except Exception as img_err:
                                print(f"      [Lỗi chèn ảnh vào Word] {img_path}: {img_err}")
                    
            doc.add_paragraph()  # Blank line between lessons
            
    doc.save(output_path)
    print(f"Đã lưu file Word thành công tại: {output_path}")

# Hàm main dùng để thực hiện logic nghiệp vụ chính và trả kết quả cho luồng gọi; cần bảo toàn hợp đồng đầu vào và giá trị trả về của luồng gọi.

def main():
    if not CRAWLED_THEORY_JSON.exists():
        print(f"Lỗi: Không tìm thấy file lý thuyết thô tại {CRAWLED_THEORY_JSON}")
        sys.exit(1)
        
    if not GRADE4_THEORY_JSON.exists():
        print(f"Lỗi: Không tìm thấy file lý thuyết Lớp 4 đã sinh tại {GRADE4_THEORY_JSON}")
        sys.exit(1)
        
    # 1. Load data
    with open(CRAWLED_THEORY_JSON, "r", encoding="utf-8") as f:
        crawled_list = json.load(f)
    
    with open(GRADE4_THEORY_JSON, "r", encoding="utf-8") as f:
        grade4_list = json.load(f)
        
    crawled_by_id = {item["lesson_id"]: item for item in crawled_list}
    
    # 2. Map images to first card
    print("Đang map hình ảnh từ crawled_theory_cards vào grade4_theory_cards...")
    updated_count = 0
    for lesson in grade4_list:
        lesson_id = lesson["lesson_id"]
        crawled_item = crawled_by_id.get(lesson_id)
        
        # Collect all images from crawled cards
        images = []
        if crawled_item:
            for card in crawled_item.get("theory_cards", []):
                if card.get("images"):
                    for img in card["images"]:
                        if 'alt_text' not in img:
                            img['alt_text'] = img.get('alt', 'Hình minh họa lý thuyết')
                        images.append(img)
                        
        theory_cards = lesson.get("theory_cards", [])
        if theory_cards:
            # Set images array for all cards (empty by default, attach to first one)
            for i, card in enumerate(theory_cards):
                if i == 0:
                    card["images"] = images
                else:
                    card["images"] = []
            updated_count += 1
            
    # 3. Save updated JSON
    with open(GRADE4_THEORY_JSON, "w", encoding="utf-8") as f:
        json.dump(grade4_list, f, ensure_ascii=False, indent=2)
    print(f"Đã cập nhật {updated_count} bài học với hình vẽ vào file {GRADE4_THEORY_JSON}.")
    
    # 4. Rebuild Word
    build_docx_with_images(grade4_list, OUTPUT_DOCX)

# Khối này tập trung xử lý nhánh nghiệp vụ và bảo toàn các điều kiện an toàn.
if __name__ == "__main__":
    main()
