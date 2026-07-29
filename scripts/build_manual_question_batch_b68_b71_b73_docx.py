"""Dàn trang bộ 66 câu đã biên soạn thủ công thành DOCX.

Script chỉ đọc nội dung có sẵn trong JSON, chèn 33 ảnh đã được duyệt và áp dụng
định dạng Word. Script không sinh hoặc viết lại câu hỏi, đáp án hay lời giải.
"""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "manual_question_batches" / "grade5_b68_b71_b73.json"
ASSETS = ROOT / "output" / "doc" / "assets" / "grade5-b68-b71-b73" / "final"
OUTPUT = ROOT / "output" / "doc" / "Ngan_hang_cau_hoi_Toan_5_Bai_68_71_73_66_cau.docx"

LESSON_TITLES = {
    68: "Bài 68. Ôn tập về các đơn vị đo thời gian",
    71: "Bài 71. Ôn tập về các phép tính với số đo thời gian",
    73: "Bài 73. Vận tốc",
}


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=110, start=150, bottom=110, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(31, 78, 121)
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.size = Pt(17)
    doc.styles["Heading 2"].font.size = Pt(13)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("NGÂN HÀNG CÂU HỎI TOÁN 5 - CÁNH DIỀU")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc, payload):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(92)
    r = p.add_run("NGÂN HÀNG CÂU HỎI BỔ SUNG")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TOÁN 5 - CÁNH DIỀU")
    r.bold = True
    r.font.size = Pt(29)
    r.font.color.rgb = RGBColor(237, 125, 49)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Bài 68 • Bài 71 • Bài 73")
    r.bold = True
    r.font.size = Pt(19)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    r = p.add_run("66 CÂU HỎI TỰ BIÊN SOẠN")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("22 câu mỗi bài • 33 câu có hình minh họa riêng biệt • 50% câu có hình\n").bold = True
    p.add_run("Mỗi câu có đề bài, phần trả lời, đáp án và lời giải chi tiết.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run(payload["source"])
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_page_break()


def add_summary(doc, questions):
    doc.add_heading("THÔNG TIN BỘ CÂU HỎI", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Bài học", "Số câu", "Câu có hình", "Tỉ lệ"]
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = text
        set_shading(cell, "D9EAF7")
        cell.paragraphs[0].runs[0].bold = True
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for lesson in (68, 71, 73):
        items = [q for q in questions if q["lesson"] == lesson]
        illustrated = sum(1 for q in items if q.get("image"))
        cells = table.add_row().cells
        cells[0].text = LESSON_TITLES[lesson]
        cells[1].text = str(len(items))
        cells[2].text = str(illustrated)
        cells[3].text = f"{illustrated / len(items):.0%}"
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

    doc.add_paragraph()
    notes = [
        "Nội dung được biên soạn thủ công, bám theo dạng bài trong SGK và không sao chép nguyên văn.",
        "Mỗi câu có hình sử dụng một ảnh riêng để có thể đảo thứ tự câu hỏi độc lập.",
        "Dữ kiện về thời gian, quãng đường và số lượng được ghi trực tiếp trên hình.",
        "Các mặt đồng hồ kim đều có đủ 12 số giờ theo đúng thứ tự.",
        "Ảnh minh họa được tạo bằng AI, nền sáng, không logo và không watermark.",
    ]
    for note in notes:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(note)
    doc.add_page_break()


def add_label_box(doc, label, text, fill, label_color):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_shading(cell, fill)
    set_cell_margins(cell, top=105, start=150, bottom=105, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(label_color)
    p.add_run(text)


def add_question(doc, question):
    lesson = question["lesson"]
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(LESSON_TITLES[lesson])
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(89, 89, 89)

    heading = doc.add_heading(f"Câu {question['number']}. {question['type']}", level=2)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True

    prompt = doc.add_paragraph()
    prompt.paragraph_format.space_after = Pt(5)
    prompt.add_run(question["question"]).bold = True

    if question.get("image"):
        image_path = ASSETS / question["image"].replace(".png", ".jpg")
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        p.add_run().add_picture(str(image_path), width=Cm(13.6))

    if question.get("choices"):
        table = doc.add_table(rows=0, cols=1)
        for choice in question["choices"]:
            cell = table.add_row().cells[0]
            set_cell_margins(cell, top=45, start=180, bottom=45, end=100)
            cell.paragraphs[0].add_run(choice)
    else:
        add_label_box(doc, "Phần trả lời: ", question["response"], "F2F2F2", "595959")

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_label_box(doc, "Đáp án: ", question["answer"], "E2F0D9", "008000")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_label_box(doc, "Lời giải chi tiết: ", question["explanation"], "EAF2F8", "1F4E78")


def validate(questions):
    if len(questions) != 66:
        raise ValueError(f"Cần đúng 66 câu, hiện có {len(questions)}")
    image_names = []
    for lesson in (68, 71, 73):
        items = [q for q in questions if q["lesson"] == lesson]
        if len(items) != 22:
            raise ValueError(f"Bài {lesson} phải có đúng 22 câu")
        illustrated = [q for q in items if q.get("image")]
        if len(illustrated) != 11:
            raise ValueError(f"Bài {lesson} phải có đúng 11 câu có hình")
        image_names.extend(q["image"] for q in illustrated)
    if len(image_names) != len(set(image_names)):
        raise ValueError("Mỗi câu có hình phải dùng một tệp ảnh riêng")


def main():
    payload = json.loads(DATA.read_text(encoding="utf-8"))
    questions = payload["questions"]
    validate(questions)

    doc = Document()
    configure(doc)
    add_cover(doc, payload)
    add_summary(doc, questions)

    for index, question in enumerate(questions):
        add_question(doc, question)
        if index < len(questions) - 1:
            doc.add_page_break()

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
