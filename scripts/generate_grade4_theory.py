import os
import sys
import json
import time
import re
import requests
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure utf-8 output encoding for Windows terminal
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
ENV_PATH = ROOT / ".env"
LESSONS_JSON = ROOT / "output" / "doc" / "current_lessons_for_theory.json"
CACHE_JSON = ROOT / "output" / "doc" / "grade4_theory_cards.json"
CRAWLED_THEORY_JSON = ROOT / "output" / "doc" / "crawled_theory_cards.json"
OUTPUT_DOCX = ROOT / "output" / "doc" / "toan_4_ly_thuyet_canh_dieu.docx"

def load_env(env_path):
    env_vars = {}
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                parts = line.split("=", 1)
                if len(parts) == 2:
                    env_vars[parts[0].strip()] = parts[1].strip()
    return env_vars

env = load_env(ENV_PATH)
API_KEY = env.get("NVIDIA_NIM_API_KEY", "")
BASE_URL = env.get("NVIDIA_NIM_BASE_URL", "https://integrate.api.nvidia.com/v1").rstrip("/")
MODEL = env.get("NVIDIA_NIM_MODEL", "meta/llama-3.3-70b-instruct")

def call_nvidia_nim(prompt, retries=3):
    url = f"{BASE_URL}/chat/completions"
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {API_KEY}"
    }
    payload = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": "You are an educational assistant that outputs raw JSON content only. Do not enclose the JSON in markdown code blocks like ```json ... ```."},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.2
    }
    
    for attempt in range(1, retries + 1):
        try:
            response = requests.post(url, headers=headers, json=payload, timeout=45)
            if response.status_code == 200:
                data = response.json()
                content = data.get("choices", [{}])[0].get("message", {}).get("content", "").strip()
                return content
            else:
                print(f"      [Attempt {attempt}] HTTP Error {response.status_code}: {response.text}")
        except Exception as e:
            print(f"      [Attempt {attempt}] Exception: {str(e)}")
        
        if attempt < retries:
            time.sleep(3 * attempt)
            
    raise Exception("Failed to call API after all retries.")

def fix_json_backslashes(text):
    marker_double_backslash = "___DOUBLE_BACKSLASH_MARKER___"
    marker_escaped_quote = "___ESCAPED_QUOTE_MARKER___"
    
    processed = text.replace("\\\\", marker_double_backslash)
    processed = processed.replace('\\"', marker_escaped_quote)
    
    # Now all remaining backslashes are single backslashes that need to be doubled
    processed = processed.replace("\\", "\\\\")
    
    # Restore the valid escapes
    processed = processed.replace(marker_escaped_quote, '\\"')
    processed = processed.replace(marker_double_backslash, "\\\\")
    
    return processed

def parse_ai_response(content):
    # Try to find JSON array brackets
    start_idx = content.find("[")
    end_idx = content.rfind("]")
    
    if start_idx == -1 or end_idx == -1 or end_idx < start_idx:
        # Check if it's a single JSON object instead of an array
        start_idx = content.find("{")
        end_idx = content.rfind("}")
        if start_idx != -1 and end_idx != -1 and end_idx > start_idx:
            json_str = fix_json_backslashes(content[start_idx:end_idx+1])
            obj = json.loads(json_str, strict=False)
            return [obj]
        raise Exception("No JSON array or object found in response.")
        
    json_str = fix_json_backslashes(content[start_idx:end_idx+1])
    return json.loads(json_str, strict=False)

def get_mathematical_context(lesson_name):
    ln = lesson_name.lower()
    
    # 1. Measurement Units
    if any(k in ln for k in ["yến", "tạ", "tấn"]):
        return """SỰ THẬT TOÁN HỌC (QUY ĐỔI ĐƠN VỊ KHỐI LƯỢNG):
- $1\\text{ yến} = 10\\text{ kg}$
- $1\\text{ tạ} = 100\\text{ kg} = 10\\text{ yến}$
- $1\\text{ tấn} = 1000\\text{ kg} = 10\\text{ tạ} = 100\\text{ yến}$"""
        
    if "giây" in ln:
        return """SỰ THẬT TOÁN HỌC (QUY ĐỔI THỜI GIAN):
- $1\\text{ phút} = 60\\text{ giây}$
- $1\\text{ giờ} = 60\\text{ phút} = 3600\\text{ giây}$"""
        
    if "thế kỉ" in ln or "thế kỷ" in ln:
        return """SỰ THẬT TOÁN HỌC (QUY ĐỔI THỜI GIAN):
- $1\\text{ thế kỉ} = 100\\text{ năm}$
- Để xác định thế kỉ của một năm:
  + Từ năm $1$ đến năm $100$ là thế kỉ $1$ (I).
  + Từ năm $101$ đến năm $200$ là thế kỉ $2$ (II).
  + Tổng quát: Năm có dạng $ab\\dots xy$, nếu $xy > 00$, thế kỉ = $ab\\dots + 1$. Nếu $xy = 00$, thế kỉ = $ab\\dots$. Ví dụ năm $1900$ thuộc thế kỉ $19$, năm $2026$ thuộc thế kỉ $21$ ($20+1$)."""

    if "mét vuông" in ln or "đề-xi-mét vuông" in ln or "mi-li-mét vuông" in ln or "diện tích" in ln:
        return """SỰ THẬT TOÁN HỌC (QUY ĐỔI ĐƠN VỊ DIỆN TÍCH):
- $1\\text{ m}^2 = 100\\text{ dm}^2 = 10\\ 000\\text{ cm}^2 = 1\\ 000\\ 000\\text{ mm}^2$
- $1\\text{ dm}^2 = 100\\text{ cm}^2 = 10\\ 000\\text{ mm}^2$
- $1\\text{ cm}^2 = 100\\text{ mm}^2$"""

    # 2. Geometry
    if "góc nhọn" in ln or "góc tù" in ln or "góc bẹt" in ln or "góc" in ln:
        return """SỰ THẬT TOÁN HỌC (CÁC LOẠI GÓC):
- Góc vuông: Bằng $90^\\circ$ (góc đỉnh ê-ke).
- Góc nhọn: Bé hơn góc vuông (nhỏ hơn $90^\\circ$).
- Góc tù: Lớn hơn góc vuông nhưng bé hơn góc bẹt (nằm giữa $90^\\circ$ và $180^\\circ$).
- Góc bẹt: Bằng hai góc vuông (bằng $180^\\circ$)."""

    if "vuông góc" in ln:
        return """SỰ THẬT TOÁN HỌC (ĐƯỜNG THẲNG VUÔNG GÓC):
- Hai đường thẳng cắt nhau tạo thành một góc vuông ($90^\\circ$) được gọi là hai đường thẳng vuông góc với nhau."""

    if "song song" in ln:
        return """SỰ THẬT TOÁN HỌC (ĐƯỜNG THẲNG SONG SONG):
- Hai đường thẳng song song là hai đường thẳng không bao giờ cắt nhau (dù kéo dài vô tận)."""

    if "bình hành" in ln:
        return """SỰ THẬT TOÁN HỌC (HÌNH BÌNH HÀNH):
- Hình bình hành có hai cặp cạnh đối diện song song và bằng nhau.
- Công thức tính diện tích hình bình hành: $S = a \\times h$ (với $a$ là độ dài đáy, $h$ là chiều cao tương ứng)."""

    if "thoi" in ln:
        return """SỰ THẬT TOÁN HỌC (HÌNH THOI):
- Hình thoi có 4 cạnh bằng nhau, hai đường chéo vuông góc với nhau tại trung điểm của mỗi đường.
- Công thức tính diện tích hình thoi: $S = \\frac{m \\times n}{2}$ (với $m$ và $n$ là độ dài hai đường chéo)."""

    # 3. Arithmetic Word Problems
    if "trung bình cộng" in ln:
        return """SỰ THẬT TOÁN HỌC (TÌM SỐ TRUNG BÌNH CỘNG):
- Số trung bình cộng của một nhóm số = (Tổng các số đó) : (Số lượng các số hạng)
- Công thức: $\\text{TBC} = (a_1 + a_2 + \\dots + a_n) : n$"""

    if "tổng và hiệu" in ln:
        return """SỰ THẬT TOÁN HỌC (TÌM HAI SỐ KHI BIẾT TỔNG VÀ HIỆU):
- Công thức tìm Số lớn: $\\text{Số lớn} = (\\text{Tổng} + \\text{Hiệu}) : 2$
- Công thức tìm Số bé: $\\text{Số bé} = (\\text{Tổng} - \\text{Hiệu}) : 2$
- Có thể tính Số bé trước rồi tính Số lớn: $\\text{Số lớn} = \\text{Số bé} + \\text{Hiệu}$ hoặc $\\text{Số lớn} = \\text{Tổng} - \\text{Số bé}$"""

    # 4. Fractions
    if "phân số bằng nhau" in ln or "tính chất cơ bản của phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (TÍNH CHẤT CƠ BẢN CỦA PHÂN SỐ):
- Nếu nhân cả tử số và mẫu số của một phân số với cùng một số tự nhiên khác $0$ thì được một phân số bằng phân số đã cho: $\\frac{a}{b} = \\frac{a \\times c}{b \\times c}$ ($c \\neq 0$).
- Nếu chia cả tử số và mẫu số của một phân số cho cùng một số tự nhiên khác $0$ (nếu cả tử và mẫu đều chia hết cho số đó) thì được một phân số bằng phân số đã cho: $\\frac{a}{b} = \\frac{a : c}{b : c}$ ($c \\neq 0$)."""

    if "rút gọn phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (RÚT GỌN PHÂN SỐ):
- Chia cả tử số và mẫu số của phân số cho cùng một số tự nhiên lớn hơn $1$ mà cả tử và mẫu đều chia hết, làm như vậy cho đến khi không rút gọn được nữa (ta được phân số tối giản).
- Phân số tối giản là phân số mà tử số và mẫu số không cùng chia hết cho số tự nhiên nào lớn hơn $1$."""

    if "quy đồng" in ln:
        return """SỰ THẬT TOÁN HỌC (QUY ĐỒNG MẪU SỐ PHÂN SỐ):
- Quy đồng mẫu số hai phân số là biến đổi chúng thành hai phân số mới bằng hai phân số đã cho nhưng có cùng mẫu số.
- Cách quy đồng cơ bản: nhân cả tử và mẫu của phân số thứ nhất với mẫu của phân số thứ hai, nhân cả tử và mẫu của phân số thứ hai với mẫu của phân số thứ nhất."""

    if "so sánh hai phân số" in ln or "so sánh phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (SO SÁNH PHÂN SỐ):
- So sánh cùng mẫu số: Phân số nào có tử số lớn hơn thì phân số đó lớn hơn.
- So sánh khác mẫu số: Ta phải quy đồng mẫu số hai phân số đó về cùng một mẫu số rồi so sánh hai phân số cùng mẫu số mới tạo thành.
- So sánh với $1$: Phân số có tử số lớn hơn mẫu số thì lớn hơn $1$; phân số có tử số bé hơn mẫu số thì bé hơn $1$; phân số có tử số bằng mẫu số thì bằng $1$."""

    if "cộng các phân số" in ln or "cộng hai phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (CỘNG PHÂN SỐ):
- Cộng cùng mẫu số: Ta cộng hai tử số với nhau và giữ nguyên mẫu số: $\\frac{a}{c} + \\frac{b}{c} = \\frac{a+b}{c}$.
- Cộng khác mẫu số: Ta quy đồng mẫu số hai phân số đó về cùng mẫu số rồi thực hiện cộng hai phân số cùng mẫu số mới tạo thành."""

    if "trừ các phân số" in ln or "trừ hai phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (TRỪ PHÂN SỐ):
- Trừ cùng mẫu số: Ta lấy tử số của phân số thứ nhất trừ đi tử số của phân số thứ hai và giữ nguyên mẫu số: $\\frac{a}{c} - \\frac{b}{c} = \\frac{a-b}{c}$.
- Trừ khác mẫu số: Ta quy đồng mẫu số hai phân số đó về cùng mẫu số rồi thực hiện trừ hai phân số cùng mẫu số mới tạo thành."""

    if "nhân phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (NHÂN PHÂN SỐ):
- Muốn nhân hai phân số, ta lấy tử số nhân với tử số, mẫu số nhân với mẫu số: $\\frac{a}{b} \\times \\frac{c}{d} = \\frac{a \\times c}{b \\times d}$."""

    if "chia phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (CHIA PHÂN SỐ):
- Muốn chia một phân số cho một phân số, ta lấy phân số thứ nhất nhân với phân số thứ hai đảo ngược: $\\frac{a}{b} : \\frac{c}{d} = \\frac{a}{b} \\times \\frac{d}{c} = \\frac{a \\times d}{b \\times c}$."""

    if "phân số của một số" in ln:
        return """SỰ THẬT TOÁN HỌC (TÌM PHÂN SỐ CỦA MỘT SỐ):
- Muốn tìm phân số $\\frac{m}{n}$ của số $A$, ta tính $A \\times \\frac{m}{n}$ (lấy $A$ nhân với tử số $m$ rồi chia cho mẫu số $n$ hoặc lấy $A$ chia cho $n$ rồi nhân với $m$)."""

    if "phân số" in ln:
        return """SỰ THẬT TOÁN HỌC (KHÁI NIỆM PHÂN SỐ):
- Phân số viết dưới dạng $\\frac{a}{b}$ với $a, b$ là các số tự nhiên, $b \\neq 0$.
- $a$ là tử số (nằm trên gạch ngang) chỉ số phần được lấy ra.
- $b$ là mẫu số (nằm dưới gạch ngang) chỉ tổng số phần bằng nhau được chia ra.
- Phép chia số tự nhiên $a : b$ ($b \\neq 0$) có thể viết dưới dạng phân số là $\\frac{a}{b}$."""

    if "các số trong phạm vi 1 000 000" in ln or "các số có nhiều chữ số" in ln or "số tự nhiên" in ln:
        return """SỰ THẬT TOÁN HỌC (HÀNG VÀ LỚP CỦA SỐ TỰ NHIÊN):
- Các hàng của Lớp Đơn Vị: Hàng đơn vị, Hàng chục, Hàng trăm.
- Các hàng của Lớp Nghìn: Hàng nghìn, Hàng chục nghìn, Hàng trăm nghìn.
- Các hàng của Lớp Triệu: Hàng triệu, Hàng chục triệu, Hàng trăm triệu.
- Để đọc số có nhiều chữ số: ta tách số thành lớp (từ phải qua trái, mỗi lớp 3 chữ số), sau đó đọc từng lớp từ trái qua phải kèm tên lớp (tên lớp triệu, nghìn; lớp đơn vị không đọc tên lớp).
- Ví dụ số $4\\ 567\\ 890$ gồm: $4$ triệu (hàng triệu), $5$ trăm nghìn (hàng trăm nghìn), $6$ chục nghìn (hàng chục nghìn), $7$ nghìn (hàng nghìn), $8$ trăm (hàng trăm), $9$ chục (hàng chục), $0$ đơn vị (hàng đơn vị).
- Số có 7 chữ số thì chữ số ngoài cùng bên trái thuộc hàng triệu (Lớp Triệu). Số $1\\ 000\\ 000$ có 7 chữ số (hàng triệu là 1, các hàng khác là 0)."""

    return ""

def generate_theory_prompt(grade, chapter_name, lesson_name, raw_context):
    math_context = get_mathematical_context(lesson_name)
    context_str = f"\n{math_context}\n" if math_context else ""
    
    raw_context_str = ""
    if raw_context.strip():
        raw_context_str = f"""
Dưới đây là nội dung lý thuyết thô thu thập được từ trang giáo dục (có thể chứa quảng cáo hoặc viết chưa đầy đủ):
---
{raw_context}
---
Dựa vào nội dung thô này, hãy làm sạch quảng cáo và biên soạn lại nội dung chính xác.
"""

    return f"""Bạn là một giáo viên dạy Toán tiểu học xuất sắc tại Việt Nam, am hiểu sâu sắc về chương trình Sách giáo khoa Toán lớp 4 (bộ sách Cánh Diều).
{context_str}
{raw_context_str}
Hãy biên soạn nội dung lý thuyết chi tiết, chuẩn xác và trực quan cho bài học sau:
- Khối lớp: Lớp {grade}
- Chương: {chapter_name}
- Bài học: {lesson_name}

YÊU CẦU CỰC KỲ QUAN TRỌNG VỀ TOÁN HỌC & CẤU TRÚC JSON (BẮT BUỘC TUÂN THỦ):
1. KHÔNG được sử dụng các ký hiệu toán học quá phức tạp vượt cấp như logarit (\\log), tổng (\\sum), giới hạn, lũy thừa lớn, v.v. Học sinh lớp 4 chỉ học các phép tính cộng, trừ, nhân, chia cơ bản.
2. Hãy cực kỳ cẩn thận về các hàng số và lớp số. Đảm bảo mọi tính toán, đếm số chữ số và quy đổi đơn vị đo lường trong ví dụ và công thức phải chính xác 100% chuẩn toán học tiểu học.
3. Mọi công thức, phép tính và ký hiệu toán học trong body, formula, example bắt buộc phải bọc trong dấu đô-la ($).
4. TUYỆT ĐỐI KHÔNG sử dụng dấu ngoặc kép (") ở bên trong nội dung văn bản của các trường JSON (như body, thinking, example). Nếu cần viết tên gọi, trích dẫn hoặc ký hiệu, bắt buộc sử dụng dấu ngoặc đơn (') thay thế để tránh lỗi cú pháp JSON.


Yêu cầu về cấu trúc của mảng JSON:
1. Chia bài học thành 1 đến 3 thẻ lý thuyết (theory cards) đại diện cho các phần kiến thức quan trọng nhất của bài học.
2. Với mỗi thẻ lý thuyết, bạn cần cung cấp:
   - "thinking": Suy nghĩ và phân tích chi tiết từng bước bằng tiếng Việt để giải quyết bài toán và lập luận toán học chính xác trước khi viết định nghĩa/ví dụ. Phải tự giải nháp ví dụ và kiểm tra lại xem các số có mấy chữ số, các phép tính có đúng hay không.
   - "title": Tên thẻ lý thuyết (ngắn gọn, tập trung vào nội dung chính, ví dụ: "Khái niệm phân số", "Quy tắc quy đồng mẫu số", "Đơn vị Yến, Tạ, Tấn").
   - "body": Nội dung định nghĩa, khái niệm hoặc lý thuyết cốt lõi được trình bày mạch lạc, dễ hiểu với học sinh lớp 4. Sử dụng danh sách liệt kê để làm rõ các ý nếu cần.
   - "formula": Quy tắc, công thức tính toán hoặc ghi nhớ cốt lõi (nếu có, bắt buộc dùng LaTeX kẹp trong dấu $).
   - "example": Một ví dụ minh họa thực tế, sinh động, kèm theo lời giải chi tiết từng bước. Ví dụ và lời giải phải sử dụng công thức LaTeX để viết các phép tính và kết quả. Tránh các ví dụ mang tính lặp lại sáo rỗng hoặc quá đơn giản (ví dụ: KHÔNG dùng kiểu "tìm số tiếp theo của 1 000 000 là 1 000 001").

Định dạng trả về:
BẮT BUỘC chỉ trả về một mảng JSON duy nhất chứa danh sách các thẻ lý thuyết, không bao gồm mã code block markdown (ví dụ không bọc trong ```json ... ```) và không có bất kỳ lời thoại hay chữ nào khác ngoài JSON.

JSON Schema mẫu:
[
  {{
    "thinking": "Phân tích logic toán học cho thẻ này...",
    "title": "Tên thẻ 1",
    "body": "Nội dung lý thuyết thẻ 1...",
    "formula": "Công thức hoặc quy tắc (sử dụng LaTeX $)...",
    "example": "Ví dụ và lời giải mẫu chi tiết (sử dụng LaTeX $)..."
  }}
]"""

def build_docx(lessons_theory, output_path):
    print(f"Đang xây dựng file Word tại: {output_path}...")
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("TỔNG HỢP LÝ THUYẾT TOÁN LỚP 4\n")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(20)
    title_run.bold = True
    
    subtitle_run = title_p.add_run("Bộ Sách Giáo Khoa Cánh Diều — Gợi Ý Học Tập")
    subtitle_run.font.name = 'Arial'
    subtitle_run.font.size = Pt(12)
    subtitle_run.italic = True
    
    doc.add_paragraph("\n")
    
    # Group by Chapter
    chapters = {}
    for item in lessons_theory:
        ch_name = item.get("chapter_name", "Khác")
        if ch_name not in chapters:
            chapters[ch_name] = []
        chapters[ch_name].append(item)
        
    for ch_name, ch_lessons in chapters.items():
        # Chapter Heading
        ch_h = doc.add_heading(ch_name, level=1)
        ch_h.style.font.name = 'Arial'
        ch_h.style.font.size = Pt(16)
        ch_h.style.font.bold = True
        
        for lesson in ch_lessons:
            # Lesson Heading
            l_h = doc.add_heading(lesson["lesson_name"], level=2)
            l_h.style.font.name = 'Arial'
            l_h.style.font.size = Pt(13)
            l_h.style.font.bold = True
            
            cards = lesson.get("theory_cards", [])
            if not cards:
                p = doc.add_paragraph("Chưa có lý thuyết cho bài học này.")
                p.italic = True
                continue
                
            for idx, card in enumerate(cards, 1):
                # Card Title
                c_p = doc.add_paragraph()
                c_run = c_p.add_run(f"• Thẻ {idx}: {card.get('title', '')}")
                c_run.font.name = 'Arial'
                c_run.font.size = Pt(11.5)
                c_run.bold = True
                
                # Card Body
                body = card.get("body", "")
                body_p = doc.add_paragraph()
                body_p.paragraph_format.left_indent = Inches(0.2)
                
                if isinstance(body, list):
                    body_text = "\n".join([f"- {item}" for item in body])
                else:
                    body_text = str(body)
                    
                body_run = body_p.add_run(body_text)
                body_run.font.name = 'Arial'
                
                # Card Formula
                formula = card.get("formula", "").strip()
                if formula:
                    form_p = doc.add_paragraph()
                    form_p.paragraph_format.left_indent = Inches(0.4)
                    form_run = form_p.add_run(f"Công thức/Quy tắc: {formula}")
                    form_run.font.name = 'Arial'
                    form_run.italic = True
                    form_run.bold = True
                    
                # Card Example
                example = card.get("example", "").strip()
                if example:
                    ex_p = doc.add_paragraph()
                    ex_p.paragraph_format.left_indent = Inches(0.4)
                    ex_run = ex_p.add_run(f"Ví dụ: {example}")
                    ex_run.font.name = 'Arial'
                    ex_run.italic = True
                    
            doc.add_paragraph()  # Blank line between lessons
            
    doc.save(output_path)
    print(f"Đã lưu file Word thành công tại: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Sinh lý thuyết Toán 4 Cánh Diều bằng NVIDIA NIM API")
    parser.add_argument("--limit", type=int, default=0, help="Giới hạn số lượng bài để chạy thử nghiệm")
    parser.add_argument("--force", action="store_true", help="Chạy lại từ đầu, không dùng cache")
    args = parser.parse_args()
    
    if not API_KEY:
        print("Lỗi: Không tìm thấy NVIDIA_NIM_API_KEY trong file .env.")
        sys.exit(1)
        
    if not LESSONS_JSON.exists():
        print(f"Lỗi: Không tìm thấy file bài học tại {LESSONS_JSON}.")
        sys.exit(1)
        
    with open(LESSONS_JSON, "r", encoding="utf-8") as f:
        all_lessons = json.load(f)
        
    grade4_lessons = [l for l in all_lessons if l.get("grade") == 4]
    print(f"Đã tìm thấy {len(grade4_lessons)} bài học thuộc Lớp 4.")
    
    if args.limit > 0:
        grade4_lessons = grade4_lessons[:args.limit]
        print(f"Chế độ giới hạn: Chỉ xử lý {len(grade4_lessons)} bài học đầu tiên.")
        
    # Load crawled theory cards for context
    crawled_by_id = {}
    if CRAWLED_THEORY_JSON.exists():
        try:
            with open(CRAWLED_THEORY_JSON, "r", encoding="utf-8") as f:
                crawled_list = json.load(f)
                for item in crawled_list:
                    crawled_by_id[item["lesson_id"]] = item
            print(f"Đã tải {len(crawled_by_id)} bài học từ file lý thuyết thô cào được.")
        except Exception as e:
            print(f"Lỗi khi đọc file lý thuyết thô: {str(e)}")
            
    # Load cache
    cache = {}
    if CACHE_JSON.exists() and not args.force:
        try:
            with open(CACHE_JSON, "r", encoding="utf-8") as f:
                cached_list = json.load(f)
                for item in cached_list:
                    cache[item["lesson_id"]] = item
            print(f"Đã tải {len(cache)} bài học từ bộ nhớ đệm (cache).")
        except Exception as e:
            print(f"Lỗi khi đọc file cache, sẽ sinh mới: {str(e)}")
            
    results = []
    success_count = 0
    
    for idx, lesson in enumerate(grade4_lessons, 1):
        lesson_id = lesson["id"]
        lesson_name = lesson["lesson_name"]
        chapter_name = lesson["chapter_name"]
        
        print(f"[{idx}/{len(grade4_lessons)}] Đang xử lý: {chapter_name} -> {lesson_name}")
        
        if lesson_id in cache and not args.force:
            print("  -> Đã có trong cache. Bỏ qua.")
            results.append(cache[lesson_id])
            success_count += 1
            continue
            
        # Get raw context
        raw_context = ""
        crawled_item = crawled_by_id.get(lesson_id)
        if crawled_item:
            raw_context_parts = []
            for card in crawled_item.get("theory_cards", []):
                title = card.get("title", "")
                body = card.get("body", "")
                if body:
                    # Clean simple HTML artifacts if any
                    body_clean = re.sub(r'<[^>]+>', '', body)
                    raw_context_parts.append(f"Tiêu đề: {title}\nNội dung: {body_clean}")
            raw_context = "\n\n".join(raw_context_parts)
            
        prompt = generate_theory_prompt(4, chapter_name, lesson_name, raw_context)
        
        try:
            raw_response = call_nvidia_nim(prompt)
            theory_cards = parse_ai_response(raw_response)
            
            lesson_entry = {
                "lesson_id": lesson_id,
                "grade": 4,
                "chapter_name": chapter_name,
                "lesson_name": lesson_name,
                "theory_cards": theory_cards
            }
            
            results.append(lesson_entry)
            cache[lesson_id] = lesson_entry
            success_count += 1
            print(f"  -> Sinh thành công {len(theory_cards)} thẻ lý thuyết.")
            
            # Save cache immediately after each success
            with open(CACHE_JSON, "w", encoding="utf-8") as f:
                json.dump(list(cache.values()), f, ensure_ascii=False, indent=2)
                
        except Exception as e:
            print(f"  -> LỖI: Không thể xử lý bài học này: {str(e)}")
            if args.limit > 0:
                print("Dừng do lỗi ở chế độ giới hạn.")
                raise e
                
        time.sleep(1.0)  # Politeness delay
        
    print(f"\nĐã hoàn thành sinh lý thuyết cho {success_count}/{len(grade4_lessons)} bài học.")
    
    # Save the final JSON
    with open(CACHE_JSON, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
        
    # Build the Word Document
    build_docx(results, OUTPUT_DOCX)

if __name__ == "__main__":
    main()
