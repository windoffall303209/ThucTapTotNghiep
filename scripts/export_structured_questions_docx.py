import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import requests
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "doc"
DEFAULT_INPUT = OUTPUT_DIR / "crawled_questions_structured_source.json"
DEFAULT_MAPPING = OUTPUT_DIR / "crawled_questions_structured_mapping.json"
DEFAULT_OUTPUT = OUTPUT_DIR / "tong_hop_cau_hoi_theo_chuong_bai_moi.docx"
DEFAULT_IMAGE_DIR = OUTPUT_DIR / "images"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )
}


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def clean_multiline_text(value):
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def book_label(index_url):
    path = urlparse(index_url or "").path.lower()
    grade_match = re.search(r"toan-(\d+)", path)
    grade = grade_match.group(1) if grade_match else ""
    series = "Cánh Diều"
    if "ket-noi" in path or "-kn" in path:
        series = "Kết nối tri thức"
    tap_match = re.search(r"tap-(\d+)", path)
    suffix = f" - Tập {tap_match.group(1)}" if tap_match else ""
    return f"Toán {grade} {series}{suffix}".strip()


def risk_label(flags):
    if not flags:
        return "OK - map rõ"
    labels = {
        "unmatched": "Chưa map được",
        "low_score": "Điểm map thấp",
        "ambiguous_match": "Có nhiều bài gần giống",
        "lesson_number_mismatch": "Số bài không khớp",
    }
    return "Cần duyệt lại: " + ", ".join(labels.get(flag, flag) for flag in flags)


def image_extension(url, content_type=""):
    path_ext = Path(urlparse(url).path).suffix.lower()
    if path_ext in {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}:
        return path_ext
    if "png" in content_type:
        return ".png"
    if "jpeg" in content_type or "jpg" in content_type:
        return ".jpg"
    if "gif" in content_type:
        return ".gif"
    if "webp" in content_type:
        return ".webp"
    return ".img"


def cached_image_path(image_dir, image_url):
    prefix = hashlib.sha1(image_url.encode("utf-8")).hexdigest()
    for item in image_dir.glob(prefix + ".*"):
        return item
    return None


def download_image(session, image_dir, image_url):
    image_dir.mkdir(parents=True, exist_ok=True)
    cached = cached_image_path(image_dir, image_url)
    if cached:
        return cached

    response = session.get(image_url, headers=HEADERS, timeout=30)
    response.raise_for_status()
    content_type = response.headers.get("content-type", "").lower()
    if not content_type.startswith("image/"):
        raise ValueError(f"URL không phải ảnh: {content_type}")

    ext = image_extension(image_url, content_type)
    target = image_dir / (hashlib.sha1(image_url.encode("utf-8")).hexdigest() + ext)
    target.write_bytes(response.content)
    return target


def docx_compatible_image(image_path, converted_dir):
    if image_path.suffix.lower() != ".webp":
        return image_path
    converted_dir.mkdir(parents=True, exist_ok=True)
    target = converted_dir / (image_path.stem + ".png")
    if not target.exists():
        with Image.open(image_path) as img:
            img.save(target, "PNG")
    return target


def set_cell_text(cell, text, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(clean_text(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)


def add_info_table(document, lesson, mapping):
    rows = [
        ("Lớp", f"Lớp {lesson.get('grade', '')}"),
        ("Sách/tập", book_label(lesson.get("index_url", ""))),
        ("Chương/chủ đề nguồn", lesson.get("chapter", "")),
        ("Bài học nguồn", lesson.get("title", "")),
        ("URL bài", lesson.get("url", "")),
        ("URL mục lục", lesson.get("index_url", "")),
    ]

    if mapping:
        best = mapping.get("bestMatch") or {}
        rows.extend(
            [
                (
                    "Gợi ý chương trong hệ thống",
                    best.get("chapterName", "Chưa map được"),
                ),
                (
                    "Gợi ý bài học trong hệ thống",
                    best.get("lessonName", "Chưa map được"),
                ),
                ("Điểm map", str(best.get("score", "")) if best else ""),
                ("Trạng thái duyệt", risk_label(mapping.get("riskFlags") or [])),
            ]
        )

    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for index, (label, value) in enumerate(rows):
        set_cell_text(table.cell(index, 0), label, bold=True)
        set_cell_text(table.cell(index, 1), value)


def add_choice(document, choice):
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(f"{choice.get('key', '')}. {clean_text(choice.get('text', ''))}")
    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_explanation(document, explanation):
    text = clean_multiline_text(explanation)
    if not text:
        return

    title = document.add_paragraph()
    title_run = title.add_run("Lời giải:")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(10.5)

    for line in text.splitlines():
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        if line.startswith("Bước ") or line.startswith("Kết luận"):
            run.bold = True


def add_image(document, session, image_dir, converted_dir, image, stats):
    url = image.get("url", "")
    if not url:
        return

    try:
        raw_path = download_image(session, image_dir, url)
        image_path = docx_compatible_image(raw_path, converted_dir)
        width_px = int(image.get("width") or 0)
        width = min(5.8, max(1.2, width_px / 96)) if width_px else 4.8
        document.add_picture(str(image_path), width=Inches(width))
        last_paragraph = document.paragraphs[-1]
        last_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        caption = clean_text(image.get("alt") or "Ảnh trong đề")
        caption_paragraph = document.add_paragraph(f"Ảnh: {caption} - {url}")
        caption_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        caption_paragraph.runs[0].font.size = Pt(8)
        stats["embedded_images"] += 1
    except Exception as error:
        document.add_paragraph(f"Không nhúng được ảnh: {url} ({error})", style="List Bullet")
        stats["failed_images"] += 1


def add_question(document, session, image_dir, converted_dir, question, stats):
    number = question.get("number", "")
    text = clean_text(question.get("text", ""))
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"Câu {number}. {text}")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10.5)

    images = question.get("images") or []
    if images:
        document.add_paragraph("Ảnh trong đề:")
        for image in images:
            add_image(document, session, image_dir, converted_dir, image, stats)

    for choice in question.get("choices") or []:
        add_choice(document, choice)

    answer = clean_text(question.get("correct_answer", ""))
    if answer:
        answer_paragraph = document.add_paragraph()
        run = answer_paragraph.add_run(f"Đáp án: {answer}")
        run.bold = True
        run.font.name = "Arial"

    explanation_images = question.get("explanation_images") or []
    if explanation_images:
        document.add_paragraph("Ảnh trong lời giải:")
        for image in explanation_images:
            add_image(document, session, image_dir, converted_dir, image, stats)

    add_explanation(document, question.get("explanation", ""))


def style_document(document):
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        styles[name].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.size = Pt(13)
    styles["Heading 4"].font.size = Pt(11)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 41, 55)


def group_lessons(lessons):
    grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    for lesson in lessons:
        grade = int(lesson.get("grade") or 0)
        book = book_label(lesson.get("index_url", ""))
        chapter = clean_text(lesson.get("chapter", "Chưa phân loại"))
        grouped[grade][book][chapter].append(lesson)
    return grouped


def add_metadata(document, payload, mappings):
    stats = payload.get("stats") or {}
    document.add_heading("Tổng hợp câu hỏi crawl theo chương bài", 0)
    document.add_paragraph(f"Thời điểm tạo file: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    document.add_paragraph(f"Thời điểm crawl nguồn: {payload.get('generated_at', '')}")
    document.add_paragraph(
        "File này dùng để duyệt dữ liệu trước khi import database. "
        "Mỗi bài được ghi rõ lớp, sách/tập, chương/chủ đề nguồn, bài học nguồn và gợi ý map vào chương/bài trong hệ thống."
    )
    document.add_paragraph(
        f"Tổng bài có câu hỏi: {stats.get('lessons', 0)}; "
        f"tổng câu hỏi: {stats.get('questions', 0)}; "
        f"tổng ảnh phát hiện: {stats.get('images', 0)}."
    )
    if mappings:
        summary = mappings.get("summary") or {}
        document.add_paragraph(
            f"Mapping hệ thống: {summary.get('okLessons', 0)} bài OK; "
            f"{summary.get('reviewLessons', 0)} bài cần duyệt; "
            f"{summary.get('unmatchedLessons', 0)} bài chưa map được."
        )


def export_docx(input_path, output_path, mapping_path, image_dir, embed_images=True):
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    mappings = {}
    if mapping_path and mapping_path.exists():
        mappings = json.loads(mapping_path.read_text(encoding="utf-8"))

    document = Document()
    style_document(document)
    add_metadata(document, payload, mappings)

    grouped = group_lessons(payload.get("lessons") or [])
    image_stats = {"embedded_images": 0, "failed_images": 0}
    converted_dir = image_dir / "_converted_docx"

    with requests.Session() as session:
        for grade in sorted(grouped):
            document.add_section(WD_SECTION.NEW_PAGE)
            document.add_heading(f"Lớp {grade}", level=1)
            for book in sorted(grouped[grade]):
                document.add_heading(book, level=2)
                for chapter, lessons in grouped[grade][book].items():
                    document.add_heading(f"Chương/chủ đề nguồn: {chapter}", level=3)
                    for lesson in lessons:
                        document.add_heading(f"Bài học nguồn: {lesson.get('title', '')}", level=4)
                        mapping = (mappings.get("mappings") or {}).get(lesson.get("url", ""))
                        add_info_table(document, lesson, mapping)
                        document.add_paragraph(f"Số câu trích xuất: {len(lesson.get('questions') or [])}")
                        for question in lesson.get("questions") or []:
                            add_question(document, session, image_dir, converted_dir, question, image_stats)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "output": str(output_path),
        "lessons": payload.get("stats", {}).get("lessons", 0),
        "questions": payload.get("stats", {}).get("questions", 0),
        "images": payload.get("stats", {}).get("images", 0),
        **image_stats,
    }


def main():
    parser = argparse.ArgumentParser(description="Xuất file Word crawl theo lớp, sách, chương, bài.")
    parser.add_argument("--input", default=str(DEFAULT_INPUT))
    parser.add_argument("--mapping", default=str(DEFAULT_MAPPING))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--image-dir", default=str(DEFAULT_IMAGE_DIR))
    parser.add_argument("--no-images", action="store_true")
    args = parser.parse_args()

    result = export_docx(
        input_path=Path(args.input),
        output_path=Path(args.output),
        mapping_path=Path(args.mapping) if args.mapping else None,
        image_dir=Path(args.image_dir),
        embed_images=not args.no_images,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
