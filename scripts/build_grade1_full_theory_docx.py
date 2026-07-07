import json
import math
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
INPUT_PATH = ROOT / "content-theory" / "grade-1-theory-blueprint.json"
OUTPUT_DIR = ROOT / "output" / "doc"
IMAGE_DIR = OUTPUT_DIR / "grade1_theory_images"
OUTPUT_PATH = OUTPUT_DIR / "ly-thuyet-lop-1-codex.docx"
FALLBACK_OUTPUT_PATH = OUTPUT_DIR / "ly-thuyet-lop-1-codex-hoan-thien.docx"

FONT_REGULAR = Path("C:/Windows/Fonts/arial.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/arialbd.ttf")

TYPE_LABELS = {
    "observe": "Quan sát",
    "concept": "Nhận biết",
    "model": "Làm mẫu",
    "quick_try": "Thử nhanh",
    "remember": "Ghi nhớ",
}

INTERACTION_LABELS = {
    "none": "Không tương tác",
    "choose": "Chọn đáp án",
    "drag_drop": "Kéo thả",
    "count": "Đếm",
    "fill_blank": "Điền số/ô trống",
    "sort": "Sắp xếp",
    "match": "Nối/Ghép cặp",
}

CARD_GUIDES = {
    "observe": "Cho học sinh nhìn tranh trước, sau đó đặt một câu hỏi ngắn để các em tự phát hiện kiến thức.",
    "concept": "Gọi tên kiến thức mới bằng câu ngắn, kèm hình minh họa rõ ràng.",
    "model": "Làm mẫu một tình huống cụ thể; học sinh nhìn cách làm rồi bắt chước ở bài luyện tập.",
    "quick_try": "Cho học sinh thao tác ngay để kiểm tra xem đã hiểu ý chính của thẻ chưa.",
    "remember": "Chốt lại bằng một câu ghi nhớ ngắn, tránh giải thích dài.",
}

WRONG_HINTS = {
    "choose": "Nếu chọn sai, yêu cầu học sinh quan sát lại tranh và thử đếm/chỉ lại từng đối tượng.",
    "drag_drop": "Nếu kéo sai, đưa đối tượng về vị trí cũ và tô sáng vùng cần quan sát.",
    "count": "Nếu đếm sai, hiển thị hiệu ứng chỉ lần lượt từng đồ vật từ trái sang phải.",
    "fill_blank": "Nếu điền sai, gợi ý bằng hình còn thiếu hoặc che bớt đáp án gây nhiễu.",
    "match": "Nếu nối sai, nhắc học sinh so sánh từng cặp hình - tên một lần nữa.",
    "none": "Không cần phản hồi sai; đây là thẻ xem mẫu hoặc ghi nhớ.",
}

STATIC_ANSWERS = {
    "Trên - dưới. Phải - trái, trước - sau. Ở giữa": [
        "Sách ở trên bàn; cặp ở dưới bàn; ghế ở bên trái bạn nhỏ; bàn học và bóng ở bên phải bạn nhỏ; cây ở phía sau bạn nhỏ; xe đồ chơi ở phía trước bạn nhỏ; bạn nhỏ đứng ở giữa ghế và bàn học.",
        "Học sinh ghép đúng từng từ chỉ vị trí với đồ vật tương ứng trong tranh: trên, dưới, trái, phải, trước, sau, ở giữa.",
        "Quả bóng ở bên phải bạn nhỏ.",
    ],
    "Hình vuông - hình tròn, hình tam giác - hình chữ nhật": [
        "Học sinh chỉ đúng các đồ vật có dạng hình vuông, hình tròn, hình tam giác và hình chữ nhật.",
        "Học sinh ghép đúng tên hình với hình: vuông, tròn, tam giác, chữ nhật.",
        "Học sinh phân loại đúng mỗi hình vào nhóm tương ứng.",
    ],
    "Các số 1, 2, 3": [
        "Nhóm 1 đồ vật chọn số 1; nhóm 2 đồ vật chọn số 2; nhóm 3 đồ vật chọn số 3.",
        "1 ứng với một đồ vật; 2 ứng với hai đồ vật; 3 ứng với ba đồ vật.",
        "Học sinh đếm đúng số quả táo/đồ vật trong tranh và chọn 1, 2 hoặc 3.",
    ],
    "Các số 4, 5, 6": [
        "Nhóm 4 đồ vật chọn số 4; nhóm 5 đồ vật chọn số 5; nhóm 6 đồ vật chọn số 6.",
        "Học sinh nối đúng số 4, 5, 6 với nhóm có số lượng tương ứng.",
        "Học sinh nêu được cách làm: đếm từng đồ vật rồi chọn số đúng.",
    ],
    "Các số 7, 8, 9": [
        "Nhóm 7 đồ vật chọn số 7; nhóm 8 đồ vật chọn số 8; nhóm 9 đồ vật chọn số 9.",
        "Học sinh ghép đúng số 7, 8, 9 vào khung có số lượng chấm tương ứng.",
    ],
    "Số 0": [
        "Đĩa/hộp không còn đồ vật nào thì chọn số 0.",
        "Học sinh nhận biết số 0 dùng để chỉ không có đồ vật nào.",
    ],
    "Số 10": [
        "Học sinh đếm đủ 10 đồ vật và đọc là mười.",
        "Học sinh chọn đúng nhóm có 10 đồ vật; nhận biết số 10 gồm chữ số 1 và chữ số 0.",
    ],
    "Nhiều hơn - ít hơn - bằng nhau": [
        "Nhóm còn thừa sau khi ghép đôi là nhóm nhiều hơn; nhóm kia ít hơn.",
        "Học sinh chọn đúng từ nhiều hơn, ít hơn hoặc bằng nhau cho từng tranh.",
    ],
    "Bé hơn, dấu <, Bằng nhau, dấu =": [
        "Học sinh chọn đúng nhóm có ít đồ vật hơn là nhóm bé hơn.",
        "Ví dụ đúng: 3 < 5; 4 = 4. Học sinh kéo dấu < hoặc = đúng vào giữa hai số.",
    ],
    "Làm quen với phép cộng, dấu cộng": [
        "Gộp 2 quả táo và 1 quả táo được 3 quả táo; phép tính: 2 + 1 = 3.",
        "Học sinh hiểu dấu + là gộp thêm; đọc được phép tính 2 + 1 = 3.",
    ],
    "Làm quen với phép cộng - dấu cộng (tiếp theo)": [
        "Có 3 bông hoa, thêm 2 bông hoa, tất cả 5 bông hoa; phép tính: 3 + 2 = 5.",
        "Học sinh chọn đúng phép cộng phù hợp với tranh gộp nhóm.",
    ],
    "Phép cộng trong phạm vi 6": [
        "4 + 2 = 6.",
        "Học sinh chọn đúng tổng của các phép cộng có kết quả không quá 6.",
    ],
    "Phép cộng trong phạm vi 6 (tiếp theo)": [
        "6 gồm 4 và 2 nên 4 + 2 = 6.",
        "4 + 2 = 6, vì cần thêm 2 để đủ 6.",
    ],
    "Phép cộng trong phạm vi 10": [
        "7 + 2 = 9.",
        "Khi cộng, gộp các nhóm lại rồi đếm tất cả.",
    ],
    "Phép cộng trong phạm vi 10 (tiếp theo)": [
        "Nếu đã có 7 ô thì cần thêm 3 ô để đủ 10.",
        "Học sinh chọn đúng phép cộng theo tranh trong phạm vi 10.",
    ],
    "Khối hộp chữ nhật, khối lập phương": [
        "Hộp sữa/hộp quà giống khối hộp chữ nhật; xúc xắc/rubik giống khối lập phương.",
        "Học sinh ghép đúng đồ vật với khối hộp chữ nhật hoặc khối lập phương.",
    ],
    "Làm quen với phép trừ, dấu trừ": [
        "Có 5 quả bóng, bớt 2 quả, còn 3 quả; phép tính: 5 - 2 = 3.",
        "Học sinh hiểu dấu - là bớt đi; đọc được phép tính 5 - 2 = 3.",
    ],
    "Phép trừ trong phạm vi 6": [
        "6 - 2 = 4.",
        "Học sinh chọn đúng số còn lại sau khi bớt trong phạm vi 6.",
    ],
    "Phép trừ trong phạm vi 6 (tiếp theo)": [
        "6 tách thành 4 và 2 nên 6 - 2 = 4.",
        "6 - 2 = 4, vì số đã bớt là 2.",
    ],
    "Phép trừ trong phạm vi 10": [
        "9 - 3 = 6.",
        "Khi trừ, bớt đi rồi đếm phần còn lại.",
    ],
    "Phép trừ trong phạm vi 10 (tiếp theo)": [
        "10 - 3 = 7, vì đã bớt 3.",
        "Học sinh chọn đúng phép trừ theo tranh trong phạm vi 10.",
    ],
    "Các số đến 100": [
        "Học sinh tìm đúng số được yêu cầu trên bảng số từ 1 đến 100.",
        "Sau 28 là 29; học sinh điền đúng số còn thiếu trong dãy.",
    ],
    "Các số có hai chữ số (Từ 21 đến 40)": [
        "23 gồm 2 chục và 3 đơn vị.",
        "3 chục và 5 đơn vị là số 35.",
    ],
    "Các số có hai chữ số (Từ 41 đến 70)": [
        "58 gồm 5 chục và 8 đơn vị.",
        "Học sinh kéo chữ số vào đúng hàng chục và hàng đơn vị để tạo số.",
    ],
    "Các số có hai chữ số (Từ 71 đến 99)": [
        "76 đọc là bảy mươi sáu.",
        "8 chục và 4 đơn vị là số 84.",
    ],
    "Chục và đơn vị": [
        "Số 35 có 3 chục và 5 đơn vị.",
        "Số 62 có 6 chục và 2 đơn vị.",
    ],
    "So sánh các số trong phạm vi 100": [
        "45 < 62 vì 4 chục bé hơn 6 chục.",
        "Học sinh điền đúng dấu <, > hoặc = giữa hai số.",
    ],
    "Dài hơn, ngắn hơn": [
        "Học sinh chọn đúng vật dài hơn và vật ngắn hơn khi hai vật được đặt cùng điểm đầu.",
        "Muốn so sánh độ dài, đặt hai vật cùng một đầu rồi quan sát đầu còn lại.",
    ],
    "Đo độ dài": [
        "Học sinh đếm đúng số đơn vị đo bằng nhau đặt sát nhau dọc theo vật.",
        "Cách đo đúng là đặt các đơn vị đo sát nhau, không chồng lên nhau, không để hở.",
    ],
    "Xăng-ti-mét": [
        "Xăng-ti-mét viết tắt là cm.",
        "Nếu bút chì đặt từ vạch 0 đến vạch 6 thì bút chì dài 6 cm.",
    ],
    "Phép cộng dạng 14 + 3": [
        "14 + 3 = 17.",
        "Giữ nguyên 1 chục, cộng 4 đơn vị với 3 đơn vị được 7 đơn vị.",
    ],
    "Phép trừ dạng 17 - 2": [
        "17 - 2 = 15.",
        "Giữ nguyên 1 chục, trừ 7 đơn vị bớt 2 đơn vị còn 5 đơn vị.",
    ],
    "Cộng trừ các số tròn chục": [
        "20 + 30 = 50.",
        "60 - 20 = 40.",
    ],
    "Phép cộng dạng 25 + 14": [
        "25 + 14 = 39.",
        "32 + 16 = 48.",
    ],
    "Phép cộng dạng 25 + 4, 25 + 40": [
        "25 + 4 = 29.",
        "25 + 40 = 65.",
    ],
    "Phép trừ dạng 39 - 15": [
        "39 - 15 = 24.",
        "48 - 26 = 22.",
    ],
    "Phép trừ dạng 27 - 4; 63 - 40": [
        "27 - 4 = 23.",
        "63 - 40 = 23.",
    ],
    "Các ngày trong tuần lễ": [
        "Một tuần có 7 ngày.",
        "Sau thứ ba là thứ tư.",
    ],
    "Đồng hồ - thời gian": [
        "Kim ngắn chỉ giờ, kim dài chỉ phút.",
        "Kim dài chỉ 12, kim ngắn chỉ 7 là 7 giờ.",
    ],
}


def font(size, bold=False):
    path = FONT_BOLD if bold and FONT_BOLD.exists() else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size) if path.exists() else ImageFont.load_default()


def safe_name(text):
    allowed = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_"
    result = []
    for char in text.lower().replace(" ", "-"):
        result.append(char if char in allowed else "-")
    return "".join(result).strip("-")[:80]


def draw_wrapped(draw, text, xy, fnt, fill="#1f2937", width=42, spacing=8):
    x, y = xy
    lines = []
    for part in str(text).splitlines():
        wrapped = textwrap.wrap(part, width=width) or [""]
        lines.extend(wrapped)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + spacing
    return y


def rounded_rect(draw, box, fill, outline=None, width=2, radius=24):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def draw_badge(draw, text, x, y, fill="#dbeafe", color="#1d4ed8"):
    f = font(24, True)
    w = draw.textlength(text, font=f)
    rounded_rect(draw, (x, y, x + w + 38, y + 46), fill=fill, outline=None, radius=18)
    draw.text((x + 19, y + 9), text, font=f, fill=color)


def draw_apple(draw, cx, cy, r=26, fill="#ef4444"):
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=fill, outline="#991b1b", width=3)
    draw.line((cx, cy - r, cx + 10, cy - r - 24), fill="#7c2d12", width=5)
    draw.ellipse((cx + 10, cy - r - 28, cx + 36, cy - r - 10), fill="#22c55e", outline="#15803d", width=2)


def draw_dot(draw, cx, cy, r=18, fill="#60a5fa"):
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=fill, outline="#2563eb", width=3)


def draw_flower(draw, cx, cy, scale=1.0):
    r = int(14 * scale)
    for angle in range(0, 360, 60):
        px = cx + math.cos(math.radians(angle)) * r * 1.4
        py = cy + math.sin(math.radians(angle)) * r * 1.4
        draw.ellipse((px - r, py - r, px + r, py + r), fill="#f9a8d4", outline="#be185d", width=2)
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill="#facc15", outline="#a16207", width=2)


def draw_pencil(draw, x, y, length=110, color="#facc15"):
    draw.rounded_rectangle((x, y, x + length, y + 24), radius=8, fill=color, outline="#a16207", width=2)
    draw.polygon([(x + length, y), (x + length + 28, y + 12), (x + length, y + 24)], fill="#fed7aa", outline="#9a3412")
    draw.polygon([(x + length + 22, y + 9), (x + length + 28, y + 12), (x + length + 22, y + 15)], fill="#111827")
    draw.rectangle((x, y, x + 18, y + 24), fill="#f9a8d4", outline="#be185d")


def draw_base10(draw, x, y, tens=2, ones=5):
    for i in range(tens):
        bx = x + i * 34
        draw.rounded_rectangle((bx, y, bx + 24, y + 230), radius=7, fill="#93c5fd", outline="#1d4ed8", width=2)
        for k in range(1, 10):
            yy = y + k * 23
            draw.line((bx, yy, bx + 24, yy), fill="#1d4ed8", width=1)
    start_x = x + max(tens, 1) * 38 + 28
    for i in range(ones):
        ox = start_x + (i % 5) * 40
        oy = y + (i // 5) * 44
        draw.rounded_rectangle((ox, oy, ox + 26, oy + 26), radius=6, fill="#fde68a", outline="#a16207", width=2)


def draw_ten_frame(draw, x, y, filled=7, crossed=0):
    idx = 0
    for row in range(2):
        for col in range(5):
            box = (x + col * 70, y + row * 70, x + col * 70 + 58, y + row * 70 + 58)
            draw.rounded_rectangle(box, radius=8, fill="#ffffff", outline="#64748b", width=3)
            if idx < filled:
                draw.ellipse((box[0] + 10, box[1] + 10, box[2] - 10, box[3] - 10), fill="#60a5fa", outline="#1d4ed8", width=2)
            if idx < crossed:
                draw.line((box[0] + 8, box[1] + 8, box[2] - 8, box[3] - 8), fill="#ef4444", width=6)
                draw.line((box[0] + 8, box[3] - 8, box[2] - 8, box[1] + 8), fill="#ef4444", width=6)
            idx += 1


def draw_equation(draw, text, x, y):
    f = font(52, True)
    rounded_rect(draw, (x, y, x + 420, y + 86), fill="#ffffff", outline="#cbd5e1", radius=20)
    draw.text((x + 35, y + 14), text, font=f, fill="#111827")


def draw_clock(draw, cx, cy, hour=7):
    r = 145
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill="#ffffff", outline="#1e3a8a", width=6)
    for i in range(1, 13):
        angle = math.radians(i * 30 - 90)
        tx = cx + math.cos(angle) * (r - 34)
        ty = cy + math.sin(angle) * (r - 34)
        label = str(i)
        f = font(24, True)
        draw.text((tx - draw.textlength(label, font=f) / 2, ty - 13), label, font=f, fill="#111827")
    hour_angle = math.radians((hour % 12) * 30 - 90)
    draw.line((cx, cy, cx + math.cos(hour_angle) * 78, cy + math.sin(hour_angle) * 78), fill="#ef4444", width=8)
    draw.line((cx, cy, cx, cy - 110), fill="#2563eb", width=6)
    draw.ellipse((cx - 9, cy - 9, cx + 9, cy + 9), fill="#111827")


def draw_calendar(draw, x, y):
    days = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "CN"]
    f = font(22, True)
    for i, day in enumerate(days):
        bx = x + i * 126
        fill = "#dbeafe" if i < 5 else "#fee2e2"
        rounded_rect(draw, (bx, y, bx + 108, y + 92), fill=fill, outline="#64748b", radius=14)
        draw.text((bx + 16, y + 32), day, font=f, fill="#111827")


def draw_ruler(draw, x, y):
    draw.rounded_rectangle((x, y, x + 670, y + 76), radius=10, fill="#fef3c7", outline="#a16207", width=4)
    f = font(22, True)
    for i in range(0, 11):
        xx = x + i * 63
        draw.line((xx, y, xx, y + (50 if i % 2 == 0 else 34)), fill="#78350f", width=3)
        draw.text((xx - 7, y + 48), str(i), font=f, fill="#78350f")


def draw_number_grid(draw, x, y, start=1):
    f = font(19, True)
    n = start
    for row in range(10):
        for col in range(10):
            bx = x + col * 54
            by = y + row * 38
            fill = "#eff6ff" if n % 10 else "#dbeafe"
            draw.rectangle((bx, by, bx + 48, by + 32), fill=fill, outline="#94a3b8")
            draw.text((bx + 9, by + 5), str(n), font=f, fill="#1e293b")
            n += 1


def draw_position_scene(draw):
    draw.rectangle((110, 300, 470, 330), fill="#8b5e34")
    draw.rectangle((150, 330, 180, 480), fill="#8b5e34")
    draw.rectangle((400, 330, 430, 480), fill="#8b5e34")
    rounded_rect(draw, (215, 230, 365, 290), fill="#93c5fd", outline="#1d4ed8", radius=16)
    draw.text((240, 246), "SÁCH", font=font(28, True), fill="#1d4ed8")
    rounded_rect(draw, (250, 385, 360, 480), fill="#fca5a5", outline="#b91c1c", radius=20)
    draw.arc((268, 350, 342, 420), 180, 360, fill="#b91c1c", width=8)
    draw.text((720, 185), "TRÁI", font=font(34, True), fill="#2563eb")
    draw.text((920, 185), "PHẢI", font=font(34, True), fill="#2563eb")
    draw.line((850, 230, 700, 230), fill="#2563eb", width=7)
    draw.polygon([(690, 230), (720, 210), (720, 250)], fill="#2563eb")
    draw.line((850, 230, 1000, 230), fill="#2563eb", width=7)
    draw.polygon([(1010, 230), (980, 210), (980, 250)], fill="#2563eb")


def draw_shapes_scene(draw):
    draw.rectangle((130, 210, 280, 360), fill="#60a5fa", outline="#1d4ed8", width=5)
    draw.ellipse((360, 210, 510, 360), fill="#f87171", outline="#b91c1c", width=5)
    draw.polygon([(665, 205), (565, 365), (765, 365)], fill="#facc15", outline="#a16207")
    draw.rectangle((835, 230, 1080, 350), fill="#86efac", outline="#15803d", width=5)
    labels = [("Vuông", 155, 385), ("Tròn", 395, 385), ("Tam giác", 590, 385), ("Chữ nhật", 890, 385)]
    for label, x, y in labels:
        draw.text((x, y), label, font=font(30, True), fill="#111827")


def draw_counting_scene(draw, counts):
    colors = ["#ef4444", "#f97316", "#22c55e", "#60a5fa", "#a78bfa"]
    y_positions = [205, 330, 455]
    for row, count in enumerate(counts[:3]):
        y = y_positions[row]
        draw.text((110, y - 18), str(count), font=font(58, True), fill="#1e3a8a")
        for i in range(count):
            draw_apple(draw, 260 + i * 72, y + 8, 24, fill=colors[row % len(colors)])


def draw_comparison_scene(draw):
    for i in range(5):
        draw_apple(draw, 210 + i * 70, 260, 24, "#ef4444")
    for i in range(4):
        draw_apple(draw, 250 + i * 70, 430, 24, "#22c55e")
    draw.text((105, 250), "Nhóm A", font=font(28, True), fill="#111827")
    draw.text((105, 420), "Nhóm B", font=font(28, True), fill="#111827")
    for i in range(4):
        draw.line((210 + i * 70, 300, 250 + i * 70, 390), fill="#94a3b8", width=3)
    draw_badge(draw, "Nhóm A nhiều hơn", 730, 315, "#dcfce7", "#166534")


def draw_sign_scene(draw):
    draw_counting_scene(draw, [3, 5])
    draw.text((590, 285), "<", font=font(96, True), fill="#dc2626")
    draw.text((735, 285), "5", font=font(78, True), fill="#1e3a8a")
    draw.text((500, 425), "4 = 4", font=font(70, True), fill="#15803d")


def draw_add_scene(draw, equation="2 + 1 = 3", left=2, right=1):
    for i in range(left):
        draw_apple(draw, 170 + i * 72, 275, 25, "#ef4444")
    draw.text((350, 250), "+", font=font(82, True), fill="#2563eb")
    for i in range(right):
        draw_apple(draw, 460 + i * 72, 275, 25, "#22c55e")
    draw.text((630, 250), "=", font=font(82, True), fill="#2563eb")
    for i in range(left + right):
        draw_apple(draw, 750 + i * 58, 275, 22, "#f97316")
    draw_equation(draw, equation, 380, 405)


def draw_sub_scene(draw, equation="5 - 2 = 3", total=5, sub=2):
    for i in range(total):
        cx = 180 + i * 82
        draw_apple(draw, cx, 285, 27, "#ef4444")
        if i >= total - sub:
            draw.line((cx - 35, 250, cx + 35, 320), fill="#dc2626", width=8)
            draw.line((cx - 35, 320, cx + 35, 250), fill="#dc2626", width=8)
    draw.text((650, 250), "còn lại", font=font(42, True), fill="#111827")
    for i in range(total - sub):
        draw_apple(draw, 840 + i * 66, 285, 24, "#22c55e")
    draw_equation(draw, equation, 390, 405)


def draw_solids_scene(draw):
    draw.polygon([(170, 285), (330, 235), (450, 300), (290, 360)], fill="#93c5fd", outline="#1d4ed8")
    draw.polygon([(170, 285), (290, 360), (290, 470), (170, 390)], fill="#60a5fa", outline="#1d4ed8")
    draw.polygon([(290, 360), (450, 300), (450, 410), (290, 470)], fill="#bfdbfe", outline="#1d4ed8")
    draw.text((180, 500), "Khối hộp chữ nhật", font=font(30, True), fill="#111827")
    draw.polygon([(760, 270), (900, 220), (1040, 270), (900, 320)], fill="#fde68a", outline="#a16207")
    draw.polygon([(760, 270), (900, 320), (900, 460), (760, 400)], fill="#facc15", outline="#a16207")
    draw.polygon([(900, 320), (1040, 270), (1040, 410), (900, 460)], fill="#fef3c7", outline="#a16207")
    draw.text((780, 500), "Khối lập phương", font=font(30, True), fill="#111827")


def draw_place_value_scene(draw, number=35):
    tens = number // 10
    ones = number % 10
    draw.text((120, 165), f"Số {number}", font=font(60, True), fill="#1e3a8a")
    rounded_rect(draw, (115, 260, 360, 345), fill="#dbeafe", outline="#2563eb", radius=18)
    rounded_rect(draw, (385, 260, 650, 345), fill="#fef3c7", outline="#a16207", radius=18)
    draw.text((155, 282), f"{tens} chục", font=font(34, True), fill="#1e3a8a")
    draw.text((425, 282), f"{ones} đơn vị", font=font(34, True), fill="#92400e")
    draw_base10(draw, 735, 205, tens=min(tens, 7), ones=ones)


def draw_length_scene(draw):
    draw_pencil(draw, 190, 230, 330, "#facc15")
    draw_pencil(draw, 190, 360, 510, "#86efac")
    draw.line((180, 215, 180, 410), fill="#64748b", width=4)
    draw.text((590, 226), "ngắn hơn", font=font(36, True), fill="#92400e")
    draw.text((790, 356), "dài hơn", font=font(36, True), fill="#166534")


def draw_week_scene(draw):
    draw_calendar(draw, 135, 285)


def draw_default_scene(draw, card, lesson):
    text = card.get("visual_prompt") or lesson.get("lesson") or ""
    rounded_rect(draw, (110, 185, 1090, 500), fill="#ffffff", outline="#cbd5e1", radius=24)
    draw_wrapped(draw, text, (155, 235), font(30, True), width=48, spacing=12)


def choose_scene(draw, lesson, card):
    name = lesson["lesson"].lower()
    title = card.get("title", "").lower()
    display = card.get("display_text", "").lower()
    all_text = f"{name} {title} {display}"

    if "trên" in name or "vị trí" in all_text:
        draw_position_scene(draw)
    elif "hình vuông" in name or "hình tròn" in name:
        draw_shapes_scene(draw)
    elif "1, 2, 3" in name:
        draw_counting_scene(draw, [1, 2, 3])
    elif "4, 5, 6" in name:
        draw_counting_scene(draw, [4, 5, 6])
    elif "7, 8, 9" in name:
        draw_counting_scene(draw, [7, 8, 9])
    elif "số 0" in name:
        draw.ellipse((260, 255, 520, 385), fill="#f8fafc", outline="#64748b", width=5)
        draw.text((620, 285), "Không có đồ vật nào", font=font(42, True), fill="#111827")
        draw.text((720, 370), "0", font=font(90, True), fill="#1e3a8a")
    elif "số 10" in name:
        draw_ten_frame(draw, 175, 245, filled=10)
        draw_equation(draw, "10", 720, 295)
    elif "nhiều hơn" in name:
        draw_comparison_scene(draw)
    elif "bé hơn" in name or "dấu <" in name:
        draw_sign_scene(draw)
    elif "cộng" in name and "100" not in name:
        draw_add_scene(draw, "2 + 1 = 3", 2, 1)
    elif "trừ" in name and "100" not in name:
        draw_sub_scene(draw, "5 - 2 = 3", 5, 2)
    elif "khối" in name:
        draw_solids_scene(draw)
    elif "các số đến 100" in name:
        draw_number_grid(draw, 330, 150)
    elif "21 đến 40" in name:
        draw_place_value_scene(draw, 35)
    elif "41 đến 70" in name:
        draw_place_value_scene(draw, 58)
    elif "71 đến 99" in name:
        draw_place_value_scene(draw, 84)
    elif "chục và đơn vị" in name:
        draw_place_value_scene(draw, 62)
    elif "so sánh các số" in name:
        draw_place_value_scene(draw, 45)
        draw.text((610, 270), "<", font=font(90, True), fill="#dc2626")
        draw_place_value_scene(draw, 62)
    elif "dài hơn" in name or "ngắn hơn" in name:
        draw_length_scene(draw)
    elif "đo độ dài" in name:
        draw_ruler(draw, 205, 355)
        draw_pencil(draw, 250, 255, 420, "#86efac")
    elif "xăng-ti-mét" in name:
        draw_ruler(draw, 205, 355)
        draw.text((330, 250), "cm", font=font(86, True), fill="#1e3a8a")
    elif "14 + 3" in name:
        draw_place_value_scene(draw, 14)
        draw_equation(draw, "14 + 3 = 17", 695, 405)
    elif "17 - 2" in name:
        draw_place_value_scene(draw, 17)
        draw_equation(draw, "17 - 2 = 15", 695, 405)
    elif "tròn chục" in name:
        draw_base10(draw, 180, 215, tens=2, ones=0)
        draw.text((330, 290), "+", font=font(72, True), fill="#2563eb")
        draw_base10(draw, 430, 215, tens=3, ones=0)
        draw_equation(draw, "20 + 30 = 50", 640, 405)
    elif "25 + 14" in name:
        draw_place_value_scene(draw, 25)
        draw_equation(draw, "25 + 14 = 39", 695, 405)
    elif "25 + 4" in name or "25 + 40" in name:
        draw_place_value_scene(draw, 25)
        draw_equation(draw, "25 + 4 = 29", 695, 360)
        draw_equation(draw, "25 + 40 = 65", 695, 455)
    elif "39 - 15" in name:
        draw_place_value_scene(draw, 39)
        draw_equation(draw, "39 - 15 = 24", 695, 405)
    elif "27 - 4" in name or "63 - 40" in name:
        draw_place_value_scene(draw, 27)
        draw_equation(draw, "27 - 4 = 23", 695, 360)
        draw_equation(draw, "63 - 40 = 23", 695, 455)
    elif "ngày" in name and "tuần" in name:
        draw_week_scene(draw)
    elif "đồng hồ" in name:
        draw_clock(draw, 570, 325, hour=7)
    else:
        draw_default_scene(draw, card, lesson)


def render_card_image(chapter_index, lesson_index, card_index, chapter, lesson, card):
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1200, 650), "#f8fafc")
    draw = ImageDraw.Draw(img)
    rounded_rect(draw, (35, 35, 1165, 615), fill="#ffffff", outline="#d0d7de", radius=34)
    draw_badge(draw, TYPE_LABELS.get(card.get("type"), card.get("type", "")), 80, 75)
    draw.text((80, 135), lesson["lesson"], font=font(34, True), fill="#111827")
    draw.line((80, 178, 1120, 178), fill="#e5e7eb", width=3)
    choose_scene(draw, lesson, card)

    display = card.get("display_text", "")
    if display:
        rounded_rect(draw, (80, 535, 1120, 595), fill="#f1f5f9", outline=None, radius=18)
        draw.text((110, 550), display, font=font(28, True), fill="#0f172a")

    filename = f"g1-c{chapter_index:02d}-l{lesson_index:02d}-card{card_index:02d}-{safe_name(card.get('title', 'the'))}.png"
    path = IMAGE_DIR / filename
    img.save(path)
    return path


def set_run_font(run, size=10.5, bold=False, italic=False, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)

    for style_name, size, color in [
        ("Title", 22, "1F4E79"),
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 11, "1F1F1F"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_meta_paragraph(doc, label, value):
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, bold=True)
    value_run = paragraph.add_run(str(value or ""))
    set_run_font(value_run)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run)


def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text or ""))
    set_run_font(run, size=size, bold=bold)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_overview_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["STT", "Chủ đề", "Số bài", "Số thẻ"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def build_card_detail(card):
    card_type = card.get("type", "")
    interaction = card.get("interaction", "none")
    return [
        f"Nội dung hiển thị: {card.get('display_text', '')}",
        f"Cách triển khai: {CARD_GUIDES.get(card_type, '')}",
        f"Nhiệm vụ học sinh: {card.get('student_task', '')}",
        f"Tương tác đề xuất: {INTERACTION_LABELS.get(interaction, interaction)}",
        f"Gợi ý khi học sinh sai: {WRONG_HINTS.get(interaction, '')}",
    ]


def expected_answer(lesson, card_index):
    answers = STATIC_ANSWERS.get(lesson["lesson"], [])
    if card_index - 1 < len(answers):
        return answers[card_index - 1]
    if answers:
        return answers[-1]
    return "Học sinh hoàn thành đúng nhiệm vụ của thẻ theo tranh minh họa."


def theory_takeaway(lesson, card):
    card_type = card.get("type", "")
    if card_type == "observe":
        return "Quan sát tranh để tự nhận ra đặc điểm hoặc tình huống toán học."
    if card_type == "concept":
        return "Gọi tên đúng kiến thức mới và liên hệ với hình minh họa."
    if card_type == "model":
        return "Làm theo mẫu: nhìn tranh, xác định dữ kiện, rồi chọn phép tính/kết quả."
    if card_type == "quick_try":
        return "Tự thao tác một bước ngắn để kiểm tra đã hiểu kiến thức vừa học."
    if card_type == "remember":
        return "Ghi nhớ câu chốt ngắn để dùng khi làm bài luyện tập."
    return "Nắm ý chính của thẻ học."


def build_docx():
    data = json.loads(INPUT_PATH.read_text(encoding="utf-8"))
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    image_paths = {}
    for chapter_index, chapter in enumerate(data["chapters"], 1):
        for lesson_index, lesson in enumerate(chapter["lessons"], 1):
            for card_index, card in enumerate(lesson.get("cards", []), 1):
                key = (chapter_index, lesson_index, card_index)
                image_paths[key] = render_card_image(chapter_index, lesson_index, card_index, chapter, lesson, card)

    doc = Document()
    set_document_defaults(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("LÝ THUYẾT TOÁN LỚP 1")
    set_run_font(title_run, size=24, bold=True, color="1F4E79")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Bản học liệu đầy đủ có tranh minh họa - Codex")
    set_run_font(subtitle_run, size=13, italic=True)

    doc.add_paragraph()
    add_meta_paragraph(doc, "Nguồn tham khảo", data.get("source_textbook", "SGK Toán 1 Cánh Diều"))
    add_meta_paragraph(doc, "Phạm vi", "Lớp 1 - cấp Tiểu học")
    add_meta_paragraph(doc, "Ghi chú", "Tranh minh họa trong tài liệu là hình vẽ gốc được tạo riêng cho bản nháp, không sao chép tranh trong SGK.")

    profile = data["cognitive_profile"]
    doc.add_heading("Định hướng thiết kế", level=1)
    add_bullets(
        doc,
        [
            "Mỗi bài được chia thành các thẻ học ngắn để học sinh quan sát, nhận biết, làm mẫu và thử nhanh.",
            "Mỗi thẻ có tranh minh họa, câu hiển thị ngắn, nhiệm vụ học sinh và gợi ý xử lý khi học sinh làm sai.",
            "Nội dung phù hợp học sinh lớp 1: ít chữ, nhiều hình, ưu tiên đếm và thao tác trực quan.",
            profile["ai_policy"],
        ],
    )

    doc.add_heading("Tổng quan nội dung", level=1)
    overview_rows = []
    for index, chapter in enumerate(data["chapters"], 1):
        lesson_count = len(chapter["lessons"])
        card_count = sum(len(lesson["cards"]) for lesson in chapter["lessons"])
        overview_rows.append([index, chapter["title"], lesson_count, card_count])
    add_overview_table(doc, overview_rows)

    for chapter_index, chapter in enumerate(data["chapters"], 1):
        doc.add_page_break()
        doc.add_heading(f"Chủ đề {chapter_index}: {chapter['title']}", level=1)
        for lesson_index, lesson in enumerate(chapter["lessons"], 1):
            doc.add_heading(f"{chapter_index}.{lesson_index}. {lesson['lesson']}", level=2)
            add_meta_paragraph(doc, "Mục tiêu bài học", lesson.get("objective", ""))
            add_meta_paragraph(doc, "Cách học phù hợp", "Học sinh quan sát tranh, thao tác theo yêu cầu ngắn, sau đó làm bài thử nhanh ngay trên thẻ.")

            for card_index, card in enumerate(lesson.get("cards", []), 1):
                doc.add_heading(
                    f"Thẻ {card_index}: {TYPE_LABELS.get(card.get('type'), card.get('type', ''))} - {card.get('title', '')}",
                    level=3,
                )
                picture_paragraph = doc.add_paragraph()
                picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = picture_paragraph.add_run()
                run.add_picture(str(image_paths[(chapter_index, lesson_index, card_index)]), width=Inches(6.7))
                add_meta_paragraph(doc, "Kiến thức cần rút ra", theory_takeaway(lesson, card))
                add_bullets(doc, build_card_detail(card))
                add_meta_paragraph(doc, "Đáp án/Kết quả mong đợi", expected_answer(lesson, card_index))
                add_meta_paragraph(doc, "Mô tả tranh", card.get("visual_prompt", ""))
            doc.add_paragraph()

    doc.add_page_break()
    doc.add_heading("Gợi ý chuẩn hóa để đưa vào hệ thống", level=1)
    add_bullets(
        doc,
        [
            "Mỗi thẻ nên được lưu với các trường: type, title, display_text, image_url, student_task, interaction, wrong_hint.",
            "Tranh minh họa nên được thay dần bằng bộ asset thống nhất về phong cách khi làm giao diện chính thức.",
            "Với lớp 1, không nên mở Chat AI tự do; chỉ dùng phản hồi cố định, ví dụ: Đếm lại từng hình, Quan sát nhóm bên trái, Thử chọn dấu khác.",
            "Sau khi duyệt bản Word này, có thể xuất lại thành JSON import vào bảng Lessons.theory_cards.",
        ],
    )

    try:
        doc.save(OUTPUT_PATH)
        return OUTPUT_PATH
    except PermissionError:
        doc.save(FALLBACK_OUTPUT_PATH)
        return FALLBACK_OUTPUT_PATH


if __name__ == "__main__":
    print(build_docx())
